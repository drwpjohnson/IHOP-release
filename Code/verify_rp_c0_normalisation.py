"""verify_rp_c0_normalisation.py -- confirm that the retention-profile error bars in
UnfavorableMaster.xlsx are computed with each replicate normalised by its own C0.

WHY THIS EXISTS.  Replicates at one nominal condition were run at genuinely different injection
concentrations -- C0 spans about a factor of six across the six replicated conditions.  Retained
numbers scale with injected numbers, so a column run at six times the concentration holds about six
times as many colloids at every depth.  A raw standard deviation across replicates would therefore
mix that intended difference into an error bar meant to show experimental scatter.

unfav_master_fit.py ALREADY handles this (see `normalize=(expkey=='rp_exp')` in its per-condition
sheet writer): each replicate is divided by its own C0 and rescaled by the group's geometric-mean
C0 before the mean and standard deviation are taken, and the sheet header records the C0avg used.
The breakthrough curve needs no such treatment, since C/C0 has already divided C0 out.

This script recomputes both versions straight from the tidy CSV and checks the master against them.
It was written after a mistaken claim that the figures were showing RAW scatter; the check refuted
that in one run, and is kept so the question does not have to be re-opened from memory.

WHAT IT REPORTS, per replicated study-condition and depth:
  master sd  -- as written in UnfavorableMaster.xlsx
  sd norm    -- recomputed with C0 normalisation  (should equal master sd)
  sd raw     -- recomputed without it             (what the bars would be if unnormalised)

Exit status is 1 if any master sd departs from the normalised recomputation by more than 1e-3 dex.

Usage:
    python3 verify_rp_c0_normalisation.py            # summary only
    python3 verify_rp_c0_normalisation.py --detail   # every depth
"""
import csv, collections, argparse, sys
import numpy as np
import openpyxl

CSV = "../Data/LiTong_experimental_data_tidy.csv"
MASTER = "../Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx"
RP_HEADER_ROW = 101          # per-condition sheets: BTEC block at row 5, RP block at row 101
TOL = 1e-3


def sizeclass(s):
    s = float(s)
    return 1.1 if 0.9 <= s <= 1.15 else round(s, 2)


def load():
    rp = collections.defaultdict(list); meta = {}
    for r in csv.DictReader(open(CSV)):
        if r["curve"] != "RP" or r["chemistry"] != "unfavorable":
            continue
        k = (r["source"], r["medium"], r["workbook_col"])
        rp[k].append((float(r["x"]), float(r["value"])))
        meta[k] = dict(medium=r["medium"], size=float(r["colloid_um"]),
                       vel=float(r["velocity_mday"]),
                       IS=(None if r["IS_mM"] == "" else float(r["IS_mM"])),
                       C0=float(r["C0_per_mL"]))
    for k in rp:
        rp[k] = np.array(sorted(rp[k]))
    return rp, meta


def master_block(wb, med, size, vel, IS, study):
    """(x, mean, sd) as written in the master's per-condition sheet, or None."""
    name = f"{'gl' if med == 'glass' else 'qu'}_{sizeclass(size)}um_{vel:.0f}md_{IS:.0f}mM"
    if name not in wb.sheetnames:
        return None
    ws = wb[name]
    cx = cy = cs = None
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(RP_HEADER_ROW, c).value or "")
        if h.strip() == f"{study} x":
            cx = c
        elif h.startswith(f"{study} mean log10"):
            cy = c
        elif h.strip() == f"{study} sd dex" and cy is not None and c > cy:
            cs = c
    if not (cx and cy and cs):
        return None
    out = []
    for r in range(RP_HEADER_ROW + 1, ws.max_row + 1):
        x, y, s = ws.cell(r, cx).value, ws.cell(r, cy).value, ws.cell(r, cs).value
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            out.append((x, y, s if isinstance(s, (int, float)) else np.nan))
        elif out:
            break
    return np.array(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--detail", action="store_true", help="print every depth, not just medians")
    args = ap.parse_args()

    rp, meta = load()
    wb = openpyxl.load_workbook(MASTER, data_only=True)

    grp = collections.defaultdict(list)
    for k, m in meta.items():
        if m["IS"] is not None and m["IS"] <= 1.0:
            continue
        grp[(k[0], m["medium"], sizeclass(m["size"]), m["vel"], m["IS"])].append(k)
    reps = {g: sorted(ks, key=lambda k: k[2]) for g, ks in grp.items() if len(ks) > 1}

    worst = 0.0; rows = []
    for g, ks in sorted(reps.items()):
        lab = f"{g[0]} {g[1]} {g[2]}um {g[3]:.0f}md {g[4]:.0f}mM"
        blk = master_block(wb, g[1], g[2], g[3], g[4], g[0])
        if blk is None:
            print(f"  (no master mean block) {lab}"); continue
        # the master interpolates every replicate onto the FIRST column's depths
        xref = rp[ks[0]][:, 0]
        C0avg = float(np.exp(np.mean(np.log([meta[k]["C0"] for k in ks]))))
        lin_raw = np.vstack([10 ** np.interp(xref, rp[k][:, 0], rp[k][:, 1]) for k in ks])
        lin_nrm = np.vstack([10 ** np.interp(xref, rp[k][:, 0], rp[k][:, 1]) / meta[k]["C0"] * C0avg
                             for k in ks])
        sd_raw = np.log10(lin_raw).std(0, ddof=1)
        sd_nrm = np.log10(lin_nrm).std(0, ddof=1)
        n = min(len(blk), len(xref))
        d = np.abs(blk[:n, 2] - sd_nrm[:n])
        worst = max(worst, float(np.nanmax(d)))
        rows.append((lab, len(ks), float(np.median(blk[:n, 2])), float(np.median(sd_nrm[:n])),
                     float(np.median(sd_raw[:n])), float(np.nanmax(d)),
                     float(np.log10(max(meta[k]["C0"] for k in ks) / min(meta[k]["C0"] for k in ks)))))
        if args.detail:
            print(f"\n{lab}   C0avg = {C0avg:.3g}")
            print(f"{'depth (m)':>10}{'master sd':>11}{'sd norm':>10}{'sd raw':>9}")
            for i in range(n):
                print(f"{xref[i]:10.4f}{blk[i,2]:11.4f}{sd_nrm[i]:10.4f}{sd_raw[i]:9.4f}")

    print(f"\n{'condition':32}{'n':>3}{'master sd':>11}{'sd norm':>10}{'sd raw':>9}"
          f"{'|max diff|':>12}{'C0 spread':>11}")
    for lab, n, ms, sn, sr, mx, c0 in rows:
        print(f"{lab:32}{n:>3}{ms:11.3f}{sn:10.3f}{sr:9.3f}{mx:12.2e}{c0:11.2f}")
    med = lambda i: float(np.median([r[i] for r in rows]))
    print(f"\n  median across conditions: master {med(2):.3f} dex, recomputed normalised "
          f"{med(3):.3f} dex, unnormalised would be {med(4):.3f} dex")
    print(f"  largest departure of master from the normalised recomputation: {worst:.2e} dex")
    if worst > TOL:
        print(f"  FAIL: exceeds {TOL} dex -- the master is NOT C0-normalising as expected")
        sys.exit(1)
    print("  PASS: the master's RP error bars are C0-normalised.")


if __name__ == "__main__":
    main()
