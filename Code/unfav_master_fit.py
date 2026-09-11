"""unfav_master_fit.py -- fit Serial-3 interception-history model to EVERY unfavorable BTEC/RP column in the
cleaned tidy CSV, one column at a time, and assemble the single unfavorable master table + a plottable
worksheet per physical condition.

Method (W.P.J., 2026-08):
  * Fit each unfavorable column separately (its own C0). Replicates are averaged at the PARAMETER level
    (mean over the columns), NOT by averaging curves -- so no C0-normalization is needed.
  * r_s is held FIXED at the favorable interception rate (chemistry-independent, grounded in CFT). The
    +-0.6 dex band was dropped: with a_s free, the RP level fixes only the product r_s*a_s, so a banded r_s
    slides across the band and absorbs the IS-dependence (defeating chemistry-independence). Fixing r_s makes
    a_s/a_m/f_x/k_r carry all the chemistry (clean monotonic a_s(IS)). r_s source per condition:
      - fitted favorable r_s /m (Favorable_data_and_sim.xlsx) where a favorable full-model fit exists;
      - favorable RP-slope seed for the c<QL small-glass columns (0.1/0.2 um: no full fit possible);
      - TE (Tufenkji-Elimelech 2004) x 0.77 ONLY where no favorable data exist (glass 0.5/8, quartz 0.5/4),
        flagged with '*'.
  * FREE (4 params): a_s (via 1-a_s), a_m, f_x, k_r. v_ns pinned 5%.
  * BTEC floored at log10 C/C0 = -6 (detection limit) for both data and model.
  * DROP IS <= 1 mM (retention gradient below model resolution) and (upstream) downgradient / DI / perturbation.
  * RP-SHAPE INLET WINDOW IS BRANCH-AWARE (2026-08-24, W.P.J.; canonical Records/plateau_rs_decision.md s6).
    (a) CLASSIFY each column from its FIRST TWO measured RP points -- point2 > point1 -> peaked.
        This is Al-Zghoul (2025) Eqn 15's inlet-slope criterion applied to the DATA, not to fitted alphas.
        It agrees with an argmax over all 10 points on 28 of 29 columns; the one exception (Tong.B,
        quartz 0.5 um 20 mM) has pt2-pt1 = -0.004 and an argmax only +0.013 log above pt1 -- a flat inlet,
        not a peak, which the two-point test correctly declines to call peaked.
    (b) A CONDITION takes ONE branch by MAJORITY over its replicate columns; ties -> non-peaking. The only
        split case is Li quartz 1.1 um 4 m/day 20 mM (M/P/S -> non-peaking, non-peaking, peaked; S rises
        by just +0.038 log) -> majority non-peaking.
    (c) SCORE with the half-split if non-peaking, the first NIN_PEAKED=4 points if peaked. The half
        split spans x = 1-11 cm here and STRADDLES an interior peak at 3-7 cm, averaging rise against
        fall and penalising a model that reproduces the peak. General rule: the window must focus on
        where the change is most dramatic; NIN_PEAKED=4 is specific to this 10-point / 2 cm grid.
    ** Weighted costs are comparable only WITHIN one window convention -- see the Fit quality sheet. **
  * Objective: RP level (W_RP) + two-segment RP slope (W_SHAPE, branch-aware window) + plateau MEAN-log target over 1.2-4 PV
    with +-0.25 dex tolerance (W_PLAT) + tail level (W_TAIL) + tail slope (W_TSLOPE). RP amplitude FIXED by
    C0 (never floated), per data_inventory s0. NOTE: C0 corrected for R & O (quartz 0.5/8/50) to 8.2e6
    (recorded 3.36E5 was the same bad placeholder flagged for O; R=O = same run). A +-0.25 dex r_s band was
    explored to improve B/U and the RP inlets but REJECTED (it dragged glass plateaus low and reopened the
    quartz r_s<->IS degeneracy); r_s stays FIXED. Steep glass RP inlets (hyperexponential) and sharp quartz
    RP peaks are a single-population limitation, not fixable by weighting -- see Records/plateau_rs_decision.md.

Workbook (UnfavorableMaster.xlsx):
  * 'Master table'  -- one row per (study/author, condition): mean params, numbers-only, units in header row,
    replicate columns listed as letters in one cell, r_s (fixed; +* if TE-sourced).
  * 'Per-column fits' -- every column's individual params + r_s source + RP/shelf+tail RMS (the spread).
  * one sheet per condition (medium/size/velocity/IS, ACROSS authors) -- plottable BTEC & RP data + fits for
    all its columns (both studies and all replicates), so data and model overlay like the diagnostic figure.

Run from Code/ with the cleaned CSV in ../Data/ (see extract_tidy_data.py). Canonical data description: Records/data_inventory.md.
"""
import csv, numpy as np, collections, warnings, functools
warnings.filterwarnings("ignore"); print=functools.partial(print,flush=True)
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
import openpyxl; from openpyxl.styles import Font
CSV="../Data/LiTong_experimental_data_tidy.csv"
DAY=86400.; L=0.2; REV,T0,Vref=22.801836559387397,3.58,0.1667; INJPV=T0/(L/Vref)
KR_MED={"glass":5.58e-5,"quartz":1.87e-5}  # MULTISTART SEED ONLY (k_r is fitted free); values superseded by 3.75e-5/1.82e-5 -- deliberately NOT refreshed, since changing a seed perturbs every converged fit for no gain
THETA={"glass":0.375,"quartz":0.36}
W_RP,W_SHAPE,W_PLAT,W_TAIL,W_TSLOPE=6.0,2.5,20.0,3.0,16.0  # W_PLAT = weight on plateau MEAN-log target (1.2-4 PV)
PLAT_TOL=0.25  # mean-centered plateau tolerance (log dex): penalize only when |model_mean - data_mean| > PLAT_TOL (avoids over-fitting the f_x<->k_r degeneracy)
kB=1.381e-23;gg=9.806;RHO_C=1055.;RHO_F=998.;MU=9.8e-4;TK=298.2;DP=0.510e-3;HAM={"glass":7.17e-21,"quartz":1.96e-20}
def TE_rs(size_um,vel,med):
    dc=size_um*1e-6;U=vel/DAY;th=THETA[med];Ui=U*th;gam=(1-th)**(1/3);Hm=HAM[med]
    As=2*(1-gam**5)/(2-3*gam+3*gam**5-2*gam**6);D=kB*TK/(6*np.pi*MU*(dc/2))
    NR=dc/DP;NvdW=Hm/(kB*TK);NG=2*(dc/2)**2*(RHO_C-RHO_F)*gg/(9*MU*Ui);NA=Hm/(12*np.pi*MU*(DP/2)**2*Ui);NPe=Ui*DP/D
    TE=2.4*As**(1/3)*NR**(-0.081)*NPe**(-0.715)*NvdW**0.052+0.55*As*NR**1.675*NA**0.125+0.22*NR**(-0.24)*NG**1.11*NvdW**0.053
    return -1.5*(gam/DP)*np.log(1-TE)*0.77
def sizeclass(s): s=float(s); return 1.1 if 0.9<=s<=1.15 else round(s,2)
# ---- fitted-favorable r_s /m (from the favorable-condition fits) by
# (medium,sizeclass,vel). glass 1.1/4 was a mean(Tong AB 30.2, Li E 28.8) computed before Tong AB and Li E were
# known to be the SAME experiment; now that Li E's corrected r_s (30.2) equals Tong AB's exactly, it's one value.
FAVFIT={("glass",0.5,4.0):47.5,("glass",1.1,4.0):30.2,("glass",2.0,8.0):15.2,("glass",1.1,2.0):49.9,("glass",1.1,8.0):29.6,
        ("quartz",0.5,8.0):56.7,("quartz",1.1,2.0):120.6,("quartz",1.1,4.0):62.3,("quartz",1.1,8.0):68.9}  # quartz 0.5/8 anchor = O favorable r_s (56.7 at corrected C0 8.2e6; ~unchanged from 56.9 -- r_s set by RP slope, C0-independent)
SEED_CQL={("glass",0.2,4.0):49.0,("glass",0.1,8.0):62.4,("glass",0.2,8.0):39.1}  # favorable RP-slope seed, c<QL (no full fit)
def seed_rs(med,size,vel):
    k=(med,sizeclass(size),vel)
    if k in FAVFIT: return FAVFIT[k],"fav"
    if k in SEED_CQL: return SEED_CQL[k],"favRP"
    return TE_rs(size,vel,med),"TE"
class Eng:
    def __init__(s,vmday):s.v=vmday/DAY;s.dx=L/90;s.dt=s.dx/s.v;s.pv=L/s.v
    def run(s,k1,a_s,a_m,fx,kr,vns=0.05,tot=10):
        kf=k1;k2=fx*kf;kmw=a_m*kf;kmg=a_m*vns*kf;G=np.zeros((4,4))
        G[0,0]-=kf;G[3,0]+=a_s*kf;G[1,0]+=(1-a_s)*kf;G[1,1]-=(kmw+k2);G[3,1]+=kmw;G[2,1]+=k2;G[2,2]-=kmg;G[3,2]+=kmg;G[1,2]+=kr;G[2,2]-=kr
        E=expm(G*s.dt);D=s.v*L/150;r=D*s.dt/s.dx**2;mn=(1+2*r)*np.ones(90);of=-r*np.ones(89);mn[0]=1+r;mn[-1]=1+r
        lu=splu(csc_matrix(diags([of,mn,of],[-1,0,1],format="csc")))
        nt=int(round(tot*90));ti=INJPV*s.pv;c=vns;Y=np.zeros((4,90));C=np.zeros(nt);t=0.
        for i in range(nt):
            Y=E@Y;C[i]=Y[0,-1]+Y[1,-1]+c*Y[2,-1];Y[0,1:]=Y[0,:-1];Y[0,0]=0;Y[1,1:]=Y[1,:-1];Y[1,0]=0
            ym=Y[2].copy();Y[2,1:]=ym[1:]-c*(ym[1:]-ym[:-1]);Y[2,0]=ym[0]*(1-c);t+=s.dt
            if t<ti:Y[0,0]+=1.
            Y[0,:]=lu.solve(Y[0,:]);Y[1,:]=lu.solve(Y[1,:])
        tp=(np.arange(nt)+1)*s.dt/s.pv;x=(np.arange(90)+0.5)*s.dx
        # ---- RETENTION PROFILE = ACCUMULATED solid phase at EXCISION (W.P.J., 2026-08-25) ----
        # Was the Johnson 2018 eq (4) injection-window RATE SNAPSHOT:
        #     rp=(a_s*kf*Yst[0,:]+kmw*Yst[1,:]+kmg*Yst[2,:])       # <-- RETIRED, invalid
        # The columns are excised at 10 PV. Injection is 2.984 PV, so 7.016 PV -- 70% of the
        # experiment -- elapses after that snapshot, and eq (4) accounts for none of it. What is
        # measured is the ACCUMULATED attached population, which is Y[3] at the end of the run.
        # Dividing by ti keeps the units the amplitude anchor logK expects, and makes this
        # identical to Code/HYDEQ/hydeq_engine.py IhopReferenceEngine's `rp` (verified to 0.000e+00).
        #   Why eq (4) looked acceptable for so long: at the MEASURED depths the two differ by
        #   <=0.033 log, essentially all at the first point, and on the column-integrated total by
        #   <=0.010 log -- assuming steady deposition across the whole injection overestimates by
        #   about what ignoring post-injection deposition underestimates. The section-5 level-and-
        #   tilt check measured exactly the quantities that survive that cancellation.
        #   Cost of the switch, measured over all 29 columns (Code/accum_refit_check.py):
        #   a_s moves <=0.090 dex, a_m <=0.180 dex, RP peak depth moves >0.5 cm on NO column. The
        #   only large shifts are f_x (2.6 dex) and k_r (3.8 dex) on Li.M, whose cost is unchanged
        #   to 2 dp -- that is the known {f_x, v_ns, k_r} degeneracy on a column whose k_r rails,
        #   not a real change.
        rp=(Y[3,:]/ti);m=(tp>0.5*INJPV)&(tp<INJPV)
        return dict(tp=tp,C=np.maximum(C,1e-300),x=x,rp=np.maximum(rp,1e-300),plat=float(C[m].mean()))
NIN_PEAKED=4  # inlet points used to SCORE a peaked profile -- see Records/plateau_rs_decision.md s6
def sl(x,y,nin=None):
    """Inlet and outlet RP log-slopes.  nin=None -> the historical half-split (non-peaking profiles).
    nin=k     -> inlet regressed over the first k points, outlet over the remainder (peaked profiles).
    The half-split spans x = 1-11 cm on this grid, which STRADDLES an interior peak at 3-7 cm and
    averages the rise against the fall, penalising a model that reproduces the peak.  General rule
    (W.P.J.): the scoring window must focus on where the change is most dramatic.  NIN_PEAKED=4 is
    specific to THIS grid (10 points at 2 cm) and must be re-derived for other data densities."""
    if nin is None:
        h=len(x)//2; return float(np.polyfit(x[:h+1],y[:h+1],1)[0]),float(np.polyfit(x[h:],y[h:],1)[0])
    return float(np.polyfit(x[:nin],y[:nin],1)[0]),float(np.polyfit(x[nin-1:],y[nin-1:],1)[0])
def rp_branch(rlog):
    """CLASSIFY the branch from the FIRST TWO measured points -- Al-Zghoul (2025) Eqn 15's inlet-slope
    criterion applied to the DATA, not to fitted alphas.  point2 > point1 -> peaked.  Agrees with
    an argmax over all 10 points on 28 of 29 unfavorable columns; the exception is Tong.B (quartz 0.5 um
    20 mM), where pt2-pt1 = -0.004 but argmax sits +0.013 log above pt1 -- a FLAT inlet, not a peak, which
    the two-point test correctly declines to call peaked.  Local by construction.  Do NOT use
    these two points as the SCORING window: tested, it chases one point-pair difference and degrades
    every column (glass 6 mM RP RMS 0.051 -> 0.113)."""
    return "peaked" if rlog[1]>rlog[0] else "non-peaking"
def condition_branch(branches):
    """A CONDITION gets ONE branch: the MAJORITY over its replicate columns (W.P.J., 2026-08-24).
    Ties -> 'non-peaking' (conservative; preserves the historical half-split).  Rationale: replicates of
    one physical condition cannot be on different branches, and averaging parameters fitted under two
    different windows would be meaningless.  The only split case in this dataset is Li quartz 1.1 um
    4 m/day 20 mM (M/P/S): M and P are non-peaking, S rises by only +0.038 log between its first two
    points -- an order of magnitude below every other peaked column, and plausibly inside the
    measurement scatter that the tidy CSV does not carry.  Majority therefore calls it non-peaking."""
    nm=sum(1 for b in branches if b=="peaked")
    return "peaked" if nm>len(branches)-nm else "non-peaking"
def ts(pv,lc,lo=6,hi=10):
    m=(np.asarray(pv)>=lo)&(np.asarray(pv)<=hi);return float("nan") if np.sum(m)<2 else float(np.polyfit(np.asarray(pv)[m],np.asarray(lc)[m],1)[0])
# ---- read CSV ----
rows=list(csv.DictReader(open(CSV)))
# GUARD, 2026-09-01: the shared master CSV was silently overwritten with a 352-row/Li-only a_mg-probe
# subset for a month (2026-08-29 to 09-01) before anyone noticed -- see Records/CLAUDE.md. Fail loudly
# instead of fitting a partial dataset in silence. A genuinely smaller, deliberate working set belongs
# under its own filename, never this shared path.
_srcs=set(r['source'] for r in rows)
assert len(rows)>=1200 and 'Li' in _srcs and 'Tong' in _srcs, (
    f"CSV at {CSV!r} looks like a subset ({len(rows)} rows, sources={_srcs}), not the full master "
    "(~1529 rows, both Li and Tong) -- refusing to fit a silently-truncated dataset.")
cols=collections.defaultdict(lambda:{"BTEC":[],"RP":[]}); meta={}
# BTEC_ND (added to the CSV 2026-08-28, task #32, wired in here 2026-09-01): non-detect tail points, real
# PV, value = the already-substituted 0.5*QL = 5e-6 floor (log10 -5.301). Merged straight into the BTEC
# array -- not a separate bucket -- because they carry a real PV and a defensible value, and because the
# fit already floors every BTEC point at -6 (`bt[:,1]=np.maximum(bt[:,1],-6.0)` below) rather than
# treating detection-limited data specially, so this is the SAME convention, not a new one. Their PV
# cutoff (>=4.5, extract_tidy_data.py) sits inside the tail_level/tail_slope window (pv>4.2) this script
# already scores, so they just extend the tail with real points instead of leaving it short.
for r in rows:
    curve='BTEC' if r['curve']=='BTEC_ND' else r['curve']
    k=(r['source'],r['medium'],r['workbook_col']); cols[k][curve].append((float(r['x']),float(r['value'])))
    meta[k]=dict(medium=r['medium'],chem=r['chemistry'],size=float(r['colloid_um']),vel=float(r['velocity_mday']),
                 IS=(None if r['IS_mM']=='' else float(r['IS_mM'])),C0=float(r['C0_per_mL']))
for k in cols:
    for cu in ("BTEC","RP"): cols[k][cu]=np.array(sorted(cols[k][cu])) if cols[k][cu] else np.zeros((0,2))
# ---- fit one unfavorable column ----
def fit_col(k,nin_override='auto'):
    m=meta[k]; bt=cols[k]['BTEC']; rp=cols[k]['RP']
    if len(bt)<3 or len(rp)<3: return None
    med=m['medium'];size=m['size'];vel=m['vel'];C0=m['C0'];theta=THETA[med];V_MS=vel/DAY;eng=Eng(vel);krs=np.log10(KR_MED[med])
    logK=np.log10(REV*T0*theta*C0*Vref); rs0,src=seed_rs(med,size,vel); kk=rs0*V_MS   # r_s FIXED at favorable anchor
    bt=bt.copy(); bt[:,1]=np.maximum(bt[:,1],-6.0)   # -6 = detection floor for log10 C/C0 (W.P.J.)
    pv=bt[:,0];lc=bt[:,1];rx=rp[:,0];rlog=rp[:,1]
    branch=rp_branch(rlog)                       # this COLUMN's own two-point classification
    if nin_override=='auto': nin=NIN_PEAKED if branch=='peaked' else None
    else: nin=nin_override                       # condition-level majority window (see condition_branch)
    si,so=sl(rx,rlog,nin)
    wm=(pv>1.2)&(pv<4.0);pvw=pv[wm] if wm.sum() else pv;dmean=float(lc[wm].mean()) if wm.sum() else float(lc.mean());tm=pv>4.2;sld=ts(pv,lc)
    def resid(p):
        a_s=1-10**p[0];am=10**p[1];fx=10**p[2];krr=10**p[3];r=eng.run(kk,a_s,am,fx,krr)
        rl=np.interp(rx,r['x'],np.maximum(r['rp'],1e-300))/V_MS;lS=logK+np.log10(np.maximum(rl,1e-300))
        rpx=W_RP*(lS-rlog);sm,sn=sl(rx,lS,nin);shp=W_SHAPE*np.array([sm-si,sn-so])
        mlog=np.mean(np.log10(np.maximum(np.interp(pvw,r['tp'],r['C']),1e-6)));dev=mlog-dmean;pl=W_PLAT*np.array([np.sign(dev)*max(0.,abs(dev)-PLAT_TOL)])  # plateau: mean-log target over 1.2-4 PV, +-PLAT_TOL tolerance
        tl=W_TAIL*(np.log10(np.maximum(np.interp(pv[tm],r['tp'],r['C']),1e-6))-lc[tm]) if tm.sum() else np.array([0.])
        slm=ts(r['tp'],np.log10(np.maximum(r['C'],1e-6)));tsl=W_TSLOPE*np.array([slm-sld if sld==sld else 0.])
        return np.concatenate([rpx,shp,pl,tl,tsl])
    lo=[-3.3,-3.3,np.log10(1e-4),np.log10(1e-6)];hi=[np.log10(1-1e-4),np.log10(.999),np.log10(1.0),np.log10(.1)]
    best=None
    for s0 in [[-1.5,-1.,np.log10(.005),krs],[-.5,-.3,np.log10(.02),krs],[-2.3,-1.5,np.log10(.001),krs-0.5]]:
        rr=least_squares(resid,s0,bounds=(lo,hi),max_nfev=80)
        if best is None or rr.cost<best.cost: best=rr
    a_s=1-10**best.x[0];am=10**best.x[1];fx=10**best.x[2];krf=10**best.x[3];rsf=rs0;kr_pv=krf*(L/eng.v);kr_s=krf
    r1=eng.run(kk,a_s,am,fx,krf)
    # model curves for plotting
    idx=np.linspace(0,len(r1['tp'])-1,90).astype(int)
    btf=np.column_stack([r1['tp'][idx],np.log10(np.maximum(r1['C'][idx],1e-6))])
    rpf=np.column_stack([r1['x'],logK+np.log10(np.maximum(r1['rp']/V_MS,1e-300))])
    lS_at=logK+np.log10(np.maximum(np.interp(rx,r1['x'],r1['rp'])/V_MS,1e-300))
    rpRMS=float(np.sqrt(np.mean((lS_at-rlog)**2)))
    stm=pv>1.2   # shelf+tail region; excludes the steep breakthrough front (log-axis front timing inflates RMS)
    btRMS=float(np.sqrt(np.mean((np.log10(np.maximum(np.interp(pv[stm],r1['tp'],r1['C']),1e-6))-lc[stm])**2))) if stm.sum() else float('nan')
    lSf=logK+np.log10(np.maximum(r1['rp']/V_MS,1e-300)); _im=int(np.argmax(lSf))
    pk_m=float(r1['x'][_im]*100) if _im>0 else 0.0
    _id=int(np.argmax(rlog)); pk_d=float(rx[_id]*100) if _id>0 else 0.0
    v=resid(best.x); n1=len(rx); ntl=int(tm.sum()) if tm.sum() else 1
    hh=lambda a:0.5*float(np.sum(np.asarray(a)**2))
    blk=dict(rp_level=hh(v[:n1]),rp_shape=hh(v[n1:n1+2]),plateau=hh(v[n1+2:n1+3]),
             tail_level=hh(v[n1+3:n1+3+ntl]),tail_slope=hh(v[n1+3+ntl:]))
    return dict(rs=rsf,rs_seed=rs0,seed_src=src,a_s=a_s,a_m=am,fx=fx,kr_pv=kr_pv,kr_s=kr_s,cost=float(best.cost),
                rpRMS=rpRMS,btRMS=btRMS,bt_exp=bt,rp_exp=rp,bt_fit=btf,rp_fit=rpf,
                branch=branch,nin=(nin if nin else 0),pk_m=pk_m,pk_d=pk_d,**blk)
# ---- fit all unfavorable (drop IS<=1 mM: retention gradient below model resolution, W.P.J.) ----
unf=[k for k,mm in meta.items() if mm['chem']=='unfavorable' and not (mm['IS'] is not None and mm['IS']<=1.0)]
# PASS 1 -- classify every column from its first two RP points (no fitting)
colbr={}
for k in unf:
    rp=cols[k]['RP']
    if len(rp)>=3 and len(cols[k]['BTEC'])>=3: colbr[k]=rp_branch(rp[:,1])
# PASS 2 -- one branch per CONDITION by majority, then fit every column with that condition's window
cbr={}
_grp=collections.defaultdict(list)
for k in colbr: mm=meta[k]; _grp[(k[0],mm['medium'],mm['size'],mm['vel'],mm['IS'])].append(k)
for key,ks in _grp.items():
    b=condition_branch([colbr[k] for k in ks]); cbr[key]=b
    if len(set(colbr[k] for k in ks))>1:
        print(f"  NOTE replicates disagree: {key} -> columns "
              f"{{{', '.join(f'{k[2]}:{colbr[k]}' for k in sorted(ks,key=lambda z:z[2]))}}} ; majority = {b}")
res={}
for k in unf:
    mm=meta[k]; key=(k[0],mm['medium'],mm['size'],mm['vel'],mm['IS'])
    if key not in cbr: continue
    nin=NIN_PEAKED if cbr[key]=='peaked' else None
    v=fit_col(k,nin_override=nin)
    if v: v['cond_branch']=cbr[key]; res[k]=v
print(f"fitted {len(res)}/{len(unf)} unfavorable columns")
# ================= workbook =================
Bf=Font(bold=True); IT=Font(italic=True)
wb=openpyxl.Workbook(); wb.remove(wb.active)
def hdr(ws,r,vals):
    for j,v in enumerate(vals,1): c=ws.cell(r,j,v); c.font=Bf
# ---- Master table (per author-condition) ----
ms=wb.create_sheet("Master table")
ms.cell(1,1,"Unfavorable master table -- Serial-3 interception-history fits (one row per study x condition; replicates averaged over columns; r_s FIXED at favorable anchor)").font=Bf
MHDR=["study","column(s)","medium","size (um)","IS (mM)","v (m/day)","r_s (/m, fixed)","alpha_s","alpha_m","f_x","k_r (/s)","RP branch (measured)","inlet pts scored"]
hdr(ms,3,MHDR)
byrow=collections.defaultdict(list)
for k in res: mm=meta[k]; byrow[(k[0],mm['medium'],mm['size'],mm['vel'],mm['IS'])].append(k)
rr=4; percol=[]
for key in sorted(byrow, key=lambda t:(t[1],t[2],t[3],(t[4] or 0),t[0])):
    src,med,sz,vel,IS=key; ks=sorted(byrow[key], key=lambda k:k[2])
    def mean(f): return float(np.mean([res[k][f] for k in ks]))
    letters="/".join(k[2] for k in ks)
    rsv=res[ks[0]]['rs_seed']; ssrc=res[ks[0]]['seed_src']; rstxt=f"{rsv:.1f}{'*' if ssrc=='TE' else ''}"
    vals=[src,letters,med,sz,(IS if IS is not None else ""),vel,rstxt,round(mean('a_s'),4),
          round(mean('a_m'),3),round(mean('fx'),4),float(f"{mean('kr_s'):.3e}"),
          res[ks[0]]['cond_branch'],(res[ks[0]]['nin'] or 'half')]
    for j,v in enumerate(vals,1): ms.cell(rr,j,v)
    rr+=1
    for k in ks: percol.append((key,k))
ms.cell(rr+1,1,"r_s held FIXED at the favorable interception rate (chemistry-independent, grounded in CFT); alpha_s/alpha_m/f_x/k_r carry the chemistry. r_s source: fitted favorable r_s /m, RP-slope favorable seed for c<QL 0.1-0.2 um glass, and * = Tufenkji-Elimelech (2004) x 0.77 where no favorable data exist (glass 0.5/8, quartz 0.5/4).").font=IT
for col,w in zip("ABCDEFGHIJKLM",[7,11,7,8,8,9,13,9,9,9,10,20,16]): ms.column_dimensions[col].width=w
# ---- Per-column fits ----
pc=wb.create_sheet("Per-column fits")
hdr(pc,1,["study","column","medium","size (um)","IS (mM)","v (m/day)","r_s (/m, fixed)","r_s source","alpha_s","alpha_m","f_x","k_r (/s)","RP RMS (log)","shelf+tail RMS (log)","branch (this column)","branch (condition)","inlet pts scored"])
prr=2
for key,k in sorted(percol, key=lambda t:(t[0][1],t[0][2],t[0][3],(t[0][4] or 0),t[0][0],t[1][2])):
    src,med,sz,vel,IS=key; d=res[k]
    row=[src,k[2],med,sz,(IS if IS is not None else ""),vel,round(d['rs'],1),d['seed_src'],round(d['a_s'],4),round(d['a_m'],4),
         round(d['fx'],4),float(f"{d['kr_s']:.3e}"),round(d['rpRMS'],3),round(d['btRMS'],3),
         d['branch'],d['cond_branch'],(d['nin'] or 'half')]
    for j,v in enumerate(row,1): pc.cell(prr,j,v)
    prr+=1
for col,w in zip("ABCDEFGHIJKLMNOPQ",[7,7,7,8,8,9,13,10,9,9,9,10,12,13,19,19,16]): pc.column_dimensions[col].width=w
# ---- Fit quality: per-block cost decomposition, all conditions ----
fq=wb.create_sheet("Fit quality")
fq.cell(1,1,f"Fit quality -- full unfavorable set ({len(byrow)} conditions)").font=Bf
fq.cell(2,1,"Weighted cost = 1/2 sum of squared weighted residuals (lower = better). RP = level + shape; BTEC = plateau + tail-level + tail-slope. Replicate conditions = mean over columns.").font=IT
fq.cell(3,1,"** COSTS ARE COMPARABLE ONLY WITHIN ONE INLET-WINDOW CONVENTION. ** Peaked conditions are scored with a 4-point inlet window, non-peaking conditions with the half-split (Records/plateau_rs_decision.md s6). The 4-point window targets the measured inlet RISE over x = 1-7 cm, which is far steeper than the 1-11 cm half-split average, so a peaked condition is scored against a HARDER target and shows a LARGER cost even when its fit is closer to the data. Do NOT compare a peaked cost against a non-peaking one, or against a pre-2026-08-24 value. The convention-independent metrics are RP RMS and RP peak depth, given at right.").font=IT
hdr(fq,4,["condition","RP level","RP shape","RP total","Plateau","Tail level","Tail slope","BTEC total","TOTAL","window","RP RMS","peak cm (model)","peak cm (measured)"])
fqr=5; tots=[]
for key in sorted(byrow, key=lambda t:(t[1],t[3],t[2],(t[4] or 0),t[0])):
    st,med,sz,vel,IS=key; ks=byrow[key]
    mn=lambda f: float(np.mean([res[k][f] for k in ks]))
    lab=f"{st} {med} {sz:.1f}um {vel:.0f}md {(IS or 0):.0f}mM"
    rl,rsh,plt_=mn('rp_level'),mn('rp_shape'),mn('plateau'); tlv,tsv=mn('tail_level'),mn('tail_slope')
    tot=rl+rsh+plt_+tlv+tsv; tots.append((tot,lab))
    win=(res[ks[0]]['nin'] or 'half')
    for j,v in enumerate([lab,round(rl,2),round(rsh,2),round(rl+rsh,2),round(plt_,2),
                          round(tlv,2),round(tsv,2),round(plt_+tlv+tsv,2),round(tot,2),win,
                          round(mn('rpRMS'),3),round(mn('pk_m'),2),round(mn('pk_d'),2)],1): fq.cell(fqr,j,v)
    fqr+=1
mx=max(tots)
fq.cell(fqr+1,1,f"median {np.median([t for t,_ in tots]):.1f} | max {mx[0]:.1f} ({mx[1]}) | n={len(tots)}. Residual concentrates in RP-shape (steep glass hyper-exponential inlet / sharp quartz peak -- single-population limitation); plateau ~0 except declining 50 mM conditions; BTEC well captured throughout.").font=IT
for col,w in zip("ABCDEFGHIJKLM",[34,10,10,10,10,11,11,11,10,9,9,17,19]): fq.column_dimensions[col].width=w
# ---- one plottable sheet per condition (across authors) ----
bycond=collections.defaultdict(list)
for k in res: mm=meta[k]; bycond[(mm['medium'],sizeclass(mm['size']),mm['vel'],mm['IS'])].append(k)
def block(ws,r0,title,expkey,fitkey,ks,xlab,ylab):
    ws.cell(r0,1,title).font=Bf; ws.cell(r0+1,1,f"x = {xlab}").font=IT; ws.cell(r0+1,4,f"y = {ylab}").font=IT
    hr=r0+2; cc=1
    for k in ks:
        lab=f"{k[0]}.{k[2]}"
        ws.cell(hr,cc,f"{lab} x").font=Bf; ws.cell(hr,cc+1,f"{lab} exp").font=Bf
        arr=res[k][expkey]
        for i,(x,y) in enumerate(arr): ws.cell(hr+1+i,cc,round(float(x),4)); ws.cell(hr+1+i,cc+1,round(float(y),4))
        cc+=2
    ws.cell(hr,cc,"|").font=Bf; cc+=1
    for k in ks:
        lab=f"{k[0]}.{k[2]}"
        ws.cell(hr,cc,f"{lab} xfit").font=Bf; ws.cell(hr,cc+1,f"{lab} fit").font=Bf
        arr=res[k][fitkey]
        for i,(x,y) in enumerate(arr): ws.cell(hr+1+i,cc,round(float(x),4)); ws.cell(hr+1+i,cc+1,round(float(y),4))
        cc+=2
    # per-study within-replicate mean +- stdev, placed BESIDE this block (studies with >1 replicate; RP C0-normalized)
    normalize=(expkey=='rp_exp'); bystudy=collections.defaultdict(list)
    for k in ks: bystudy[k[0]].append(k)
    for study,sks in sorted(bystudy.items()):
        if len(sks)<2: continue
        ws.cell(hr,cc,"|").font=Bf; cc+=1
        xref=res[sks[0]][expkey][:,0]
        C0avg=float(np.exp(np.mean(np.log([meta[k]['C0'] for k in sks])))) if normalize else None
        lins=[]
        for k in sks:
            a=res[k][expkey]; yy=np.interp(xref,a[:,0],a[:,1]); lin=10**yy
            if normalize: lin=lin/meta[k]['C0']*C0avg
            lins.append(lin)
        A=np.vstack(lins); M=A.mean(0); SD=np.log10(np.maximum(A,1e-300)).std(0,ddof=1)
        h2=f"{study} mean log10"+(f" (C0avg {C0avg:.3g})" if normalize else "")
        ws.cell(hr,cc,f"{study} x").font=Bf; ws.cell(hr,cc+1,h2).font=Bf; ws.cell(hr,cc+2,f"{study} sd dex").font=Bf
        for i in range(len(xref)):
            ws.cell(hr+1+i,cc,round(float(xref[i]),4)); ws.cell(hr+1+i,cc+1,round(float(np.log10(M[i])),4)); ws.cell(hr+1+i,cc+2,round(float(SD[i]),4))
        cc+=3
    return hr+1+max(max(len(res[k][expkey]),len(res[k][fitkey])) for k in ks)
used=set()
for cond in sorted(bycond, key=lambda t:(t[0],t[1],t[2],(t[3] or 0))):
    med,szc,vel,IS=cond; ks=sorted(bycond[cond], key=lambda k:(k[0],k[2]))
    nm=f"{med[:2]}_{szc}um_{vel:.0f}md_{(IS if IS else 0):.0f}mM"
    nm=nm[:31]
    while nm in used: nm=nm[:29]+"_2"
    used.add(nm); ws=wb.create_sheet(nm)
    ws.cell(1,1,f"{med}  {szc} um  {vel:.0f} m/day  {(IS if IS else 0):.0f} mM  (unfavorable) -- cols: "+", ".join(f'{k[0]}.{k[2]}' for k in ks)).font=Bf
    end=block(ws,3,"BTEC","bt_exp","bt_fit",ks,"pore volumes","log10 C/C0")
    block(ws,end+3,"RP (retention profile)","rp_exp","rp_fit",ks,"distance (m)","log10 spheres")
    for col in [openpyxl.utils.get_column_letter(i) for i in range(1,30)]: ws.column_dimensions[col].width=11
wb.save("../Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx")
print(f"wrote UnfavorableMaster.xlsx : Master table ({len(byrow)} rows), Per-column fits ({len(percol)}), {len(bycond)} condition sheets")
