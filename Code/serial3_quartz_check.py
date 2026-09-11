"""serial3_quartz_check.py -- Serial-3 rerun of the medium-axis quartz cross-check of the (near-universal)
k_r. Mirrors parallel quartz_check_fits.py but with the SERIAL-3 generator. v_ns=5% pinned, r_s per
size/medium, k_r FREE per column. Question: does Serial-3's near-universal glass k_r shift down for quartz?"""
import numpy as np, openpyxl, re, warnings, functools, json
from openpyxl.utils import get_column_letter, column_index_from_string as cix
warnings.filterwarnings("ignore"); print=functools.partial(print,flush=True)
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
DATA="../Data/DataFromLi&Tong.xlsx"; DAY=86400.; L=0.2; theta=0.36
REV,T0,Vref=22.801836559387397,3.58,0.1667; INJPV=T0/(L/Vref)
W_RP,W_SHAPE,W_PLAT,W_TAIL,W_TSLOPE=6.0,2.5,10.0,3.0,16.0; LODF=np.log10(4e-5)
isp=lambda v:any(abs(v-p)<0.05 for p in {-6.,-9.94,-15.33})
book=openpyxl.load_workbook(DATA,data_only=True)
def c0h(ws,c,r):
    m=re.search(r'([\d.]+)\s*E\s*([+-]?\d+)',str(ws.cell(r,c).value));return float(m.group(1))*10**int(m.group(2)) if m else None
def c0scan(H):
    m=re.search(r'([\d.]+)\s*E\s*([+-]?\d+)'," || ".join(str(x) for x in H));return float(m.group(1))*10**int(m.group(2)) if m else None
# RP_ROWS: per-sheet row range -- see favorable_both_models.py for why a single hardcoded range for both
# Li sheets is wrong (fixed 2026-09-01, W.P.J.). This file only ever calls load_li on Quartz Sand Li today
# (latent, not live, for Glass Beads Li) -- the assert stays so it can't come back silently.
RP_ROWS = {"Microspheres Glass Beads Li": (47, 57), "Microspheres Quartz Sand Li": (41, 51)}


def load_li(sh,dc):
    ws=book[sh]
    bt=[(ws.cell(r,dc).value,ws.cell(r,dc+1).value) for r in range(10,45)];bt=[(a,b) for a,b in bt if isinstance(a,(int,float)) and isinstance(b,(int,float)) and b<0.5 and not isp(b)]
    r0,r1=RP_ROWS[sh]
    rp=[(ws.cell(r,dc).value,ws.cell(r,dc+1).value) for r in range(r0,r1)];rp=[(a,b) for a,b in rp if isinstance(a,(int,float)) and isinstance(b,(int,float)) and b>0.5 and not isp(b)]
    assert len(rp)==10, f"RP truncated: {sh} col {dc} got {len(rp)} points, expected 10 -- row range wrong for this sheet"
    return np.array(bt),np.array(rp),c0h(ws,dc,8)
def load_tong(dcol):
    ws=book["Microspheres Quartz Sand Tong"];br=None
    for r in range(6,13):
        x=ws.cell(r,dcol).value
        if isinstance(x,str) and x.strip().upper().startswith("PV"):br=r;break
    H=[str(ws.cell(r,dcol).value) for r in range(1,br) if ws.cell(r,dcol).value not in (None,"")]
    C0=c0h(ws,dcol,br) or c0scan(H);pv=[];lc=[]
    for r in range(br+1,67):
        a=ws.cell(r,dcol).value;b=ws.cell(r,dcol+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isp(b) and b<0.5:pv.append(float(a));lc.append(float(b))
    xx=[];rp=[]
    for r in range(68,78):
        a=ws.cell(r,dcol).value;b=ws.cell(r,dcol+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isp(b):xx.append(float(a));rp.append(float(b))
    # ALL RP data has exactly 10 depth points, always (W.P.J., 2026-09-01) -- assert, don't silently truncate.
    if rp: assert len(rp)==10, f"RP truncated: col {dcol} got {len(rp)} points, expected 10"
    return np.array(list(zip(pv,lc))),np.array(list(zip(xx,rp))),C0
QC=[("Li-Qtz20 (1.1µm,4m/d,20mM)","li","Microspheres Quartz Sand Li",19,41.5,4.),
    ("Li-Qtz6 (1.1µm,4m/d,6mM)","li","Microspheres Quartz Sand Li",28,41.5,4.),
    ("Tong-B (0.5µm,4m/d,20mM)","tong",None,cix("B"),59.5,4.),
    ("Tong-H (0.5µm,4m/d,50mM)","tong",None,cix("H"),59.5,4.),
    ("Tong-R (0.5µm,8m/d,50mM)","tong",None,cix("R"),43.8,8.)]
class Eng:   # SERIAL-3
    def __init__(s,v):s.v=v/DAY;s.dx=L/90;s.dt=s.dx/s.v;s.pv=L/s.v
    def run(s,k1,a_s,a_m,fx,kr,D,vns=0.05):
        kf=k1;k2=fx*kf;kmw=a_m*kf;kmg=a_m*vns*kf;G=np.zeros((4,4))
        G[0,0]-=kf;G[3,0]+=a_s*kf;G[1,0]+=(1-a_s)*kf
        G[1,1]-=(kmw+k2);G[3,1]+=kmw;G[2,1]+=k2
        G[2,2]-=kmg;G[3,2]+=kmg;G[1,2]+=kr;G[2,2]-=kr
        s._k1=k1;s._kmw=kmw;s._kmg=kmg;E=expm(G*s.dt);r=D*s.dt/s.dx**2;mn=(1+2*r)*np.ones(90);of=-r*np.ones(89);mn[0]=1+r;mn[-1]=1+r;lu=splu(csc_matrix(diags([of,mn,of],[-1,0,1],format="csc")))
        nt=900;ti=INJPV*s.pv;c=vns;Y=np.zeros((4,90));C=np.zeros(nt);t=0.;Yst=None
        for i in range(nt):
            Y=E@Y;C[i]=(Y[0,-1]+Y[1,-1]+c*Y[2,-1]);Y[0,1:]=Y[0,:-1];Y[0,0]=0;Y[1,1:]=Y[1,:-1];Y[1,0]=0
            ym=Y[2].copy();Y[2,1:]=ym[1:]-c*(ym[1:]-ym[:-1]);Y[2,0]=ym[0]*(1-c);t+=s.dt
            if t<ti:Y[0,0]+=1.
            Y[0,:]=lu.solve(Y[0,:]);Y[1,:]=lu.solve(Y[1,:])
            if t<ti and t>0.9*ti:Yst=Y.copy()
        tp=(np.arange(nt)+1)*s.dt/s.pv;x=(np.arange(90)+0.5)*s.dx;rp=(a_s*s._k1*Yst[0,:]+s._kmw*Yst[1,:]+s._kmg*Yst[2,:]);m=(tp>0.5*INJPV)&(tp<INJPV)
        return dict(tp=tp,C=C,x=x,rp=rp,plat=float(C[m].mean()))
def sl(x,y):h=len(x)//2;return float(np.polyfit(x[:h+1],y[:h+1],1)[0]),float(np.polyfit(x[h:],y[h:],1)[0])
def ts(pv,lc,lo=6,hi=10):
    m=(np.asarray(pv)>=lo)&(np.asarray(pv)<=hi);return float("nan") if np.sum(m)<2 else float(np.polyfit(np.asarray(pv)[m],np.asarray(lc)[m],1)[0])
def fit(bt,rp,C0,rs,vel):
    col=Eng(vel);V_MS=vel/DAY;logK=np.log10(REV*T0*theta*C0*Vref);D=col.v*L/150;k1s=np.log10(rs*V_MS)
    x=rp[:,0];lrp=rp[:,1];pv=bt[:,0];lc=bt[:,1];si,so=sl(x,lrp);pm=(pv>1.3)&(pv<3.5);pp=lc[pm];plo,phi=float(pp.min()),float(pp.max());tm=pv>4.;sld=ts(pv,lc)
    def rr(p):
        k1,a_s,a_m,fx,kr=[10**q for q in p];r=col.run(k1,a_s,a_m,fx,kr,D)
        rl=np.interp(x,r['x'],np.maximum(r['rp'],1e-300))/V_MS;lS=logK+np.log10(np.maximum(rl,1e-300))
        rpx=W_RP*(lS-lrp);sm,sn=sl(x,lS);sh=W_SHAPE*np.array([sm-si,sn-so]);p_=np.log10(max(r['plat'],1e-12));ex=max(0.,p_-phi)+max(0.,plo-p_);pl=W_PLAT*np.array([ex])
        tl=W_TAIL*(np.log10(np.maximum(np.interp(pv[tm],r['tp'],r['C']),1e-12))-lc[tm]);slm=ts(r['tp'],np.log10(np.maximum(r['C'],1e-12)));tsl=W_TSLOPE*np.array([slm-sld])
        return np.concatenate([rpx,sh,pl,tl,tsl])
    lo=[k1s-0.3,-3.3,-3.3,np.log10(1e-4),np.log10(1e-6)];hi=[k1s+0.3,np.log10(.95),np.log10(.999),np.log10(1.0),np.log10(.1)]
    best=None
    for s0 in [[k1s,-1.3,-.5,np.log10(.005),np.log10(1e-4)],[k1s,-.1,-.3,np.log10(.02),np.log10(3e-4)]]:
        r=least_squares(rr,s0,bounds=(lo,hi),max_nfev=70)
        if best is None or r.cost<best.cost:best=r
    k1,a_s,a_m,fx,kr=[10**q for q in best.x];r=col.run(k1,a_s,a_m,fx,kr,D)
    rl=np.interp(x,r['x'],np.maximum(r['rp'],1e-300))/V_MS;lS=logK+np.log10(np.maximum(rl,1e-300))
    return r,lS,float(best.cost),dict(a_s=a_s,a_m=a_m,fx=fx,kr_pv=kr*(L/col.v),kr_s=kr)
fig,ax=plt.subplots(5,2,figsize=(13,15));rows=[]
print(f"{'column':28}{'r_s':>6}{'k_r/PV':>8}{'k_r/s':>11}{'cost':>8}")
for i,(label,mode,sh,dc,rs,vel) in enumerate(QC):
    bt,rp,C0=(load_li(sh,dc) if mode=="li" else load_tong(dc))
    r,lS,cost,par=fit(bt,rp,C0,rs,vel);aL,aR=ax[i]
    rows.append(dict(label=label,rs=rs,kr_pv=par['kr_pv'],kr_s=par['kr_s'],cost=cost,a_s=par['a_s'],a_m=par['a_m'],fx=par['fx']))
    print(f"{label:28}{rs:>6.1f}{par['kr_pv']:>8.3f}{par['kr_s']:>11.2e}{cost:>8.1f}")
    aL.plot(bt[:,0],bt[:,1],'o',ms=4,color="#7b2d8b",alpha=.8);aL.plot(r['tp'],np.log10(np.maximum(r['C'],1e-10)),'-',lw=1.6,color="#7b2d8b")
    aL.axhline(LODF,ls=':',lw=.8,color="grey");aL.set_xlim(0,11);aL.set_ylim(-6,.6);aL.set_xlabel("PV");aL.set_ylabel("log C/C0");aL.grid(alpha=.3)
    aL.set_title(f"BTEC — {label}  k_r={par['kr_pv']:.2f}/PV c={cost:.1f}",fontsize=9)
    aR.plot(rp[:,0],rp[:,1],'o',ms=4,color="#7b2d8b",alpha=.8);aR.plot(rp[:,0],lS,'-',lw=1.6,color="#7b2d8b")
    aR.set_xlim(0,.2);aR.set_ylim(5.,8.7);aR.set_xlabel("x (m)");aR.set_ylabel("log spheres");aR.grid(alpha=.3);aR.set_title(f"RP — {label.split()[0]}",fontsize=9)
fig.suptitle("SERIAL-3 QUARTZ cross-check — BTEC+RP fits (v_ns=5%, r_s per size/medium); Li 1.1µm reliable, Tong-H tail/RP good (plateau off), Tong-R near-LOD",fontsize=10)
fig.tight_layout(rect=[0,0,1,0.99]);fig.savefig("artifacts/serial3_quartz_check.png",dpi=140)
json.dump(rows,open("artifacts/serial3_quartz_results.json","w"),indent=1)
rel=[x for x in rows if x['label'].startswith('Li')]
print("\nReliable (Li 1.1µm) quartz k_r/s:",["%.2e"%x['kr_s'] for x in rel]," (glass shared was 5.58e-5/s)")
print("wrote serial3_quartz_check.png, serial3_quartz_results.json")
