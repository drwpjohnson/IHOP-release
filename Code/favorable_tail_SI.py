"""favorable_tail_SI.py -- SI companion to favorable_tail_clean.py. Two-panel (BTEC left, RP right) fits for the
FIVE favorable columns that are fit in Favorable_data_and_sim.xlsx but omitted from the
5-column main-text figure: Glass Tong CS (2.0 um, 8 m/d), Quartz Tong O (0.5 um, 8 m/d), and Glass Li B/E/H
(0.98 um, 2/4/8 m/d). Same setup as favorable_tail_clean.py: k_r PINNED per-medium; r_s (±0.6 dex), alpha_s
(via MI 1-alpha_s), alpha_m, f_x FREE; joint BTEC + RP (level+slope) objective; RP amplitude FIXED by C0.
C0 row map: GB Li row 14, others row 8 (data_inventory §4); CS/O have no C0 in the sheet -> RP-implied
(1.157e6 / 5.762e6). Quartz Tong O has no BTEC tail region (>4.2 PV) so its tail RMS is n/a.

RESULT (re-run 2026-09-01 after the RP row-range fix -- see Records/CLAUDE.md; GB Li B/E/H's RP was
truncated to 4 of 10 points before this, and their a_m was pinned at the 0.0005 fit floor as a direct
consequence -- 4 points cannot identify a crawl-attachment fraction): CS rs=15.2 MI=1.55% a_m=0.90*
fx=0.10* rpRMS=0.120 ; O rs=56.9 MI=3.49% a_m=0.44 fx=0.10* rpRMS=0.145 ; GB Li B rs=49.9 MI=1.02%
a_m=0.075 fx=0.0116 rpRMS=0.164 ; GB Li E rs=30.2 MI=2.55% a_m=0.246 fx=0.0069 rpRMS=0.037 ; GB Li H
rs=29.6 MI=19.47% a_m=0.351 fx=0.0026 rpRMS=0.045. (* = at a fit bound; CS/O's f_x and CS's a_m remain
bound-pinned, unrelated to the RP bug. Favorable data under-constrain the 4-parameter fit in general --
read the tail-reality Fit Ratio, not the individual params.) Writes favorable_tail_SI.png.
Requires DataFromLi&Tong.xlsx under ../Data/."""
import numpy as np, openpyxl, re, os
from openpyxl.utils import column_index_from_string as cix
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
_HERE=os.path.dirname(os.path.abspath(__file__)); DATA=os.path.join(_HERE,"..","Data","DataFromLi&Tong.xlsx")
DAY=86400.; L=0.2; REV,T0,Vref=22.801836559387397,3.58,0.1667; INJPV=T0/(L/Vref)
KR_MED={"glass":5.58e-5,"quartz":1.87e-5}   # PINNED; superseded by 3.75e-5 / 1.82e-5 -- see favorable_both_models.py for why it is not updated
THETA={"glass":0.375,"quartz":0.36}
W_RP,W_SH,W_BT=6.0,2.5,3.0
C0_OVERRIDE={("Glass Tong","CS"):1.157e6,("Quartz Tong","O"):5.762e6}
PLACE={-6.,-9.94,-15.33}; isp=lambda v:any(abs(v-p)<0.05 for p in PLACE)
wb=openpyxl.load_workbook(DATA,data_only=True)
def c0cell(ws,c,r):
    m=re.search(r'([\d.]+)\s*E\s*([+-]?\d+)',str(ws.cell(r,c).value));return float(m.group(1))*10**int(m.group(2)) if m else None
# RP_ROWS: per-sheet row range -- Glass Beads Li's RP header sits 6 rows lower than Quartz Sand Li's
# (its BTEC block runs longer). Fixed 2026-09-01 (W.P.J.): a single hardcoded range for both sheets
# silently truncated Glass Beads Li's RP to 4 of 10 points. ALL RP data has exactly 10 depth points,
# always -- the assert is the contract, not optional defensive code.
RP_ROWS = {"Microspheres Glass Beads Li": (47, 57), "Microspheres Quartz Sand Li": (41, 51)}


def load_li(sh,dc):
    ws=wb[sh]
    bt=[(ws.cell(r,dc).value,ws.cell(r,dc+1).value) for r in range(10,45)];bt=[(a,b) for a,b in bt if isinstance(a,(int,float)) and isinstance(b,(int,float)) and b<0.5 and not isp(b)]
    r0,r1=RP_ROWS[sh]
    rp=[(ws.cell(r,dc).value,ws.cell(r,dc+1).value) for r in range(r0,r1)];rp=[(a,b) for a,b in rp if isinstance(a,(int,float)) and isinstance(b,(int,float)) and b>0.5 and not isp(b)]
    assert len(rp)==10, f"RP truncated: {sh} col {dc} got {len(rp)} points, expected 10 -- row range wrong for this sheet"
    return np.array(bt),np.array(rp),c0cell(ws,dc,14 if "Glass Beads Li" in sh else 8)
def load_tong(sh,c):
    ws=wb[sh];br=None
    for r in range(6,13):
        x=ws.cell(r,c).value
        if isinstance(x,str) and x.strip().upper().startswith("PV"):br=r;break
    pv=[];lc=[]
    for r in range(br+1,67):
        a=ws.cell(r,c).value;b=ws.cell(r,c+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isp(b) and b<0.5: pv.append(float(a));lc.append(float(b))
    xx=[];rp=[]
    for r in range(68,78):
        a=ws.cell(r,c).value;b=ws.cell(r,c+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isp(b): xx.append(float(a));rp.append(float(b))
    # ALL RP data has exactly 10 depth points, always (W.P.J., 2026-09-01) -- assert, don't silently truncate.
    if rp: assert len(rp)==10, f"RP truncated: {sh} col {c} got {len(rp)} points, expected 10"
    return np.array(list(zip(pv,lc))),np.array(list(zip(xx,rp))),c0cell(ws,c,8)
def rs_plateau(bt):
    pv=bt[:,0];lc=bt[:,1];m=(pv>1.2)&(pv<3.4)
    if m.sum()<2:m=(pv>0.9)&(pv<3.6)
    return -np.log(10)*float(np.median(lc[m]))/L
def sl(x,y):h=len(x)//2;return float(np.polyfit(x[:h+1],y[:h+1],1)[0]),float(np.polyfit(x[h:],y[h:],1)[0])
class Eng:
    def __init__(s,vmday):s.v=vmday/DAY;s.dx=L/90;s.dt=s.dx/s.v;s.pv=L/s.v
    def run(s,k1,a_s,a_m,fx,kr,vns=0.05,tot=10):
        kf=k1;k2=fx*kf;kmw=a_m*kf;kmg=a_m*vns*kf;G=np.zeros((4,4))
        G[0,0]-=kf;G[3,0]+=a_s*kf;G[1,0]+=(1-a_s)*kf;G[1,1]-=(kmw+k2);G[3,1]+=kmw;G[2,1]+=k2;G[2,2]-=kmg;G[3,2]+=kmg;G[1,2]+=kr;G[2,2]-=kr
        E=expm(G*s.dt);D=s.v*L/150;r=D*s.dt/s.dx**2;mn=(1+2*r)*np.ones(90);of=-r*np.ones(89);mn[0]=1+r;mn[-1]=1+r
        lu=splu(csc_matrix(diags([of,mn,of],[-1,0,1],format="csc")))
        nt=int(round(tot*90));ti=INJPV*s.pv;c=vns;Y=np.zeros((4,90));C=np.zeros(nt);t=0.;Yst=None
        for i in range(nt):
            Y=E@Y;C[i]=Y[0,-1]+Y[1,-1]+c*Y[2,-1];Y[0,1:]=Y[0,:-1];Y[0,0]=0;Y[1,1:]=Y[1,:-1];Y[1,0]=0
            ym=Y[2].copy();Y[2,1:]=ym[1:]-c*(ym[1:]-ym[:-1]);Y[2,0]=ym[0]*(1-c);t+=s.dt
            if t<ti:Y[0,0]+=1.
            Y[0,:]=lu.solve(Y[0,:]);Y[1,:]=lu.solve(Y[1,:])
            if t<ti and t>0.9*ti:Yst=Y.copy()
        tp=(np.arange(nt)+1)*s.dt/s.pv;x=(np.arange(90)+0.5)*s.dx
        rp=(a_s*kf*Yst[0,:]+kmw*Yst[1,:]+kmg*Yst[2,:])
        return dict(tp=tp,C=np.maximum(C,1e-300),x=x,rp=np.maximum(rp,1e-300))
COLS=[("Glass Tong","CS","tong","glass",2.0,8),("Quartz Tong","O","tong","quartz",0.5,8),
      ("Glass Li","B","li","glass",0.98,2),("Glass Li","E","li","glass",0.98,4),("Glass Li","H","li","glass",0.98,8)]
SHEET={"Glass Tong":"Microspheres Glass Beads Tong","Quartz Tong":"Microspheres Quartz Sand Tong","Glass Li":"Microspheres Glass Beads Li"}
fig,ax=plt.subplots(len(COLS),2,figsize=(11,2.7*len(COLS)))
print(f"{'cond':16}{'r_s':>6}{'MI%':>7}{'a_m':>7}{'f_x':>8}{'tRMS':>7}{'rpRMS':>7}")
for i,(tab,letter,layout,med,size,vel) in enumerate(COLS):
    if layout=="li": bt,rp,C0=load_li(SHEET[tab],cix(letter))
    else:            bt,rp,C0=load_tong(SHEET[tab],cix(letter))
    if C0 is None: C0=C0_OVERRIDE.get((tab,letter))
    rs_seed=rs_plateau(bt); V_MS=vel/DAY; eng=Eng(vel); kr=KR_MED[med]; theta=THETA[med]
    logK=np.log10(REV*T0*theta*Vref*C0)
    pv=bt[:,0];lc=bt[:,1];rx=rp[:,0];rlog=rp[:,1];si,so=sl(rx,rlog);k1s=np.log10(rs_seed*V_MS);tm=pv>4.2
    r0=eng.run(rs_seed*V_MS,1.0,0.,0.,0.)
    def rpm(r): return logK+np.log10(np.maximum(np.interp(rx,r['x'],r['rp'])/V_MS,1e-300))
    def resid(p):
        kk=10**p[0];a_s=1-10**p[1];am=10**p[2];fx=10**p[3];r=eng.run(kk,a_s,am,fx,kr)
        mC=np.log10(np.maximum(np.interp(pv,r['tp'],r['C']),1e-12));lS=rpm(r);sm,sn=sl(rx,lS)
        return np.concatenate([W_RP*(lS-rlog),W_SH*np.array([sm-si,sn-so]),W_BT*(mC-lc)])
    best=None
    for s0 in [[k1s,-1.5,-1.5,np.log10(.005)],[k1s,-2.0,-2.0,np.log10(.02)]]:
        rr=least_squares(resid,s0,bounds=([k1s-0.6,-3,-3.3,np.log10(1e-4)],[k1s+0.6,np.log10(0.5),np.log10(0.9),np.log10(0.1)]),max_nfev=60)
        if best is None or rr.cost<best.cost:best=rr
    k1f=10**best.x[0];a_s=1-10**best.x[1];a_m=10**best.x[2];fx=10**best.x[3];rs=k1f/V_MS
    r1=eng.run(k1f,a_s,a_m,fx,kr)
    tr=float(np.sqrt(np.mean((np.log10(np.maximum(np.interp(pv[tm],r1['tp'],r1['C']),1e-12))-lc[tm])**2))) if tm.sum()>=2 else float('nan')
    rpRMS=float(np.sqrt(np.mean((rpm(r1)-rlog)**2)))
    print(f"{tab+' '+letter:16}{rs:>6.1f}{100*(1-a_s):>7.2f}{a_m:>7.3f}{fx:>8.4f}{tr:>7.3f}{rpRMS:>7.3f}")
    col=("#e07b00" if med=="glass" else "#7b2d8b"); aL,aR=ax[i]
    aL.plot(pv,lc,"ko",ms=5,zorder=5,label="data")
    aL.plot(r0['tp'],np.log10(r0['C']),"--",color="grey",lw=1.6,label="k_f only")
    aL.plot(r1['tp'],np.log10(r1['C']),"-",color=col,lw=2.2,label=f"+{100*(1-a_s):.1f}% MI (k_r {kr*(L/eng.v):.2f}/PV)")
    aL.set_xlim(0,12);aL.set_ylim(-6.2,0.4);aL.set_xlabel("pore volumes");aL.set_ylabel("log10 C/C0")
    aL.set_title(f"{tab} {letter}: {med} {size}um {vel} m/d - BTEC (r_s={rs:.0f}/m)",fontsize=9);aL.legend(fontsize=7,loc="upper right")
    aR.plot(rx*100,rlog,"ko",ms=5,zorder=5,label="data")
    aR.plot(r1['x']*100,logK+np.log10(np.maximum(r1['rp']/V_MS,1e-300)),"-",color=col,lw=2.2,label="Serial-3 (fit)")
    aR.set_xlim(0,20);aR.set_xlabel("distance (cm)");aR.set_ylabel("log10 spheres")
    aR.set_title(f"{tab} {letter}: RP (rms {rpRMS:.3f})",fontsize=9);aR.legend(fontsize=7,loc="upper right")
fig.suptitle("SI - FAVORABLE (the 5 columns omitted from the main figure): Serial-3, k_r pinned; r_s,a_s,a_m,f_x free; BTEC+RP. CS/O use RP-implied C0.",fontsize=9.5)
fig.tight_layout(rect=[0,0,1,0.985]);fig.savefig("../Manuscript/FigsExcelsFavorable/favorable_tail_SI.png",dpi=125);print("wrote favorable_tail_SI.png")
