"""
fx_pin_test.py -- necessity test for the focus fraction f_x on the glass colloid-size series.

Question: does f_x carry real per-column information, or is its apparent variation just the f_x<->k_r
identifiability split?  For each glass size-series column, fit with f_x FREE (5 params: k_f, alpha_s,
alpha_m, f_x, k_r) vs f_x PINNED to a single constant F (4 params; k_f, alpha_s, alpha_m, k_r free to
compensate).  Compare total weighted cost.  Mirror of the k_r single-constant test in serial3_size_fit.py.

RESULT (2026): total cost free = 62.2, pinned (F = 0.0039, the free median) = 104.1  -> +67.5%
  (vs only +9% for a single shared k_r). So f_x is NOT reducible to a constant -- it carries real
  per-column information. ~90% of the excess comes from three SMALL-colloid columns (0.1-0.2 um):
    AX 0.1um +105%, BH 0.2um +143%, BN 0.2um +328%
  -- they require small f_x (4-6e-4) and cannot be forced up to 0.0039 even with k_r free. The >=0.5 um
  columns sit near the constant (0-28%). CONCLUSION: the SIZE dependence of f_x is required by the data
  (small colloids -> small f_x), mechanistically the shallow secondary minimum |U_sec| of small colloids
  (|U_sec| ~0.18 kT at 0.1um vs 3.55 kT at 2.0um; see dlvo_kramers_size.py). Medium and velocity remain
  within the identifiability scatter.  See Records/fx_trend_analysis.md.

Requires DataFromLi&Tong.xlsx under ../Data/.  Mirrors serial3_size_fit.py's engine/setup exactly.
Usage:  python3 fx_pin_test.py
"""
import functools, numpy as np, openpyxl, re, warnings, os
from openpyxl.utils import get_column_letter
print = functools.partial(print, flush=True); warnings.filterwarnings("ignore")
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
_HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(_HERE, "..", "Data", "DataFromLi&Tong.xlsx")
DAY = 86400.0; L = 0.2; theta = 0.375
REV, T0, Vref = 22.801836559387397, 3.58, 0.1667; AMP = REV * T0 * theta * Vref
W_RP, W_SHAPE, W_PLAT, W_TAIL, W_TSLOPE = 6.0, 2.5, 10.0, 3.0, 16.0; INJPV = T0 / (L / Vref)
RS = {(0.1, 8): 92.8, (0.2, 4): 87.9, (0.2, 8): 53.5, (0.5, 4): 35.6, (0.5, 8): 26.3, (1.1, 4): 22.9, (2.0, 8): 19.0}
PLACE = {-6.0, -9.94, -15.33}
def isplace(v): return any(abs(v - p) < 0.05 for p in PLACE)
wb = openpyxl.load_workbook(DATA, data_only=True); ws = wb["Microspheres Glass Beads Tong"]

def parse(c):
    br = None
    for r in range(6, 13):
        x = ws.cell(r, c).value
        if isinstance(x, str) and x.strip().upper().startswith("PV"): br = r; break
    if br is None: return None
    H = " || ".join(str(ws.cell(r, c).value) for r in range(1, br) if ws.cell(r, c).value is not None); hl = H.lower()
    d = {'cond': "unfav" if "unfav" in hl else ("fav" if "favor" in hl else "?"),
         'size': float(re.search(r'([\d.]+)\s*micron', hl).group(1)) if re.search(r'([\d.]+)\s*micron', hl) else None,
         'vel': float(re.search(r'([\d.]+)\s*m/day', hl).group(1)) if re.search(r'([\d.]+)\s*m/day', hl) else None,
         'IS': float(re.search(r'([\d.]+)\s*M\b', H).group(1)) if re.search(r'([\d.]+)\s*M\b', H) else None,
         'dg': "downgrad" in hl, 'twoPV': bool(re.search(r'2\s*pv', hl))}
    m = re.search(r'([\d.]+)\s*E\s*([+-]?\d+)', H); d['C0'] = float(m.group(1)) * 10 ** int(m.group(2)) if m else None
    pv = []; lc = []
    for r in range(br + 1, 67):
        a = ws.cell(r, c).value; b = ws.cell(r, c + 1).value
        if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isplace(b) and b < 0.5: pv.append(float(a)); lc.append(float(b))
    xx = []; rp = []
    for r in range(68, 78):
        a = ws.cell(r, c).value; b = ws.cell(r, c + 1).value
        if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isplace(b): xx.append(float(a)); rp.append(float(b))
    d['bt_pv'] = np.array(pv); d['bt_log'] = np.array(lc); d['x'] = np.array(xx); d['logrp'] = np.array(rp); d['col'] = get_column_letter(c); return d

def cols():
    o = []
    for c in range(1, ws.max_column + 1):
        for r in range(6, 13):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper().startswith("PV"): o.append(c); break
    return o
usable = [d for d in (parse(c) for c in cols()) if d and d['cond'] == "unfav" and not d['twoPV'] and not d['dg'] and d['IS'] == 0.02 and len(d['bt_pv']) > 0 and len(d['x']) > 0]

class Eng:
    def __init__(s, vmday): s.v = vmday / DAY; s.nx = 90; s.dx = L / 90; s.dt = s.dx / s.v; s.pv = L / s.v
    def run(s, k1, a_s, a_m, fx, kr, D, vns=0.05, inj=INJPV, tot=10):
        kf = k1; k2 = fx * kf; kmw = a_m * kf; kmg = a_m * vns * kf; G = np.zeros((4, 4))
        G[0, 0] -= kf; G[3, 0] += a_s * kf; G[1, 0] += (1 - a_s) * kf
        G[1, 1] -= (kmw + k2); G[3, 1] += kmw; G[2, 1] += k2; G[2, 2] -= kmg; G[3, 2] += kmg; G[1, 2] += kr; G[2, 2] -= kr
        s._k1 = k1; s._kmw = kmw; s._kmg = kmg; E = expm(G * s.dt)
        r = D * s.dt / s.dx ** 2; mn = (1 + 2 * r) * np.ones(90); of = -r * np.ones(89); mn[0] = 1 + r; mn[-1] = 1 + r
        lu = splu(csc_matrix(diags([of, mn, of], [-1, 0, 1], format="csc")))
        nt = int(round(tot * 90)); ti = inj * s.pv; c = vns; Y = np.zeros((4, 90)); C = np.zeros(nt); t = 0.; Yst = None
        for i in range(nt):
            Y = E @ Y; C[i] = (Y[0, -1] + Y[1, -1] + c * Y[2, -1])
            Y[0, 1:] = Y[0, :-1]; Y[0, 0] = 0; Y[1, 1:] = Y[1, :-1]; Y[1, 0] = 0
            ym = Y[2].copy(); Y[2, 1:] = ym[1:] - c * (ym[1:] - ym[:-1]); Y[2, 0] = ym[0] * (1 - c); t += s.dt
            if t < ti: Y[0, 0] += 1.0
            Y[0, :] = lu.solve(Y[0, :]); Y[1, :] = lu.solve(Y[1, :])
            if t < ti and t > 0.9 * ti: Yst = Y.copy()
        tp = (np.arange(nt) + 1) * s.dt / s.pv; x = (np.arange(90) + 0.5) * s.dx
        rp = (a_s * s._k1 * Yst[0, :] + s._kmw * Yst[1, :] + s._kmg * Yst[2, :])
        m = (tp > 0.5 * inj) & (tp < inj); return dict(tp=tp, C=C, x=x, rp=rp, plat=float(C[m].mean()))

def sl(x, y): h = len(x) // 2; return float(np.polyfit(x[:h + 1], y[:h + 1], 1)[0]), float(np.polyfit(x[h:], y[h:], 1)[0])
def ts(pv, lc, lo=6, hi=10):
    m = (np.asarray(pv) >= lo) & (np.asarray(pv) <= hi); return float("nan") if np.sum(m) < 2 else float(np.polyfit(np.asarray(pv)[m], np.asarray(lc)[m], 1)[0])

def setup(d):
    sz, v = d['size'], d['vel']; rs = RS[(sz, v)]; col = Eng(v); V_MS = v / DAY; logK = np.log10(AMP * d['C0']); D = col.v * L / 150
    si, so = sl(d['x'], d['logrp']); pm = (d['bt_pv'] > 1.3) & (d['bt_pv'] < 3.5); pp = d['bt_log'][pm]
    plo, phi = float(pp.min()), float(pp.max()); tm = d['bt_pv'] > 4.0; sld = ts(d['bt_pv'], d['bt_log']); k1s = np.log10(rs * V_MS)
    def resid(lk1, la_s, la_m, lfx, kr):
        k1 = 10 ** lk1; r = col.run(k1, 10 ** la_s, 10 ** la_m, 10 ** lfx, kr, D)
        rl = np.interp(d['x'], r['x'], np.maximum(r['rp'], 1e-300)) / V_MS; lS = logK + np.log10(np.maximum(rl, 1e-300))
        rp = W_RP * (lS - d['logrp']); sm, sn = sl(d['x'], lS); sh = W_SHAPE * np.array([sm - si, sn - so])
        p_ = np.log10(max(r['plat'], 1e-12)); ex = max(0., p_ - phi) + max(0., plo - p_); pl = W_PLAT * np.array([ex])
        tl = W_TAIL * (np.log10(np.maximum(np.interp(d['bt_pv'][tm], r['tp'], r['C']), 1e-12)) - d['bt_log'][tm])
        slm = ts(r['tp'], np.log10(np.maximum(r['C'], 1e-12))); tsl = W_TSLOPE * np.array([slm - sld])
        return np.concatenate([rp, sh, pl, tl, tsl])
    return dict(resid=resid, k1s=k1s, d=d, sz=sz, v=v)

def main():
    S = [setup(d) for d in sorted(usable, key=lambda z: (z['size'], z['vel']))]
    free = []
    for s in S:
        def rr(p): return s['resid'](p[0], p[1], p[2], p[3], 10 ** p[4])
        lo = [s['k1s'] - 0.3, -3.3, -3.3, np.log10(1e-4), np.log10(1e-6)]; hi = [s['k1s'] + 0.3, np.log10(.9), np.log10(.999), np.log10(1.0), np.log10(.1)]
        best = None
        for st in [[s['k1s'], -1.5, -1.0, np.log10(.005), np.log10(1e-4)], [s['k1s'], -2.2, -0.3, np.log10(.02), np.log10(3e-4)]]:
            r = least_squares(rr, st, bounds=(lo, hi), max_nfev=60)
            if best is None or r.cost < best.cost: best = r
        free.append((s, best.cost, 10 ** best.x[3]))
    F = float(np.median([f[2] for f in free]))
    print(f"pinned f_x constant F = {F:.4f} (median of free values)")
    print(f"{'col':>4}{'sz':>5}{'v':>3}{'fx_free':>9}{'cost_free':>10}{'cost_pin':>10}{'d%':>7}")
    tot_free = tot_pin = 0.
    lF = np.log10(F)
    for s, cfree, fxf in free:
        def rr(p): return s['resid'](p[0], p[1], p[2], lF, 10 ** p[3])
        lo = [s['k1s'] - 0.3, -3.3, -3.3, np.log10(1e-6)]; hi = [s['k1s'] + 0.3, np.log10(.9), np.log10(.999), np.log10(.1)]
        best = None
        for st in [[s['k1s'], -1.5, -1.0, np.log10(1e-4)], [s['k1s'], -2.2, -0.3, np.log10(3e-4)]]:
            r = least_squares(rr, st, bounds=(lo, hi), max_nfev=60)
            if best is None or r.cost < best.cost: best = r
        cpin = best.cost; tot_free += cfree; tot_pin += cpin
        print(f"{s['d']['col']:>4}{s['sz']:>5}{s['v']:>3.0f}{fxf:>9.4f}{cfree:>10.2f}{cpin:>10.2f}{100*(cpin-cfree)/max(cfree,1e-9):>7.1f}")
    print(f"\nTOTAL  free={tot_free:.2f}  pinned(f_x={F:.4f})={tot_pin:.2f}  penalty +{100*(tot_pin-tot_free)/tot_free:.1f}%")

if __name__ == "__main__":
    main()
