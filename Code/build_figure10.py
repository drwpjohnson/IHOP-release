#!/usr/bin/env python3
"""Build Figure 10 as a native openpyxl Excel chart: fitted k_r vs pore-water velocity, for every
(medium, size, IS) group in the unfavorable set that has a matched velocity pair (4 vs 8 m/day) --
three groups: glass 0.2 um 20 mM, glass 0.5 um 20 mM, quartz 0.5 um 50 mM. Each group's raw
per-column points are plotted, plus a line connecting that group's per-velocity geometric mean, and
the same two per-medium clean-geomean horizontal reference lines used in Figure 9.

Data source: Manuscript/FigsExcelsUnfav/k_r_trend.xlsx -- 'k_r fitted (new)' sheet, filtered to
IS_mM>1 and grouped by (medium, size, IS), keeping only groups with more than one distinct
velocity present (reproduces kr_trend.py's VG construction exactly); 'reference constants' sheet
for the glass/quartz clean geomeans (same source and same numbers as Figure 9's reference lines).

** FLAG, NOT SILENTLY FIXED: the embedded Figure 10 image (word/media/image16.png, rId26) carries
the SAME staleness already documented for Figure 9 -- it is the companion panel from the same
kr_trend.py run. Its dashed reference lines read glass ~4.5e-05/s, quartz ~2.1e-05/s (the
pre-2026-08-25, pre-accum-convention numbers); today's clean geomeans are glass 3.75e-05/s, quartz
1.82e-05/s. The three velocity-paired GROUPS and their qualitative story (glass 0.2 um and quartz
0.5 um both k_r(8)>k_r(4); glass 0.5 um shows the same direction more strongly) are unchanged --
only the two reference lines are out of date, exactly as for Figure 9. Not re-verified point-by-
point here since the underlying per-column values are the same k_r_trend.xlsx rows Figure 9 already
checked; only the additional geomean-per-velocity connecting lines are new to this figure.

Run from Code/.
"""
import collections
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text as ChartText
from openpyxl.chart.title import Title
from openpyxl.chart.legend import Legend
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (RichTextProperties, Paragraph, ParagraphProperties,
                                    CharacterProperties, Font as DrawFont, RegularTextRun)

SRC = "../Manuscript/FigsExcelsUnfav/k_r_trend.xlsx"   # run from Code/
OUTDIR = "../Manuscript/Figures"
APTOS = "Aptos"

PAL = ["E67E22", "16A085", "7B2D8B"]   # matches kr_trend.py's PAL[0:3] exactly, in sorted-group order
GREY_GLASS = "1F77B4"
GREY_QUARTZ = "7B2D8B"


def rich_run(text, sz=1400, bold=True, italic=False):
    cp = CharacterProperties(sz=sz, b=bold, i=italic)
    cp.latin = DrawFont(typeface=APTOS)
    return RegularTextRun(rPr=cp, t=text)


def axis_title(text, x, y):
    para = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True)),
                      r=[rich_run(text, sz=1400)])
    rt = RichText(bodyPr=RichTextProperties(), p=[para])
    t = Title(tx=ChartText(rich=rt))
    t.overlay = True
    t.layout = Layout(manualLayout=ManualLayout(xMode="edge", yMode="edge", x=x, y=y))
    return t


def read_source(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["k_r fitted (new)"]
    hdr = [ws.cell(4, c).value for c in range(1, 10)]
    rows = []
    r = 5
    while ws.cell(r, 1).value is not None:
        rows.append({h: ws.cell(r, c).value for c, h in enumerate(hdr, 1)})
        r += 1

    VG = collections.defaultdict(list)
    for v in rows:
        if v["IS_mM"] and v["IS_mM"] > 1:
            VG[(v["medium"], round(v["size_um"], 1), v["IS_mM"])].append(v)
    VG = {k: v for k, v in VG.items() if len({x["v_mday"] for x in v}) > 1}
    groups = [(k, VG[k]) for k in sorted(VG)]

    wref = wb["reference constants"]
    fg = wref.cell(6, 2).value
    fq = wref.cell(8, 2).value
    return groups, fg, fq


def gm(vals):
    return float(np.exp(np.mean(np.log(np.asarray(vals, float)))))


def write_data_sheet(ws, groups, fg, fq):
    c0 = 1
    meta = []
    for (med, size, IS), pts in groups:
        label = f"{med} {size:g} um, {IS:.0f} mM"
        ws.cell(1, c0, f"{label}: v (m/day)"); ws.cell(1, c0 + 1, "k_r (/s)"); ws.cell(1, c0 + 2, "identifier (study.col)")
        ws.column_dimensions[get_column_letter(c0 + 2)].width = 24
        for i, p in enumerate(pts, start=2):
            ws.cell(i, c0, p["v_mday"]); ws.cell(i, c0 + 1, p["k_r_/s"]); ws.cell(i, c0 + 2, p["col"])
        # per-velocity geomean line (2 points: v=4, v=8)
        byv = collections.defaultdict(list)
        for p in pts: byv[p["v_mday"]].append(p["k_r_/s"])
        vs = sorted(byv)
        lc = c0 + 3
        ws.cell(1, lc, f"{label}: geomean v"); ws.cell(1, lc + 1, "geomean k_r (/s)")
        for i, v in enumerate(vs, start=2):
            ws.cell(i, lc, v); ws.cell(i, lc + 1, gm(byv[v]))
        meta.append(dict(label=label, medium=med, c0=c0, n=len(pts), lc=lc, nline=len(vs)))
        c0 += 5

    xlo, xhi = 3.2, 10
    rc0 = c0
    ws.cell(1, rc0, "glass ref x"); ws.cell(1, rc0 + 1, "glass ref k_r (clean geomean)")
    ws.cell(2, rc0, xlo); ws.cell(2, rc0 + 1, fg); ws.cell(3, rc0, xhi); ws.cell(3, rc0 + 1, fg)
    ws.cell(1, rc0 + 2, "quartz ref x"); ws.cell(1, rc0 + 3, "quartz ref k_r (clean geomean)")
    ws.cell(2, rc0 + 2, xlo); ws.cell(2, rc0 + 3, fq); ws.cell(3, rc0 + 2, xhi); ws.cell(3, rc0 + 3, fq)
    return meta, rc0


def build_chart(ws, meta, rc0, fg, fq):
    ch = ScatterChart()
    ch.style = 2
    ch.scatterStyle = "marker"
    ch.varyColors = False

    def colref(c, r1, r2):
        return Reference(ws, min_col=c, min_row=r1, max_col=c, max_row=r2)

    def add_line(xcol, ycol, count, title, color, dash="solid", w=15875, smooth=False):
        if count == 0: return
        xref = colref(xcol, 2, 1 + count); yref = colref(ycol, 2, 1 + count)
        s = Series(values=yref, xvalues=xref, title=title)
        s.marker = Marker(symbol="none")
        s.graphicalProperties = GraphicalProperties()
        s.graphicalProperties.line = LineProperties(solidFill=color, w=w, prstDash=dash)
        s.smooth = smooth
        ch.series.append(s)

    def add_points(c0, count, title, symbol, size, fill_color):
        if count == 0: return
        xref = colref(c0, 2, 1 + count); yref = colref(c0 + 1, 2, 1 + count)
        s = Series(values=yref, xvalues=xref, title=title)
        s.marker = Marker(symbol=symbol, size=size)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=fill_color)
        s.marker.graphicalProperties.line = LineProperties(solidFill="000000", w=9525)
        s.graphicalProperties = GraphicalProperties()
        s.graphicalProperties.line = LineProperties(noFill=True)
        s.smooth = False
        ch.series.append(s)

    add_line(rc0, rc0 + 1, 2, "glass k_r (clean geomean)", GREY_GLASS, dash="dash")
    add_line(rc0 + 2, rc0 + 3, 2, "quartz k_r (clean geomean)", GREY_QUARTZ, dash="dash")

    for i, m in enumerate(meta):
        color = PAL[i % len(PAL)]
        symbol = "circle" if m["medium"] == "glass" else "square"
        add_line(m["lc"], m["lc"] + 1, m["nline"], f"{m['label']} (geomean)", color, dash="solid", w=12700, smooth=False)
        add_points(m["c0"], m["n"], m["label"], symbol, 9, color)

    ch.x_axis.scaling.min = 3.2
    ch.x_axis.scaling.max = 10
    ch.x_axis.scaling.logBase = 10
    ch.x_axis.delete = False
    ch.x_axis.axPos = "b"
    ch.x_axis.crossAx = 20
    ch.x_axis.majorTickMark = "out"
    ch.x_axis.minorTickMark = "none"
    ch.x_axis.tickLblPos = "nextTo"
    ch.x_axis.numFmt = "0"
    ch.x_axis.spPr = GraphicalProperties()
    ch.x_axis.spPr.line = LineProperties(solidFill="000000", w=12700)
    ch.x_axis.title = axis_title("pore velocity (m/day)", 0.36, 0.90)

    ch.y_axis.scaling.min = 5e-7
    ch.y_axis.scaling.max = 1e-2
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.delete = False
    ch.y_axis.axPos = "l"
    ch.y_axis.crossAx = 10
    ch.y_axis.majorTickMark = "out"
    ch.y_axis.minorTickMark = "none"
    ch.y_axis.tickLblPos = "nextTo"
    ch.y_axis.numFmt = "0.0000000"
    ch.y_axis.spPr = GraphicalProperties()
    ch.y_axis.spPr.line = LineProperties(solidFill="000000", w=12700)
    ch.y_axis.title = axis_title("fitted k_r (/s)", 0.01, 0.35)

    ch.x_axis.majorGridlines = None
    ch.y_axis.majorGridlines = None

    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.14, y=0.04, w=0.66, h=0.90))

    ch.legend = Legend()
    ch.legend.position = "r"
    ch.legend.overlay = False
    ch.legend.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.81, y=0.10, w=0.19, h=0.75))
    legpara = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1000, b=False)))
    legpara.pPr.defRPr.latin = DrawFont(typeface=APTOS)
    ch.legend.txPr = RichText(bodyPr=RichTextProperties(), p=[legpara])

    ch.graphical_properties = GraphicalProperties()
    ch.graphical_properties.line = LineProperties(noFill=True)
    ch.graphical_properties.noFill = True

    top_para = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True)))
    top_para.pPr.defRPr.latin = DrawFont(typeface=APTOS)
    ch.textProperties = RichText(bodyPr=RichTextProperties(), p=[top_para])

    ch.width = 19
    ch.height = 13
    return ch


def build_workbook(out_name):
    groups, fg, fq = read_source(SRC)
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    ws_data = wb.create_sheet("Data")
    meta, rc0 = write_data_sheet(ws_data, groups, fg, fq)
    ws_fig = wb.create_sheet("Figure")
    ch = build_chart(ws_data, meta, rc0, fg, fq)
    ws_fig.add_chart(ch, "A1")
    ws_fig.sheet_view.zoomScale = 100
    ws_fig.sheet_view.zoomScaleNormal = 100
    ws_fig.page_setup.orientation = "landscape"
    ws_fig.page_setup.fitToPage = True
    ws_fig.page_setup.fitToWidth = 1
    ws_fig.page_setup.fitToHeight = 1
    ws_fig.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Figure", offset=-1 * wb.sheetnames.index("Figure"))
    wb.active = 0
    out_path = f"{OUTDIR}/{out_name}"
    wb.save(out_path)
    print("saved", out_path, [(m["label"], m["n"]) for m in meta])


if __name__ == "__main__":
    build_workbook("Figure10_KrVsVelocity.xlsx")
