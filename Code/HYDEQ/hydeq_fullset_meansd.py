"""hydeq_fullset_meansd.py -- redraw the three full-set HYDEQ galleries with the measured data
shown as replicate MEAN +- SD instead of one series per replicate column.  The IHOP reference
curve is NOT plotted here: these galleries show the conventional structures only.

WHY A SEPARATE SCRIPT.  hydeq_export.py plots the replicates raw because the columns do not share
a sampling grid, so a mean would have to be manufactured by interpolation.  That objection does not
apply to the mean blocks already carried in UnfavorableMaster.xlsx: those are the same precomputed
per-study means the IHOP galleries use, so plotting them here makes the two figure families show
identical measured points rather than two different renderings of the same experiments.

Conditions with a single column are plotted as raw points -- a mean of one is the point itself --
but in the same deep red as the replicate means, so the data read as one series throughout.

STYLE follows manuscript Figure 1: structures identified by Hydeq number alone, panel identifiers
in medium-velocity-size-IS-study-column form, 18 pt type, breakthrough in pore volumes and
retention profiles in cm.  The x axis is annotated on the bottom row only.

INPUTS (no refitting, no fit cache needed):
  Manuscript/HYDEQ/HydeqComparisonFullSet.xlsx      -- measured + every model curve, per column
  Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx -- per-study mean +- sd blocks

OUTPUT: <figdir>/HydeqComparisonFullSet-<split>-b.png   (three files)

Usage:
    python3 hydeq_fullset_meansd.py
    python3 hydeq_fullset_meansd.py --figdir ../../Manuscript/HYDEQ --suffix -b
"""
import re, os, argparse, collections
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl

# Manuscript Figure 1 is an Excel chart set in 18 pt; match it so the SI galleries read the same
# size on the page.  The figure box is enlarged in proportion, or the type would crowd the axes.
FS = 18
plt.rcParams.update({"font.size": FS, "axes.labelsize": FS, "axes.titlesize": FS,
                     "xtick.labelsize": FS, "ytick.labelsize": FS, "legend.fontsize": FS})

FULL = "../../Manuscript/HYDEQ/HydeqComparisonFullSet.xlsx"
MASTER = "../../Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx"

BTEC_YLIM, RP_YLIM = (-6.4, 0.7), (5.3, 9.7)
BTEC_XLIM, RP_XLIM = (0.0, 10.0), (0.0, 20.0)     # RP in cm, as in manuscript Figure 1
MEAS_COLOR = "#c0392b"        # measured data are this deep red throughout, replicated or not
LABEL_WRAP = 38

PLOT = ["M2_2site", "M3_strain", "M4_dualpor", "M5_strain_block", "M6_strain_dp", "M7_all"]
# Hydeq 2 through 7, the same set shown in manuscript Figure 1.  All curves are DASHED, with
# distinct dash patterns and deep, high-contrast colours, so overlapping curves stay readable.
STYLE = {"M2_2site":        ("#b8860b", (0, (5, 1.5, 1, 1.5)), 2.4),   # deep gold
         "M3_strain":       ("#b30000", (0, (7, 2)),           2.4),   # deep red
         "M4_dualpor":      ("#0b3d91", (0, (3, 2)),           2.4),   # deep blue
         "M5_strain_block": ("#006d2c", (0, (1.6, 1.6)),       2.4),   # deep green
         "M6_strain_dp":    ("#4b0082", (0, (10, 2.5)),        2.4),   # deep indigo
         "M7_all":          ("#8e0152", (0, (9, 2, 1.6, 2)),   2.4)}   # deep magenta
SHORT = {"M2_2site": "Two kinetic sites, attachment and detachment",
         "M3_strain": "One kinetic site plus depth-dependent straining",
         "M4_dualpor": "One kinetic site plus dual porosity (slow-region velocity pinned)",
         "M5_strain_block": "One kinetic site, straining, and Langmuir blocking",
         "M6_strain_dp": "One kinetic site, straining, and dual porosity",
         "M7_all": "One kinetic site, straining, blocking, and dual porosity"}
HDR = dict(SHORT)                     # workbook column headers use the long names
HYDEQ_NO = {"M2_2site": "2", "M3_strain": "3", "M4_dualpor": "4",
            "M5_strain_block": "5", "M6_strain_dp": "6", "M7_all": "7"}

FIGS = [("Glass, 4 m/day", "glass-4mday", lambda m: m["medium"] == "glass" and m["vel"] == 4),
        ("Glass, 8 m/day", "glass-8mday", lambda m: m["medium"] == "glass" and m["vel"] == 8),
        ("Quartz", "quartz", lambda m: m["medium"] == "quartz")]

TITLE_RE = re.compile(r"^(?P<study>\w+)\s+(?P<medium>glass|quartz)\s+column\s+(?P<col>\w+)\s*--\s*"
                      r"(?P<size>[\d.]+)\s*um,\s*(?P<vel>[\d.]+)\s*m/day,\s*(?P<IS>[\d.]+)\s*mM")


def sizeclass(s):
    s = float(s)
    return 1.1 if 0.9 <= s <= 1.15 else round(s, 2)


def wrap_label(s):
    if len(s) <= LABEL_WRAP:
        return s
    line, out = "", []
    for p in s.split("-"):
        cand = p if not line else line + "-" + p
        if len(cand) > LABEL_WRAP and line:
            out.append(line); line = p
        else:
            line = cand
    out.append(line)
    return "\n".join(out)


def col_pairs(ws, hrow, label):
    """(x, y) arrays for the column pair whose header starts with `label`."""
    tx = ty = None
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(hrow, c).value or "")
        if h.startswith(label + " :") or h == label:
            if "pore volumes" in h or "distance" in h:
                tx = c
            else:
                ty = c
    if tx is None or ty is None:
        return None
    x, y = [], []
    for r in range(hrow + 1, ws.max_row + 1):
        a, b = ws.cell(r, tx).value, ws.cell(r, ty).value
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            x.append(a); y.append(b)
        elif x:
            break
    return np.array(x), np.array(y)


def measured(ws, hrow):
    x, y = [], []
    for r in range(hrow + 1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            x.append(a); y.append(b)
        elif x:
            break
    return np.array(x), np.array(y)


def section_rows(ws):
    """(btec_header_row, rp_header_row)"""
    bt = rp = None
    for r in range(1, ws.max_row + 1):
        v = str(ws.cell(r, 1).value or "")
        if v.startswith("pore volumes"):
            bt = r
        elif v.startswith("distance"):
            rp = r
    return bt, rp


def master_mean(um, med, size, vel, IS, study):
    """{'btec': (x, mean, sd), 'rp': (...)} per study, or None when there is no mean block."""
    name = f"{'gl' if med == 'glass' else 'qu'}_{sizeclass(size)}um_{vel:.0f}md_{IS:.0f}mM"
    if name not in um.sheetnames:
        return None
    ws = um[name]
    out = {}
    for tag, hrow, xs in (("btec", 5, 1.0), ("rp", 101, 100.0)):     # RP metres -> cm
        cx = cy = cs = None
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(hrow, c).value or "")
            if h.strip() == f"{study} x":
                cx = c
            elif h.startswith(f"{study} mean log10"):
                cy = c
            elif h.strip() == f"{study} sd dex" and cy is not None and c > cy:
                cs = c
        if not (cx and cy and cs):
            return None
        x, y, sd = [], [], []
        for r in range(hrow + 1, ws.max_row + 1):
            a = ws.cell(r, cx).value; b = ws.cell(r, cy).value; e = ws.cell(r, cs).value
            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                x.append(a * xs); y.append(b); sd.append(e if isinstance(e, (int, float)) else 0.0)
            elif isinstance(ws.cell(r, 1).value, str) and str(ws.cell(r, 1).value).strip():
                break
        out[tag] = (np.array(x), np.array(y), np.array(sd))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--full", default=FULL)
    ap.add_argument("--master", default=MASTER)
    ap.add_argument("--figdir", default="../../Manuscript/HYDEQ")
    ap.add_argument("--suffix", default="-b")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.full, data_only=True)
    um = openpyxl.load_workbook(args.master, data_only=True)

    cols = {}
    for sn in wb.sheetnames:
        if not sn.startswith("accum_"):
            continue
        ws = wb[sn]
        m = TITLE_RE.match(str(ws.cell(1, 1).value or ""))
        if not m:
            print("  (skip, unparsed title) " + sn); continue
        meta = dict(study=m["study"], medium=m["medium"], col=m["col"],
                    size=float(m["size"]), vel=float(m["vel"]), IS=float(m["IS"]))
        btr, rpr = section_rows(ws)
        rec = dict(meta=meta, btec_meas=measured(ws, btr), rp_meas=measured(ws, rpr), models={})
        for k in PLOT:
            b = col_pairs(ws, btr, HDR[k]); r = col_pairs(ws, rpr, HDR[k])
            if b and r:
                rec["models"][k] = (b, r)
        cols[sn] = rec

    cond = collections.defaultdict(list)
    for sn, rec in cols.items():
        mt = rec["meta"]
        cond[(mt["medium"], sizeclass(mt["size"]), mt["vel"], mt["IS"])].append(sn)

    for title, tag, sel in FIGS:
        rows = sorted([c for c in cond if sel(dict(medium=c[0], vel=c[2]))],
                      key=lambda c: (c[2], c[1], c[3]))
        if not rows:
            continue
        n = len(rows)
        meas_labeled = [False]                      # ONE "exp" legend entry for the whole figure
        fig, axes = plt.subplots(n, 2, figsize=(17.0, 2.55 * n + 1.9), squeeze=False,
                                 sharex="col")
        for i, ck in enumerate(rows):
            sns_ = sorted(cond[ck], key=lambda s: (cols[s]["meta"]["study"], cols[s]["meta"]["col"]))
            a0, a1 = axes[i][0], axes[i][1]
            med, size, vel, IS = ck

            studies = []
            for s in sns_:
                st = cols[s]["meta"]["study"]
                if st not in studies:
                    studies.append(st)
            studies.sort(key=lambda s: (s != "Li", s))

            # ---------------- measured: replicate mean +- sd where a mean block exists ----------
            for st in studies:
                mem = [s for s in sns_ if cols[s]["meta"]["study"] == st]
                mm = master_mean(um, med, size, vel, IS, st) if len(mem) > 1 else None
                if mm:
                    x, y, sd = mm["btec"]
                    a0.errorbar(x, y, yerr=sd, fmt="o", ms=6.5, color=MEAS_COLOR, mew=0, lw=0,
                                capsize=4.5, capthick=1.4, elinewidth=1.4, zorder=5,
                                label=None if meas_labeled[0] else "exp")
                    meas_labeled[0] = True
                    x, y, sd = mm["rp"]
                    a1.errorbar(x, y, yerr=sd, fmt="s", ms=7.0, color=MEAS_COLOR, mew=0, lw=0,
                                capsize=4.5, capthick=1.4, elinewidth=1.4, zorder=5)
                else:
                    for s in mem:
                        bx, by = cols[s]["btec_meas"]; rx, ry = cols[s]["rp_meas"]
                        a0.plot(bx, by, "o", ms=6.5, color=MEAS_COLOR, mew=0, zorder=5,
                                label=None if meas_labeled[0] else "exp")
                        meas_labeled[0] = True
                        a1.plot(np.asarray(rx) * 100.0, ry, "s", ms=7.0, color=MEAS_COLOR,
                                mew=0, zorder=5)

            # ---------------- models: averaged over the condition's columns ---------------------
            for mdl in PLOT:
                cs_, ls, lw = STYLE[mdl]
                bt = [cols[s]["models"][mdl][0] for s in sns_ if mdl in cols[s]["models"]]
                rp = [cols[s]["models"][mdl][1] for s in sns_ if mdl in cols[s]["models"]]
                if not bt:
                    continue
                a0.plot(bt[0][0], np.mean([b[1] for b in bt], axis=0), linestyle=ls, color=cs_,
                        lw=lw, zorder=3, label=HYDEQ_NO[mdl] if i == 0 else None)
                a1.plot(np.asarray(rp[0][0]) * 100.0, np.mean([r[1] for r in rp], axis=0),
                        linestyle=ls, color=cs_, lw=lw, zorder=3)

            # ---------------- axes, and the panel identifier in BOTH panels ---------------------
            letters = "-".join(cols[s]["meta"]["col"] for s in sns_)
            lab = wrap_label(f"{med}-{vel:.0f}-{size}-{IS:.0f}-{'-'.join(studies)}-{letters}")
            for ax, (xl, yl, xlab, ylab) in ((a0, (BTEC_XLIM, BTEC_YLIM, "Pore Volumes",
                                                   r"Log$_{10}$ C/C$_0$")),
                                             (a1, (RP_XLIM, RP_YLIM, "Distance (cm)",
                                                   r"Log$_{10}$ #"))):
                ax.set_xlim(*xl); ax.set_ylim(*yl)
                ax.set_ylabel(ylab, fontsize=FS)
                # x axis annotated on the BOTTOM row only; the rows above give that space back to
                # the data, which keeps the gallery from running too tall for its width.
                if i == n - 1:
                    ax.set_xlabel(xlab, fontsize=FS)
                else:
                    ax.tick_params(axis="x", labelbottom=False)
                ax.tick_params(labelsize=FS, length=6, width=1.2)
                ax.grid(alpha=0.25, lw=0.6)
                for sp in ax.spines.values():
                    sp.set_linewidth(1.2)
                ax.text(0.985, 0.955, lab, transform=ax.transAxes, ha="right", va="top",
                        fontsize=FS, fontweight="bold",
                        bbox=dict(fc="white", ec="0.7", lw=0.6, boxstyle="round,pad=0.3"))

        fig.suptitle(f"Conventional (Hydeq) structures — {title}", fontsize=FS + 2, y=0.997)
        h, lb = [], []
        for row in axes:
            for ax in row:
                for hh, ll in zip(*ax.get_legend_handles_labels()):
                    if ll not in lb:
                        h.append(hh); lb.append(ll)
        order = ["exp"] + [HYDEQ_NO[m] for m in PLOT]
        pairs = sorted(zip(h, lb), key=lambda t: order.index(t[1]) if t[1] in order else 99)
        h = [p[0] for p in pairs]; lb = [p[1] for p in pairs]
        fig.legend(h, lb, loc="lower center", ncol=7, fontsize=FS, frameon=False,
                   handlelength=3.0, columnspacing=2.0, bbox_to_anchor=(0.5, 0.0))
        fig.tight_layout(rect=[0, 0.05, 1, 0.978])
        fig.subplots_adjust(hspace=0.10)
        fn = os.path.join(args.figdir, f"HydeqComparisonFullSet-{tag}{args.suffix}.png")
        fig.savefig(fn, dpi=150, bbox_inches="tight"); plt.close(fig)
        print("wrote " + fn)


if __name__ == "__main__":
    main()
