"""
serial3_size_fit.py -- Serial-3 rerun of the Tong glass colloid-size series (unfavorable, 20 mM, elution-normal).
Mirror of parallel tong_size_fit2.py but with the SERIAL-3 generator (variant b, +k_r): c->w->g->a,
alpha_s once at c, FAST wall attach k_mw=a_m*kf, focus k2=f_x*kf (w->g), SLOW crawl attach k_mg=a_m*v_ns*kf,
diffusive release k_r (g->w). v_ns pinned 5%, r_s per size (same RS table as parallel).
Goal: (1) Serial-3 per-column free-k_r fits; (2) JOINT single shared-k_r across all 10 columns
-> Serial-3's OWN consistent k_r constant, compared to parallel (5.79e-5/s).
Requires DataFromLi&Tong.xlsx uploaded to the sandbox.
"""
import functools, numpy as np, openpyxl, re, warnings
from openpyxl.utils import get_column_letter
print=functools.partial(print,flush=True); warnings.filterwarnings("ignore")
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
DATA="../Data/DataFromLi&Tong.xlsx"; DAY=86400.0; L=0.2; theta=0.375
REV,T0,Vref=22.801836559387397,3.58,0.1667; AMP=REV*T0*theta*Vref
W_RP,W_SHAPE,W_PLAT,W_TAIL,W_TSLOPE=6.0,2.5,10.0,3.0,16.0; INJPV=T0/(L/Vref)
RS={(0.1,8):92.8,(0.2,4):87.9,(0.2,8):53.5,(0.5,4):35.6,(0.5,8):26.3,(1.1,4):22.9,(2.0,8):19.0}
PLACE={-6.0,-9.94,-15.33}
def isplace(v): return any(abs(v-p)<0.05 for p in PLACE)
wb=openpyxl.load_workbook(DATA,data_only=True); ws=wb["Microspheres Glass Beads Tong"]
def parse(c):
    br=None
    for r in range(6,13):
        x=ws.cell(r,c).value
        if isinstance(x,str) and x.strip().upper().startswith("PV"): br=r;break
    if br is None:return None
    H=" || ".join(str(ws.cell(r,c).value) for r in range(1,br) if ws.cell(r,c).value is not None);hl=H.lower()
    d={'cond':"unfav" if "unfav" in hl else ("fav" if "favor" in hl else "?"),
       'size':float(re.search(r'([\d.]+)\s*micron',hl).group(1)) if re.search(r'([\d.]+)\s*micron',hl) else None,
       'vel':float(re.search(r'([\d.]+)\s*m/day',hl).group(1)) if re.search(r'([\d.]+)\s*m/day',hl) else None,
       'IS':float(re.search(r'([\d.]+)\s*M\b',H).group(1)) if re.search(r'([\d.]+)\s*M\b',H) else None,
       'dg':"downgrad" in hl,'twoPV':bool(re.search(r'2\s*pv',hl))}
    m=re.search(r'([\d.]+)\s*E\s*([+-]?\d+)',H);d['C0']=float(m.group(1))*10**int(m.group(2)) if m else None
    pv=[];lc=[]
    for r in range(br+1,67):
        a=ws.cell(r,c).value;b=ws.cell(r,c+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isplace(b) and b<0.5: pv.append(float(a));lc.append(float(b))
    xx=[];rp=[]
    for r in range(68,78):
        a=ws.cell(r,c).value;b=ws.cell(r,c+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isplace(b): xx.append(float(a));rp.append(float(b))
    d['bt_pv']=np.array(pv);d['bt_log']=np.array(lc);d['x']=np.array(xx);d['logrp']=np.array(rp);d['col']=get_column_letter(c);return d
def starts():
    o=[]
    for c in range(1,ws.max_column+1):
        for r in range(6,13):
            v=ws.cell(r,c).value
            if isinstance(v,str) and v.strip().upper().startswith("PV"):o.append(c);break
    return o
usable=[d for d in (parse(c) for c in starts()) if d and d['cond']=="unfav" and not d['twoPV'] and not d['dg'] and d['IS']==0.02 and len(d['bt_pv'])>0 and len(d['x'])>0]

class Eng:   # SERIAL-3 (variant b)
    def __init__(s,vmday):s.v=vmday/DAY;s.nx=90;s.dx=L/90;s.dt=s.dx/s.v;s.pv=L/s.v
    def run(s,k1,a_s,a_m,fx,kr,D,vns=0.05,inj=INJPV,tot=10):
        kf=k1;k2=fx*kf;kmw=a_m*kf;kmg=a_m*vns*kf;G=np.zeros((4,4))
        G[0,0]-=kf;G[3,0]+=a_s*kf;G[1,0]+=(1-a_s)*kf                 # c: attach once, else graze->w
        G[1,1]-=(kmw+k2);G[3,1]+=kmw;G[2,1]+=k2                      # w: FAST attach, focus->g
        G[2,2]-=kmg;G[3,2]+=kmg                                      # g: SLOW crawl attach
        G[1,2]+=kr;G[2,2]-=kr                                        # release g->w
        s._k1=k1;s._kmw=kmw;s._kmg=kmg;E=expm(G*s.dt)
        r=D*s.dt/s.dx**2;mn=(1+2*r)*np.ones(90);of=-r*np.ones(89);mn[0]=1+r;mn[-1]=1+r
        lu=splu(csc_matrix(diags([of,mn,of],[-1,0,1],format="csc")))
        nt=int(round(tot*90));ti=inj*s.pv;c=vns;Y=np.zeros((4,90));C=np.zeros(nt);t=0.;Yst=None
        for i in range(nt):
            Y=E@Y;C[i]=(Y[0,-1]+Y[1,-1]+c*Y[2,-1])
            Y[0,1:]=Y[0,:-1];Y[0,0]=0;Y[1,1:]=Y[1,:-1];Y[1,0]=0
            ym=Y[2].copy();Y[2,1:]=ym[1:]-c*(ym[1:]-ym[:-1]);Y[2,0]=ym[0]*(1-c);t+=s.dt
            if t<ti:Y[0,0]+=1.0
            Y[0,:]=lu.solve(Y[0,:]);Y[1,:]=lu.solve(Y[1,:])
            if t<ti and t>0.9*ti:Yst=Y.copy()
        tp=(np.arange(nt)+1)*s.dt/s.pv;x=(np.arange(90)+0.5)*s.dx
        rp=(a_s*s._k1*Yst[0,:]+s._kmw*Yst[1,:]+s._kmg*Yst[2,:])
        m=(tp>0.5*inj)&(tp<inj);return dict(tp=tp,C=C,x=x,rp=rp,plat=float(C[m].mean()))
def sl(x,y):h=len(x)//2;return float(np.polyfit(x[:h+1],y[:h+1],1)[0]),float(np.polyfit(x[h:],y[h:],1)[0])
def ts(pv,lc,lo=6,hi=10):
    m=(np.asarray(pv)>=lo)&(np.asarray(pv)<=hi);return float("nan") if np.sum(m)<2 else float(np.polyfit(np.asarray(pv)[m],np.asarray(lc)[m],1)[0])
def setup(d):
    sz,v=d['size'],d['vel'];rs=RS[(sz,v)];col=Eng(v);V_MS=v/DAY;logK=np.log10(AMP*d['C0']);D=col.v*L/150
    si,so=sl(d['x'],d['logrp']);pm=(d['bt_pv']>1.3)&(d['bt_pv']<3.5);pp=d['bt_log'][pm]
    plo,phi=float(pp.min()),float(pp.max());tm=d['bt_pv']>4.0;sld=ts(d['bt_pv'],d['bt_log']);k1s=np.log10(rs*V_MS)
    def resid(lk1,la_s,la_m,lfx,kr):
        k1=10**lk1;r=col.run(k1,10**la_s,10**la_m,10**lfx,kr,D)
        rl=np.interp(d['x'],r['x'],np.maximum(r['rp'],1e-300))/V_MS;lS=logK+np.log10(np.maximum(rl,1e-300))
        rp=W_RP*(lS-d['logrp']);sm,sn=sl(d['x'],lS);sh=W_SHAPE*np.array([sm-si,sn-so])
        p_=np.log10(max(r['plat'],1e-12));ex=max(0.,p_-phi)+max(0.,plo-p_);pl=W_PLAT*np.array([ex])
        tl=W_TAIL*(np.log10(np.maximum(np.interp(d['bt_pv'][tm],r['tp'],r['C']),1e-12))-d['bt_log'][tm])
        slm=ts(r['tp'],np.log10(np.maximum(r['C'],1e-12)));tsl=W_TSLOPE*np.array([slm-sld])
        return np.concatenate([rp,sh,pl,tl,tsl]),r,lS
    return dict(col=col,V_MS=V_MS,logK=logK,D=D,k1s=k1s,resid=resid,d=d,sz=sz,v=v,rs=rs)
S=[setup(d) for d in sorted(usable,key=lambda z:(z['size'],z['vel']))]
print(f"SERIAL-3 size series | {len(S)} columns")
print(f"{'col':>4}{'sz':>5}{'v':>3}{'a_s':>8}{'a_m':>8}{'f_x':>8}{'k_r/PV':>8}{'cost':>7}")
percol=[]
for s in S:
    def rr(p):return s['resid'](p[0],p[1],p[2],p[3],10**p[4])[0]
    lo=[s['k1s']-0.3,-3.3,-3.3,np.log10(1e-4),np.log10(1e-6)];hi=[s['k1s']+0.3,np.log10(.9),np.log10(.999),np.log10(1.0),np.log10(.1)]
    best=None
    for st in [[s['k1s'],-1.5,-1.0,np.log10(.005),np.log10(1e-4)],[s['k1s'],-2.2,-0.3,np.log10(.02),np.log10(3e-4)]]:
        r=least_squares(rr,st,bounds=(lo,hi),max_nfev=60)
        if best is None or r.cost<best.cost:best=r
    p=best.x;kr=10**p[4];_,rmod,lS=s['resid'](p[0],p[1],p[2],p[3],kr);krpv=kr*(L/s['col'].v)
    xpk=float(rmod['x'][int(np.argmax(rmod['rp']))])
    percol.append(dict(s=s,p=p,kr=kr,krpv=krpv,cost=float(best.cost),rmod=rmod,lS=lS,xpk=xpk))
    print(f"{s['d']['col']:>4}{s['sz']:>5}{s['v']:>3.0f}{10**p[1]:>8.4f}{10**p[2]:>8.4f}{10**p[3]:>8.4f}{krpv:>8.3f}{best.cost:>7.2f}")
tot_free=sum(pc['cost'] for pc in percol);print(f"  total per-column cost (free k_r) = {tot_free:.2f}")
# JOINT shared-k_r
seed=[np.log10(np.median([pc['kr'] for pc in percol]))]
for pc in percol: seed+=[pc['p'][0],pc['p'][1],pc['p'][2],pc['p'][3]]
lo=[np.log10(1e-6)];hi=[np.log10(.1)]
for s in S: lo+=[s['k1s']-0.3,-3.3,-3.3,np.log10(1e-4)];hi+=[s['k1s']+0.3,np.log10(.9),np.log10(.999),np.log10(1.0)]
def jresid(P):
    kr=10**P[0];out=[]
    for j,s in enumerate(S):
        lk1,las,lam,lfx=P[1+4*j:1+4*j+4];out.append(s['resid'](lk1,las,lam,lfx,kr)[0])
    return np.concatenate(out)
jr=least_squares(jresid,seed,bounds=(lo,hi),max_nfev=120);kr_sh=10**jr.x[0];tot_joint=float(jr.cost)
print(f"\nSERIAL-3 JOINT shared-k_r: k_r = {kr_sh:.3e}/s  (parallel was 5.79e-5/s)")
for v in sorted({s['v'] for s in S}): print(f"    = {kr_sh*(L/(v/DAY)):.3f} /PV at {v:.0f} m/day")
print(f"  total joint cost (shared k_r) = {tot_joint:.2f}  vs per-column free = {tot_free:.2f}  (D={tot_joint-tot_free:+.2f}, {100*(tot_joint-tot_free)/tot_free:+.1f}%)")
import json; json.dump({"kr_shared_per_s":kr_sh,"tot_joint":tot_joint,"tot_free":tot_free,
    "percol":[{"col":pc['s']['d']['col'],"size":pc['s']['sz'],"vel":pc['s']['v'],"kr_per_s":pc['kr'],
    "kr_per_PV":pc['krpv'],"cost":pc['cost'],"xpeak_cm":100*pc['xpk']} for pc in percol]},
    open("artifacts/serial3_size_results.json","w"),indent=1)
print("wrote serial3_size_results.json")
