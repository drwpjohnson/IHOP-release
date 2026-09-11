"""weight_robustness.py -- do the manuscript's CONCLUSIONS survive a change of objective weights?

THE QUESTION.  The production objective is a five-block weighted log10 residual with
W_RP = 6.0, W_SHAPE = 2.5, W_PLAT = 20.0, W_TAIL = 3.0, W_TSLOPE = 16.0 and a +-0.25 dex plateau
dead zone.  Those five numbers are the most visible discretionary choice left in the paper and no
systematic check of them was ever recorded (the working record's open-items list names them first).
A reviewer will ask.  This script answers the question that actually matters, which is NOT whether
the fitted parameters are stationary under reweighting -- they will not be -- but whether the
statements the manuscript makes about them survive.

WHAT IS SCORED, therefore, is the CONCLUSION SET, not the parameter values:
    T1  alpha_s rises with IS            (Figure 7)
    T2  k_r vs colloid size              (Figure 9 -- manuscript claims NO strong dependence)
    T3  k_r vs velocity                  (Figure 10 -- manuscript claims NO strong dependence)
    T4  f_x falls with |U_sec|           (Figure 8 closure)
    T5  RP RMS stays small               (fit quality, reported in dex)
    T6  branch/peak reproduction         (the structural claim against the conventional models)
Each is reported as a Spearman rho with its sign, plus the medians.  A conclusion "survives" if its
sign and rough magnitude are unchanged; the script does not decide that for you, it tabulates it.

WHY THE BASELINE RUN MATTERS.  Variant `prod` re-fits with the production weights through this
script's own code path.  If it does not reproduce UnfavorableMaster.xlsx, then differences seen in
the other variants are this script's, not the weights'.  Check `prod` before reading anything else.

THE VARIANTS, and why each is here:
    prod        production weights -- the control
    flat        all five weights = 1.0 -- the extreme.  Surviving this is a much stronger statement
                than surviving a +-20% perturbation, and it is one run.
    shape_half  W_SHAPE 2.5 -> 1.25
    shape_x2    W_SHAPE 2.5 -> 5.0
    unitfix     RP-shape residuals multiplied by L = 0.2 m so they are in DEX rather than dex/m.
                ** This is the variant to read first. **  The shape block's residuals are LOG-SLOPES
                (d log10 S / dx), whose raw magnitude on this grid is ~1.1 dex/m -- fifteen times the
                RP-level residuals (~0.07 dex) and fifty times the tail-slope residuals.  Because the
                blocks are in different units, the weight numbers badly misrepresent the pull: at the
                production fits the shape block carries 83% of the realised cost while holding only
                1.1% of the nominal w^2*n budget, and the plateau block -- nominal weight 20, the
                largest -- contributes 1.9% because its dead zone leaves it inactive on 26 of 29
                columns.  See Records/weight_robustness.md and scratch script block_pull.py.
    plat_x2     W_PLAT 20 -> 40      (tests the dead-zone block)
    tslope_half W_TSLOPE 16 -> 8

WHAT THIS SCRIPT DELIBERATELY DOES NOT DO.  It does not re-run the replicate transfer test under
each variant.  Transfer RMS is scored under the production objective and the headline result (transfer
sitting at the reproducibility floor) is conditional on the production weights; re-deriving it per
variant would multiply the work without changing the robustness question.  State that conditionality
in the manuscript rather than chasing it here.

RUNTIME.  About 2 s per column per variant, so ~60 s per variant over the 29 unfavorable columns and
~8 min for all seven.  That exceeds some tool call caps, so variants can be run in batches and merged:
each run dumps its fits with --json, and --merge rebuilds the combined workbook from those dumps
without refitting.

Usage:
    python3 weight_robustness.py --variants prod,flat,unitfix        # the three that matter
    python3 weight_robustness.py --variants all --out wr.xlsx        # ~8 min
    python3 weight_robustness.py --variants prod --columns Li.O,Li.R # quick smoke test
    python3 weight_robustness.py --variants prod,flat --json a.json  # run in batches ...
    python3 weight_robustness.py --merge a.json,b.json --out wr.xlsx # ... then merge, no refit
"""
import csv, collections, argparse, json, sys, time
import numpy as np
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
from scipy.stats import spearmanr
import openpyxl
from openpyxl.styles import Font

CSV = "../Data/LiTong_experimental_data_tidy.csv"
MASTER = "../Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx"
FXTREND = "../Manuscript/FigsExcelsUnfav/fx_trend.xlsx"
OUT = "../Manuscript/FigsExcelsUnfav/WeightRobustness.xlsx"

DAY = 86400.; L = 0.2
REV, T0, Vref = 22.801836559387397, 3.58, 0.1667
INJPV = T0 / (L / Vref)
THETA = {"glass": 0.375, "quartz": 0.36}
KR_MED = {"glass": 5.58e-5, "quartz": 1.87e-5}   # multistart SEED only, copied verbatim from unfav_master_fit.py; superseded by 3.75e-5/1.82e-5 as the reported constants and deliberately NOT refreshed here, because changing a seed perturbs every converged fit and would break the baseline reproduction this script depends on
PLAT_TOL = 0.25
NIN_PEAKED = 4

# name -> (W_RP, W_SHAPE, W_PLAT, W_TAIL, W_TSLOPE, shape_scale)
# shape_scale multiplies the RP-shape residual; 1.0 is production, L converts dex/m -> dex.
VARIANTS = {
    "prod":        (6.0, 2.50, 20.0, 3.0, 16.0, 1.0),
    "flat":        (1.0, 1.00,  1.0, 1.0,  1.0, 1.0),
    "shape_half":  (6.0, 1.25, 20.0, 3.0, 16.0, 1.0),
    "shape_x2":    (6.0, 5.00, 20.0, 3.0, 16.0, 1.0),
    "unitfix":     (6.0, 2.50, 20.0, 3.0, 16.0, L),
    "plat_x2":     (6.0, 2.50, 40.0, 3.0, 16.0, 1.0),
    "tslope_half": (6.0, 2.50, 20.0, 3.0,  8.0, 1.0),
}

kB = 1.381e-23; gg = 9.806; RHO_C = 1055.; RHO_F = 998.; MU = 9.8e-4; TK = 298.2; DP = 0.510e-3
HAM = {"glass": 7.17e-21, "quartz": 1.96e-20}


def TE_rs(size_um, vel, med):
    dc = size_um * 1e-6; U = vel / DAY; th = THETA[med]; Ui = U * th
    gam = (1 - th) ** (1 / 3); Hm = HAM[med]
    As = 2 * (1 - gam ** 5) / (2 - 3 * gam + 3 * gam ** 5 - 2 * gam ** 6)
    D = kB * TK / (6 * np.pi * MU * (dc / 2))
    NR = dc / DP; NvdW = Hm / (kB * TK)
    NG = 2 * (dc / 2) ** 2 * (RHO_C - RHO_F) * gg / (9 * MU * Ui)
    NA = Hm / (12 * np.pi * MU * (DP / 2) ** 2 * Ui); NPe = Ui * DP / D
    TE = (2.4 * As ** (1 / 3) * NR ** (-0.081) * NPe ** (-0.715) * NvdW ** 0.052
          + 0.55 * As * NR ** 1.675 * NA ** 0.125
          + 0.22 * NR ** (-0.24) * NG ** 1.11 * NvdW ** 0.053)
    return -1.5 * (gam / DP) * np.log(1 - TE) * 0.77


def sizeclass(s):
    s = float(s)
    return 1.1 if 0.9 <= s <= 1.15 else round(s, 2)


FAVFIT = {("glass", 0.5, 4.0): 47.5, ("glass", 1.1, 4.0): 30.2, ("glass", 2.0, 8.0): 15.2,
          ("glass", 1.1, 2.0): 49.9, ("glass", 1.1, 8.0): 29.6, ("quartz", 0.5, 8.0): 56.7,
          ("quartz", 1.1, 2.0): 120.6, ("quartz", 1.1, 4.0): 62.3, ("quartz", 1.1, 8.0): 68.9}
SEED_CQL = {("glass", 0.2, 4.0): 49.0, ("glass", 0.1, 8.0): 62.4, ("glass", 0.2, 8.0): 39.1}


def seed_rs(med, size, vel):
    k = (med, sizeclass(size), vel)
    if k in FAVFIT: return FAVFIT[k], "fav"
    if k in SEED_CQL: return SEED_CQL[k], "favRP"
    return TE_rs(size, vel, med), "TE"


class Eng:
    def __init__(s, vmday):
        s.v = vmday / DAY; s.dx = L / 90; s.dt = s.dx / s.v; s.pv = L / s.v

    def run(s, k1, a_s, a_m, fx, kr, vns=0.05, tot=10):
        kf = k1; k2 = fx * kf; kmw = a_m * kf; kmg = a_m * vns * kf
        G = np.zeros((4, 4))
        G[0, 0] -= kf; G[3, 0] += a_s * kf; G[1, 0] += (1 - a_s) * kf
        G[1, 1] -= (kmw + k2); G[3, 1] += kmw; G[2, 1] += k2
        G[2, 2] -= kmg; G[3, 2] += kmg; G[1, 2] += kr; G[2, 2] -= kr
        E = expm(G * s.dt); D = s.v * L / 150
        r = D * s.dt / s.dx ** 2
        mn = (1 + 2 * r) * np.ones(90); of = -r * np.ones(89); mn[0] = 1 + r; mn[-1] = 1 + r
        lu = splu(csc_matrix(diags([of, mn, of], [-1, 0, 1], format="csc")))
        nt = int(round(tot * 90)); ti = INJPV * s.pv; c = vns
        Y = np.zeros((4, 90)); C = np.zeros(nt); t = 0.
        for i in range(nt):
            Y = E @ Y; C[i] = Y[0, -1] + Y[1, -1] + c * Y[2, -1]
            Y[0, 1:] = Y[0, :-1]; Y[0, 0] = 0; Y[1, 1:] = Y[1, :-1]; Y[1, 0] = 0
            ym = Y[2].copy(); Y[2, 1:] = ym[1:] - c * (ym[1:] - ym[:-1]); Y[2, 0] = ym[0] * (1 - c)
            t += s.dt
            if t < ti: Y[0, 0] += 1.
            Y[0, :] = lu.solve(Y[0, :]); Y[1, :] = lu.solve(Y[1, :])
        tp = (np.arange(nt) + 1) * s.dt / s.pv; x = (np.arange(90) + 0.5) * s.dx
        return dict(tp=tp, C=np.maximum(C, 1e-300), x=x, rp=np.maximum(Y[3, :] / ti, 1e-300))


def sl(x, y, nin=None):
    if nin is None:
        h = len(x) // 2
        return float(np.polyfit(x[:h + 1], y[:h + 1], 1)[0]), float(np.polyfit(x[h:], y[h:], 1)[0])
    return float(np.polyfit(x[:nin], y[:nin], 1)[0]), float(np.polyfit(x[nin - 1:], y[nin - 1:], 1)[0])


def rp_branch(rlog):
    return "peaked" if rlog[1] > rlog[0] else "non-peaking"


def condition_branch(bs):
    nm = sum(1 for b in bs if b == "peaked")
    return "peaked" if nm > len(bs) - nm else "non-peaking"


def ts(pv, lc, lo=6, hi=10):
    m = (np.asarray(pv) >= lo) & (np.asarray(pv) <= hi)
    return float("nan") if np.sum(m) < 2 else float(np.polyfit(np.asarray(pv)[m], np.asarray(lc)[m], 1)[0])


# ---------------- data ----------------
def load():
    rows = list(csv.DictReader(open(CSV)))
    srcs = set(r["source"] for r in rows)
    assert len(rows) >= 1200 and "Li" in srcs and "Tong" in srcs, (
        f"CSV at {CSV!r} looks like a subset ({len(rows)} rows, sources={srcs}) -- refusing to fit "
        "a silently-truncated dataset (same guard as unfav_master_fit.py).")
    cols = collections.defaultdict(lambda: {"BTEC": [], "RP": []}); meta = {}
    for r in rows:
        cu = "BTEC" if r["curve"] == "BTEC_ND" else r["curve"]
        k = (r["source"], r["medium"], r["workbook_col"])
        cols[k][cu].append((float(r["x"]), float(r["value"])))
        meta[k] = dict(medium=r["medium"], chem=r["chemistry"], size=float(r["colloid_um"]),
                       vel=float(r["velocity_mday"]),
                       IS=(None if r["IS_mM"] == "" else float(r["IS_mM"])),
                       C0=float(r["C0_per_mL"]))
    for k in cols:
        for cu in ("BTEC", "RP"):
            cols[k][cu] = np.array(sorted(cols[k][cu])) if cols[k][cu] else np.zeros((0, 2))
    return cols, meta


def unfav_keys(cols, meta):
    """The same set unfav_master_fit fits: unfavorable, IS > 1 mM, both curves >= 3 points."""
    ks = [k for k, m in meta.items()
          if m["chem"] == "unfavorable" and m["IS"] is not None and m["IS"] > 1.0
          and len(cols[k]["BTEC"]) >= 3 and len(cols[k]["RP"]) >= 3]
    return sorted(ks)


def condition_windows(cols, meta, ks):
    """Condition-level majority branch -> the nin each column is SCORED with."""
    grp = collections.defaultdict(list)
    for k in ks:
        m = meta[k]
        grp[(m["medium"], sizeclass(m["size"]), m["vel"], m["IS"])].append(k)
    nin = {}
    for g, gk in grp.items():
        b = condition_branch([rp_branch(cols[k]["RP"][:, 1]) for k in gk])
        for k in gk:
            nin[k] = NIN_PEAKED if b == "peaked" else None
    return nin


# ---------------- fit one column under one weight set ----------------
def fit_col(k, cols, meta, nin, wt):
    W_RP, W_SHAPE, W_PLAT, W_TAIL, W_TSLOPE, SHP_SC = wt
    m = meta[k]
    bt = cols[k]["BTEC"].copy(); rp = cols[k]["RP"]
    med = m["medium"]; vel = m["vel"]; C0 = m["C0"]
    V_MS = vel / DAY; eng = Eng(vel); krs = np.log10(KR_MED[med])
    logK = np.log10(REV * T0 * THETA[med] * C0 * Vref)
    rs0, src = seed_rs(med, m["size"], vel); kk = rs0 * V_MS
    bt[:, 1] = np.maximum(bt[:, 1], -6.0)
    pv = bt[:, 0]; lc = bt[:, 1]; rx = rp[:, 0]; rlog = rp[:, 1]
    si, so = sl(rx, rlog, nin)
    wm = (pv > 1.2) & (pv < 4.0); pvw = pv[wm] if wm.sum() else pv
    dmean = float(lc[wm].mean()) if wm.sum() else float(lc.mean())
    tm = pv > 4.2; sld = ts(pv, lc)

    def resid(p):
        a_s = 1 - 10 ** p[0]; am = 10 ** p[1]; fx = 10 ** p[2]; krr = 10 ** p[3]
        r = eng.run(kk, a_s, am, fx, krr)
        rl = np.interp(rx, r["x"], np.maximum(r["rp"], 1e-300)) / V_MS
        lS = logK + np.log10(np.maximum(rl, 1e-300))
        rpx = W_RP * (lS - rlog)
        sm, sn = sl(rx, lS, nin)
        shp = W_SHAPE * SHP_SC * np.array([sm - si, sn - so])
        mlog = np.mean(np.log10(np.maximum(np.interp(pvw, r["tp"], r["C"]), 1e-6)))
        dev = mlog - dmean
        pl = W_PLAT * np.array([np.sign(dev) * max(0., abs(dev) - PLAT_TOL)])
        tl = (W_TAIL * (np.log10(np.maximum(np.interp(pv[tm], r["tp"], r["C"]), 1e-6)) - lc[tm])
              if tm.sum() else np.array([0.]))
        slm = ts(r["tp"], np.log10(np.maximum(r["C"], 1e-6)))
        tsl = W_TSLOPE * np.array([slm - sld if sld == sld else 0.])
        return np.concatenate([rpx, shp, pl, tl, tsl])

    lo = [-3.3, -3.3, np.log10(1e-4), np.log10(1e-6)]
    hi = [np.log10(1 - 1e-4), np.log10(.999), np.log10(1.0), np.log10(.1)]
    best = None
    for s0 in [[-1.5, -1., np.log10(.005), krs],
               [-.5, -.3, np.log10(.02), krs],
               [-2.3, -1.5, np.log10(.001), krs - 0.5]]:
        rr = least_squares(resid, s0, bounds=(lo, hi), max_nfev=80)
        if best is None or rr.cost < best.cost:
            best = rr
    a_s = 1 - 10 ** best.x[0]; am = 10 ** best.x[1]
    fx = 10 ** best.x[2]; kr = 10 ** best.x[3]

    # weight-free diagnostics
    r = eng.run(kk, a_s, am, fx, kr)
    lS = logK + np.log10(np.maximum(np.interp(rx, r["x"], np.maximum(r["rp"], 1e-300)) / V_MS, 1e-300))
    rp_rms = float(np.sqrt(np.mean((lS - rlog) ** 2)))
    smask = pv > 1.2
    bt_rms = float(np.sqrt(np.mean((np.log10(np.maximum(np.interp(pv[smask], r["tp"], r["C"]), 1e-6))
                                    - lc[smask]) ** 2)))
    xm = r["x"] <= float(rx.max())
    peak_model = float(r["x"][xm][int(np.argmax(np.log10(r["rp"][xm])))]) * 100.0
    peak_meas = float(rx[int(np.argmax(rlog))]) * 100.0
    return dict(study=k[0], medium=med, column=k[2], size=m["size"], vel=vel, IS=m["IS"],
                rs=rs0, rs_src=src, a_s=a_s, a_m=am, fx=fx, kr_s=kr,
                kr_pv=kr * (L / eng.v), cost=float(best.cost), rp_rms=rp_rms, bt_rms=bt_rms,
                peak_model_cm=peak_model, peak_meas_cm=peak_meas,
                window=("half" if nin is None else f"first{nin}"))


# ---------------- conclusion scoring ----------------
def usec_map():
    """|U_sec| in kT per column, from fx_trend.xlsx ('f_x fitted')."""
    try:
        ws = openpyxl.load_workbook(FXTREND, data_only=True)["f_x fitted"]
    except Exception as e:
        print(f"  (no |U_sec| available: {e}) -- T4 will be skipped")
        return {}
    hdr = None; out = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "study":
            hdr = [str(v) for v in row]; continue
        if hdr and row and row[0] in ("Li", "Tong"):
            d = dict(zip(hdr, row))
            col = str(d.get("col", "")).split(".")[-1]
            try:
                out[(d["study"], d["medium"], col)] = float(d["|U_sec|_kT"])
            except Exception:
                pass
    return out


def conclusions(fits, usec):
    """Spearman rho for each trend claim.  Returns {label: (rho, p, n)}."""
    def sp(xs, ys):
        xs = np.asarray(xs, float); ys = np.asarray(ys, float)
        ok = np.isfinite(xs) & np.isfinite(ys)
        if ok.sum() < 4: return (float("nan"), float("nan"), int(ok.sum()))
        r, p = spearmanr(xs[ok], ys[ok])
        return (float(r), float(p), int(ok.sum()))

    out = {}
    for med in ("glass", "quartz"):
        f = [d for d in fits if d["medium"] == med]
        out[f"T1 alpha_s vs IS [{med}]"] = sp([d["IS"] for d in f], [d["a_s"] for d in f])
    out["T2 k_r vs size"] = sp([d["size"] for d in fits], [d["kr_s"] for d in fits])
    out["T3 k_r vs velocity"] = sp([d["vel"] for d in fits], [d["kr_s"] for d in fits])
    uu = [(usec.get((d["study"], d["medium"], d["column"])), d["fx"]) for d in fits]
    uu = [(a, b) for a, b in uu if a is not None]
    out["T4 f_x vs |U_sec|"] = sp([a for a, _ in uu], [b for _, b in uu])
    return out


def report(allfits, names, usec, out_path):
    ks = allfits[names[0]]

    # ---- console summary ----
    print("\n=== CONCLUSION SET (Spearman rho; sign is what matters) ===")
    concs = {v: conclusions(allfits[v], usec) for v in names}
    labels = list(concs[names[0]])
    print(f"{'claim':28}" + "".join(f"{v:>14}" for v in names))
    for lab in labels:
        cells = []
        for v in names:
            r, p, n = concs[v][lab]
            cells.append(f"{r:+.2f} (p{p:.2f})" if r == r else "n/a")
        print(f"{lab:28}" + "".join(f"{c:>14}" for c in cells))
    print(f"\n{'median RP RMS (dex)':28}" + "".join(
        f"{np.median([d['rp_rms'] for d in allfits[v]]):>14.3f}" for v in names))
    print(f"{'median BTEC RMS (dex)':28}" + "".join(
        f"{np.median([d['bt_rms'] for d in allfits[v]]):>14.3f}" for v in names))
    print(f"{'peak within 1 cm':28}" + "".join(
        f"{sum(1 for d in allfits[v] if abs(d['peak_model_cm']-d['peak_meas_cm'])<=1.0):>10d}/{len(ks):<3d}"
        for v in names))

    # ---- workbook ----
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    B = Font(bold=True); IT = Font(italic=True)
    ws = wb.create_sheet("Conclusions")
    ws.cell(1, 1, "Objective-weight robustness -- do the manuscript's conclusions survive reweighting?").font = B
    ws.cell(2, 1, "Spearman rho between the stated variables, one row per claim, one column per weight "
                  "variant. SIGN and rough magnitude are what matter; parameter values are expected to "
                  "move. Reproducer: Code/weight_robustness.py. Record: Records/weight_robustness.md.").font = IT
    ws.cell(3, 1, "'unitfix' rescales the RP-shape residual from dex/m to dex; at production weights that "
                  "block carries 83% of the realised cost on 1.1% of the nominal budget.").font = IT
    r0 = 5
    ws.cell(r0, 1, "claim").font = B
    for j, v in enumerate(names):
        ws.cell(r0, 2 + 2 * j, f"{v} rho").font = B
        ws.cell(r0, 3 + 2 * j, f"{v} p").font = B
    for i, lab in enumerate(labels):
        ws.cell(r0 + 1 + i, 1, lab)
        for j, v in enumerate(names):
            rr, pp, nn = concs[v][lab]
            ws.cell(r0 + 1 + i, 2 + 2 * j, None if rr != rr else round(rr, 4))
            ws.cell(r0 + 1 + i, 3 + 2 * j, None if pp != pp else round(pp, 4))
    rr = r0 + 2 + len(labels)
    for lab, fn in (("median RP RMS (dex)", lambda f: np.median([d["rp_rms"] for d in f])),
                    ("median BTEC RMS (dex)", lambda f: np.median([d["bt_rms"] for d in f])),
                    ("median cost", lambda f: np.median([d["cost"] for d in f])),
                    ("peak within 1 cm (of %d)" % len(ks),
                     lambda f: sum(1 for d in f if abs(d["peak_model_cm"] - d["peak_meas_cm"]) <= 1.0))):
        ws.cell(rr, 1, lab).font = B
        for j, v in enumerate(names):
            ws.cell(rr, 2 + 2 * j, round(float(fn(allfits[v])), 4))
        rr += 1
    ws.column_dimensions["A"].width = 32

    ws = wb.create_sheet("Weights")
    ws.cell(1, 1, "variant").font = B
    for j, h in enumerate(["W_RP", "W_SHAPE", "W_PLAT", "W_TAIL", "W_TSLOPE", "shape scale"]):
        ws.cell(1, 2 + j, h).font = B
    for i, v in enumerate(VARIANTS):
        ws.cell(2 + i, 1, v)
        for j, x in enumerate(VARIANTS[v]):
            ws.cell(2 + i, 2 + j, x)

    for v in names:
        ws = wb.create_sheet(f"fits {v}"[:31])
        hdr = ["study", "medium", "column", "size (um)", "v (m/day)", "IS (mM)", "r_s (/m, fixed)",
               "r_s source", "alpha_s", "alpha_m", "f_x", "k_r (/s)", "k_r (/PV)", "cost",
               "RP RMS (dex)", "BTEC RMS (dex)", "peak cm (model)", "peak cm (measured)", "window"]
        for j, h in enumerate(hdr): ws.cell(1, 1 + j, h).font = B
        for i, d in enumerate(allfits[v]):
            for j, key in enumerate(["study", "medium", "column", "size", "vel", "IS", "rs", "rs_src",
                                     "a_s", "a_m", "fx", "kr_s", "kr_pv", "cost", "rp_rms", "bt_rms",
                                     "peak_model_cm", "peak_meas_cm", "window"]):
                x = d[key]
                ws.cell(2 + i, 1 + j, round(x, 6) if isinstance(x, float) else x)

    wb.save(out_path)
    print(f"\nwrote {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--variants", default="prod,flat,unitfix",
                    help="comma list, or 'all' (default: prod,flat,unitfix)")
    ap.add_argument("--columns", default="", help="restrict to e.g. Li.O,Li.R (smoke test)")
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--json", default="", help="also dump raw fits here")
    ap.add_argument("--merge", default="",
                    help="comma list of --json dumps to combine into one workbook; no refitting")
    args = ap.parse_args()

    if args.merge:
        allfits = {}
        for f in args.merge.split(","):
            allfits.update(json.load(open(f.strip())))
        names = [v for v in VARIANTS if v in allfits]     # canonical order
        print(f"merged {len(names)} variants from {args.merge}: {', '.join(names)}\n")
        report(allfits, names, usec_map(), args.out)
        return

    names = list(VARIANTS) if args.variants == "all" else [v.strip() for v in args.variants.split(",")]
    for v in names:
        if v not in VARIANTS:
            sys.exit(f"unknown variant {v!r}; choose from {list(VARIANTS)}")

    cols, meta = load()
    ks = unfav_keys(cols, meta)
    nin = condition_windows(cols, meta, ks)
    if args.columns:
        want = set(args.columns.split(","))
        ks = [k for k in ks if f"{k[0]}.{k[2]}" in want]
    usec = usec_map()
    print(f"{len(ks)} unfavorable columns; variants: {', '.join(names)}\n")

    allfits = {}
    for v in names:
        t0 = time.time()
        fits = [fit_col(k, cols, meta, nin[k], VARIANTS[v]) for k in ks]
        allfits[v] = fits
        print(f"  {v:12} {len(fits)} columns in {time.time()-t0:.0f}s   "
              f"median RP RMS {np.median([d['rp_rms'] for d in fits]):.3f} dex   "
              f"median BTEC RMS {np.median([d['bt_rms'] for d in fits]):.3f} dex")

    if args.json:
        json.dump(allfits, open(args.json, "w"), indent=1)
        print(f"wrote {args.json}")
    report(allfits, names, usec, args.out)


if __name__ == "__main__":
    main()
