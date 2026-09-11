#!/usr/bin/env python3
"""Figure 6 (IHOP k_r=0 vs k_r-fit) driver.

2026-09-01 (W.P.J.): switched from plotting one condition-level MEAN-parameter fit per panel to
plotting each replicate's OWN individual fit (the source workbook, fits_canonical_IHOP.xlsx, already
has both -- MEAN columns and per-column "<Author>.<col> no-k_r fit" / "<Author>.<col> +k_r fit"
columns; the original build only read the MEAN ones).

SAME COLUMN LAYOUT AS THE ORIGINAL (W.P.J.: keep the format identical -- exp / k_r=0 / k_r fit, one
column pair each, BLOCK_W=7, `build_chart` from build_fig4_si5.py unchanged). A first attempt gave
each replicate its own column pair (variable block width, a new chart-builder) -- reverted; that
changed the format for no reason. Instead, a condition with N replicates now stacks each replicate's
fit curve as successive rows in the SAME column pair, with one BLANK ROW between replicates so the
scatter chart's connecting line breaks there instead of drawing a spurious diagonal from the end of
one replicate's curve to the start of the next (Excel's scatter default is to leave a gap on a blank
cell, not connect through it). Experimental points are unchanged: one combined, sorted series per
panel (markers only, no connecting line, so merging replicates there was always harmless).
"""
import openpyxl, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))  # run from Code/
from build_fig4_si5 import build_chart

SRC = "../Manuscript/FigsExcelsUnfav/fits_canonical_IHOP.xlsx"   # run from Code/
OUTDIR = "../Manuscript/HYDEQ"                                    # Figure 6 lives here, NOT FigsExcelsFavorable
wb_src = openpyxl.load_workbook(SRC, data_only=True)

# panel spec, matching Figure 1's exact row order/identifiers
PANELS = [
    ("gl_1.1um_4md_6mM",  "glass-4-1.1-6-Li-R"),
    ("gl_1.1um_4md_20mM", "glass-4-1.1-20-Li-O"),
    ("qu_1.1um_4md_3mM",  "quartz-4-1.1-3-Li-AE"),
    ("qu_1.1um_4md_6mM",  "quartz-4-1.1-6-Li-AB"),
    ("qu_1.1um_4md_20mM", "quartz-4-1.1-20-Li-S"),
]
LEGEND = ("k_r = 0", "k_r fit")


def read_block(ws, hdr_row, data_row_start, unit=1.0):
    """exp_pts: combined across replicates, as always. nokr_segs/krfit_segs: ONE point-list PER
    REPLICATE (not merged), in column order -- MEAN columns are skipped entirely."""
    hdr = [ws.cell(hdr_row, c).value for c in range(1, ws.max_column + 1)]
    exp_pts = []
    nokr_segs = []   # [[ (x,y), ... ], ...] one list per replicate
    krfit_segs = []

    def read_series(c):
        r = data_row_start; pts = []
        while True:
            x = ws.cell(r, c - 1).value; y = ws.cell(r, c).value
            if x is None and y is None: break
            if isinstance(x, (int, float)) and isinstance(y, (int, float)): pts.append((x * unit, y))
            r += 1
        return pts

    for c, h in enumerate(hdr, start=1):
        if not h: continue
        if h.endswith("exp"):
            exp_pts.extend(read_series(c))
        elif h == "MEAN no-k_r fit" or h == "MEAN +k_r fit":
            continue  # deliberately not used -- individual replicate fits used instead
        elif h.endswith("no-k_r fit"):
            nokr_segs.append(read_series(c))
        elif h.endswith("+k_r fit"):
            krfit_segs.append(read_series(c))
    exp_pts.sort()
    return exp_pts, nokr_segs, krfit_segs


def read_panel(sheet_name):
    ws = wb_src[sheet_name]
    rp_hdr_row = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith("RP"):
            rp_hdr_row = r
            break
    btec = read_block(ws, 6, 7, unit=1.0)
    rp = read_block(ws, rp_hdr_row + 2, rp_hdr_row + 3, unit=100.0)  # m -> cm
    return dict(btec=btec, rp=rp)


DATA = {k: read_panel(k) for k, _ in PANELS}

BTEC_X = (0, 10, 1)
BTEC_Y = (-6, 0, 1)
RP_X = (0, 20, 2)
RP_Y = (5, 10, 1)

BLOCK_W = 7   # unchanged from the original: exp(2) + k_r=0(2) + k_r fit(2) + 1 gap
ROW0 = 3
DATA_ROW = ROW0 + 2


def write_segments(ws, cx, cy, segs):
    """Write one or more replicate point-lists into columns cx/cy as successive rows, with a single
    blank row between replicates (breaks the chart's connecting line there instead of joining the
    end of one replicate's curve to the start of the next). Returns the total row span, blanks
    included, which is exactly what build_chart's n_m1/n_m2 (a contiguous Reference range) needs."""
    r = DATA_ROW
    for i, seg in enumerate(segs):
        for x, y in seg:
            ws.cell(r, cx, x); ws.cell(r, cy, y); r += 1
        if i < len(segs) - 1:
            r += 1  # blank separator row
    return r - DATA_ROW


def write_data_sheet(ws, panels):
    col = 1
    for key, corner_id in panels:
        d = DATA[key]
        for kind, (exp_pts, nokr_segs, krfit_segs) in (("BTEC", d["btec"]), ("RP", d["rp"])):
            c0 = col
            ws.cell(ROW0 - 2, c0, corner_id)
            ws.merge_cells(start_row=ROW0 - 2, start_column=c0, end_row=ROW0 - 2, end_column=c0 + 5)
            ws.cell(ROW0 - 1, c0, kind)
            ws.merge_cells(start_row=ROW0 - 1, start_column=c0, end_row=ROW0 - 1, end_column=c0 + 5)
            labels = ["Experiment", "", LEGEND[0], "", LEGEND[1], ""]
            for i, lab in enumerate(labels):
                ws.cell(ROW0, c0 + i, lab)
            xlab = "Pore Volumes" if kind == "BTEC" else "Distance (cm)"
            ylab = "Log10 C/C0" if kind == "BTEC" else "Log10 #"
            for i in range(0, 6, 2):
                ws.cell(ROW0 + 1, c0 + i, xlab)
                ws.cell(ROW0 + 1, c0 + i + 1, ylab)
            r = DATA_ROW
            for x, y in exp_pts:
                ws.cell(r, c0, x); ws.cell(r, c0 + 1, y); r += 1
            n_kr0 = write_segments(ws, c0 + 2, c0 + 3, nokr_segs)
            n_krfit = write_segments(ws, c0 + 4, c0 + 5, krfit_segs)
            DATA[key][kind.lower() + "_counts"] = (len(exp_pts), n_kr0, n_krfit)
            col += BLOCK_W


def build_workbook(panels, out_name):
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    ws_data = wb.create_sheet("Data")
    write_data_sheet(ws_data, panels)
    ws_fig = wb.create_sheet("Figure")
    n = len(panels)
    row_anchor = 0
    for i, (key, corner_id) in enumerate(panels):
        bottom = (i == n - 1)
        with_legend = (i == 1)
        col0 = 1 + i * BLOCK_W * 2
        n_exp_b, n_kr0_b, n_krfit_b = DATA[key]["btec_counts"]
        n_exp_r, n_kr0_r, n_krfit_r = DATA[key]["rp_counts"]
        ch_btec = build_chart(ws_data, "BTEC", corner_id, *BTEC_X, *BTEC_Y,
                               col0=col0, data_row=DATA_ROW, n_exp=n_exp_b, n_m1=n_kr0_b,
                               n_m2=n_krfit_b, bottom_row=bottom, with_legend=with_legend,
                               legend_labels=LEGEND)
        ch_rp = build_chart(ws_data, "RP", corner_id, *RP_X, *RP_Y,
                             col0=col0 + BLOCK_W, data_row=DATA_ROW, n_exp=n_exp_r, n_m1=n_kr0_r,
                             n_m2=n_krfit_r, bottom_row=bottom, with_legend=False,
                             legend_labels=LEGEND)
        ws_fig.add_chart(ch_btec, f"A{row_anchor+1}")
        ws_fig.add_chart(ch_rp, f"J{row_anchor+1}")
        row_anchor += 13 if not bottom else 17
    ws_fig.sheet_view.zoomScale = 60
    ws_fig.sheet_view.zoomScaleNormal = 60
    ws_fig.page_setup.orientation = "landscape"
    ws_fig.page_setup.fitToPage = True
    ws_fig.page_setup.fitToWidth = 1
    ws_fig.page_setup.fitToHeight = 0
    ws_fig.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Figure", offset=-1 * wb.sheetnames.index("Figure"))
    wb.active = 0
    out_path = f"{OUTDIR}/{out_name}"
    wb.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    build_workbook(PANELS, "Figure6_IHOP.xlsx")
