#!/usr/bin/env python3
"""Build Figure 7 as a native openpyxl Excel chart: alpha_m vs alpha_s for all 19 unfavorable
conditions (Master table of UnfavorableMaster.xlsx -- condition-averaged over replicates, matching
the parameter table's own convention), medium-coded (glass=circle, quartz=square), plus the analytic RP-branch
boundary curve a_m = a_s/(1-a_s) (Al-Zghoul et al. 2025 Eqn 15). Log-log axes, matching the existing
alpha_trends.py matplotlib version (Code/alpha_trends.py, ~line 261) that this replaces for manuscript
use -- colors/markers/axis ranges chosen to match that reference exactly (verified against the image
already embedded in the manuscript as Figure 7, extracted from JohnsonetalWRR2026-IHOP.docx).

Data source: Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx, 'Master table' sheet (already-computed
fits; no new fitting performed here). Run from Code/.
"""
import numpy as np
import openpyxl
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text as ChartText
from openpyxl.chart.title import Title
from openpyxl.chart.legend import Legend
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (RichTextProperties, Paragraph, ParagraphProperties,
                                    CharacterProperties, Font as DrawFont, RegularTextRun)

SRC = "../Manuscript/FigsExcelsUnfav/UnfavorableMaster.xlsx"   # run from Code/
OUTDIR = "../Manuscript/Figures"
APTOS = "Aptos"

GLASS_COLOR = "1F77B4"   # matches alpha_trends.py's matplotlib color exactly
QUARTZ_COLOR = "7B2D8B"


def rich_run(text, sz=1400, bold=True, italic=False):
    cp = CharacterProperties(sz=sz, b=bold, i=italic)
    cp.latin = DrawFont(typeface=APTOS)
    return RegularTextRun(rPr=cp, t=text)


def axis_title(text, x, y, italic_alpha=False):
    para = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True)),
                      r=[rich_run(text, sz=1400)])
    rt = RichText(bodyPr=RichTextProperties(), p=[para])
    t = Title(tx=ChartText(rich=rt))
    t.overlay = True
    t.layout = Layout(manualLayout=ManualLayout(xMode="edge", yMode="edge", x=x, y=y))
    return t


def read_master(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Master table"]
    rows = []
    r = 4
    while ws.cell(r, 1).value is not None:
        study = ws.cell(r, 1).value; cols = ws.cell(r, 2).value
        size_ = ws.cell(r, 4).value; IS = ws.cell(r, 5).value; v = ws.cell(r, 6).value
        rows.append(dict(medium=ws.cell(r, 3).value, a_s=ws.cell(r, 8).value, a_m=ws.cell(r, 9).value,
                          identifier=f"{study}-{cols}", identifier2=f"{v}-{size_}-{IS}"))
        r += 1
    return rows


def write_data_sheet(ws, rows):
    # column layout: glass A,B,C,D (a_s,a_m,ident,v-size-IS); quartz E,F,G,H (same); boundary I,J
    glass = [(r["a_s"], r["a_m"], r["identifier"], r["identifier2"]) for r in rows if r["medium"] == "glass"]
    quartz = [(r["a_s"], r["a_m"], r["identifier"], r["identifier2"]) for r in rows if r["medium"] == "quartz"]
    boundary_x = np.logspace(np.log10(5e-4), np.log10(0.95), 100)
    boundary_y = boundary_x / (1 - boundary_x)

    ws.cell(1, 1, "glass alpha_s"); ws.cell(1, 2, "glass alpha_m")
    ws.cell(1, 3, "glass identifier (study-column(s))"); ws.cell(1, 4, "glass v-size-IS (m/day-um-mM)")
    ws.cell(1, 5, "quartz alpha_s"); ws.cell(1, 6, "quartz alpha_m")
    ws.cell(1, 7, "quartz identifier (study-column(s))"); ws.cell(1, 8, "quartz v-size-IS (m/day-um-mM)")
    ws.cell(1, 9, "boundary alpha_s"); ws.cell(1, 10, "boundary alpha_m (=a_s/(1-a_s))")
    for i, (x, y, ident, ident2) in enumerate(glass, start=2):
        ws.cell(i, 1, x); ws.cell(i, 2, y); ws.cell(i, 3, ident); ws.cell(i, 4, ident2)
    for i, (x, y, ident, ident2) in enumerate(quartz, start=2):
        ws.cell(i, 5, x); ws.cell(i, 6, y); ws.cell(i, 7, ident); ws.cell(i, 8, ident2)
    for i, (x, y) in enumerate(zip(boundary_x, boundary_y), start=2):
        ws.cell(i, 9, float(x)); ws.cell(i, 10, float(y))
    for col in ("C", "D", "G", "H"):
        ws.column_dimensions[col].width = 30
    return len(glass), len(quartz), len(boundary_x)


def build_chart(ws, n_glass, n_quartz, n_bnd):
    ch = ScatterChart()
    ch.style = 2
    ch.scatterStyle = "marker"
    ch.varyColors = False

    def colref(c, r1, r2):
        return Reference(ws, min_col=c, min_row=r1, max_col=c, max_row=r2)

    # boundary curve first (drawn under the points), then glass, then quartz
    xref = colref(9, 2, 1 + n_bnd); yref = colref(10, 2, 1 + n_bnd)
    s_b = Series(values=yref, xvalues=xref, title="branch boundary")
    s_b.marker = Marker(symbol="none")
    s_b.graphicalProperties = GraphicalProperties()
    s_b.graphicalProperties.line = LineProperties(solidFill="000000", w=12700)
    s_b.smooth = True
    ch.series.append(s_b)

    xref = colref(1, 2, 1 + n_glass); yref = colref(2, 2, 1 + n_glass)
    s_g = Series(values=yref, xvalues=xref, title="glass")
    s_g.marker = Marker(symbol="circle", size=7)
    s_g.marker.graphicalProperties = GraphicalProperties(solidFill=GLASS_COLOR)
    s_g.marker.graphicalProperties.line = LineProperties(solidFill="000000", w=9525)
    s_g.graphicalProperties = GraphicalProperties()
    s_g.graphicalProperties.line = LineProperties(noFill=True)
    s_g.smooth = False
    ch.series.append(s_g)

    xref = colref(5, 2, 1 + n_quartz); yref = colref(6, 2, 1 + n_quartz)
    s_q = Series(values=yref, xvalues=xref, title="quartz")
    s_q.marker = Marker(symbol="square", size=7)
    s_q.marker.graphicalProperties = GraphicalProperties(solidFill=QUARTZ_COLOR)
    s_q.marker.graphicalProperties.line = LineProperties(solidFill="000000", w=9525)
    s_q.graphicalProperties = GraphicalProperties()
    s_q.graphicalProperties.line = LineProperties(noFill=True)
    s_q.smooth = False
    ch.series.append(s_q)

    ch.x_axis.scaling.min = 1e-4
    ch.x_axis.scaling.max = 2.0
    ch.x_axis.scaling.logBase = 10
    ch.x_axis.delete = False
    ch.x_axis.axPos = "b"
    ch.x_axis.crossAx = 20
    ch.x_axis.majorTickMark = "out"
    ch.x_axis.minorTickMark = "none"
    ch.x_axis.tickLblPos = "nextTo"
    ch.x_axis.numFmt = "0.0000"
    ch.x_axis.spPr = GraphicalProperties()
    ch.x_axis.spPr.line = LineProperties(solidFill="000000", w=12700)
    ch.x_axis.title = axis_title("αs", 0.46, 0.90)

    ch.y_axis.scaling.min = 1e-4
    ch.y_axis.scaling.max = 50
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.delete = False
    ch.y_axis.axPos = "l"
    ch.y_axis.crossAx = 10
    ch.y_axis.majorTickMark = "out"
    ch.y_axis.minorTickMark = "none"
    ch.y_axis.tickLblPos = "nextTo"
    ch.y_axis.numFmt = "0.0000"
    ch.y_axis.spPr = GraphicalProperties()
    ch.y_axis.spPr.line = LineProperties(solidFill="000000", w=12700)
    ch.y_axis.title = axis_title("αm", 0.01, 0.42)

    ch.x_axis.majorGridlines = None
    ch.y_axis.majorGridlines = None

    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.15, y=0.04, w=0.80, h=0.82))

    ch.legend = Legend()
    ch.legend.position = "t"
    ch.legend.overlay = False
    legpara = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1200, b=False)))
    legpara.pPr.defRPr.latin = DrawFont(typeface=APTOS)
    ch.legend.txPr = RichText(bodyPr=RichTextProperties(), p=[legpara])

    ch.graphical_properties = GraphicalProperties()
    ch.graphical_properties.line = LineProperties(noFill=True)
    ch.graphical_properties.noFill = True

    top_para = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True)))
    top_para.pPr.defRPr.latin = DrawFont(typeface=APTOS)
    ch.textProperties = RichText(bodyPr=RichTextProperties(), p=[top_para])

    ch.width = 16
    ch.height = 13
    return ch


def build_workbook(out_name):
    rows = read_master(SRC)
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    ws_data = wb.create_sheet("Data")
    n_glass, n_quartz, n_bnd = write_data_sheet(ws_data, rows)
    ws_fig = wb.create_sheet("Figure")
    ch = build_chart(ws_data, n_glass, n_quartz, n_bnd)
    ws_fig.add_chart(ch, "A1")
    ws_fig.sheet_view.zoomScale = 100
    ws_fig.sheet_view.zoomScaleNormal = 100
    ws_fig.page_setup.orientation = "landscape"
    ws_fig.page_setup.fitToPage = True
    ws_fig.page_setup.fitToWidth = 1
    ws_fig.page_setup.fitToHeight = 1
    ws_fig.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Figure", offset=-1 * wb.sheetnames.index("Figure"))
    wb.active = 0
    out_path = f"{OUTDIR}/{out_name}"
    wb.save(out_path)
    print("saved", out_path, f"({n_glass} glass, {n_quartz} quartz points)")


if __name__ == "__main__":
    build_workbook("Figure7_AlphaBoundary.xlsx")
