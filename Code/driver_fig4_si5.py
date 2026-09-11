#!/usr/bin/env python3
import openpyxl, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_fig4_si5 import build_chart, SRC, OUTDIR
from openpyxl.utils import get_column_letter

wb_src = openpyxl.load_workbook(SRC, data_only=True)

def read_panel(sheet_name):
    ws = wb_src[sheet_name]
    btec_hdr = rp_hdr = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v == "PV":
            btec_hdr = r
        if v == "Distance (m)":
            rp_hdr = r
    btec = []
    r = btec_hdr + 1
    while ws.cell(r, 1).value is not None:
        btec.append(tuple(ws.cell(r, c).value for c in range(1, 5)))
        r += 1
    rp = []
    r = rp_hdr + 1
    while r <= ws.max_row and ws.cell(r, 1).value is not None:
        rp.append(tuple(ws.cell(r, c).value for c in range(1, 5)))
        r += 1
    title1 = ws.cell(1, 1).value
    return dict(btec=btec, rp=rp, title=title1)

PANELS_SRC = {
    "GlassLi_B": "GlassLi_B", "GlassTong_AB": "GlassTong_AB", "GlassLi_H": "GlassLi_H",
    "GlassTong_CS": "GlassTong_CS", "QuartzLi_I": "QuartzLi_I",
    "GlassTong_L": "GlassTong_L", "QuartzLi_B": "QuartzLi_B", "QuartzLi_E": "QuartzLi_E",
    "QuartzTong_O": "QuartzTong_O",
}
DATA = {k: read_panel(v) for k, v in PANELS_SRC.items()}

# panel spec: (data_key, corner_id, legend_labels)
FIG4_PANELS = [
    ("GlassLi_B", "glass-2-0.98-10-Li-B", ("k_f only", "IHOP")),
    ("GlassTong_AB", "glass-4-1.0-50-Tong-AB", ("k_f only", "IHOP")),  # data plotted IS the Tong AB sheet (10-pt RP); same experiment as Li E (Records/CLAUDE.md), but corner ID follows the sheet actually read, matching every other panel's single-study convention
    ("GlassLi_H", "glass-8-0.98-10-Li-H", ("k_f only", "IHOP")),
    ("GlassTong_CS", "glass-8-2.0-50-Tong-CS", ("k_f only", "IHOP")),
    ("QuartzLi_I", "quartz-8-0.98-10-Li-I", ("k_f only", "IHOP")),
]
SI5_PANELS = [
    ("GlassTong_L", "glass-4-0.5-50-Tong-L", ("k_f only", "IHOP")),
    ("QuartzLi_B", "quartz-2-0.98-10-Li-B", ("k_f only", "IHOP")),
    ("QuartzLi_E", "quartz-4-0.98-10-Li-E", ("k_f only", "IHOP")),
    ("QuartzTong_O", "quartz-8-0.5-50-Tong-O", ("k_f only", "IHOP")),
]

BTEC_X = (0, 10, 1)
BTEC_Y = (-12, 0, 2)
RP_X = (0, 20, 2)
RP_Y = (5, 10, 1)

BLOCK_W = 7  # 6 data cols + 1 gap
ROW0 = 3     # header/label row start
DATA_ROW = ROW0 + 2  # =5, same convention as Figure 1


def write_data_sheet(ws, panels):
    col = 1
    for key, corner_id, leglabels in panels:
        d = DATA[key]
        for kind, arr, ymin, ymax in (("BTEC", d["btec"], *BTEC_Y[:2]), ("RP", d["rp"], *RP_Y[:2])):
            c0 = col
            # row1 merged panel id, row2 merged kind, row3 series labels, row4 axis labels
            ws.cell(ROW0 - 2, c0, corner_id)
            ws.merge_cells(start_row=ROW0 - 2, start_column=c0, end_row=ROW0 - 2, end_column=c0 + 5)
            ws.cell(ROW0 - 1, c0, kind)
            ws.merge_cells(start_row=ROW0 - 1, start_column=c0, end_row=ROW0 - 1, end_column=c0 + 5)
            labels = ["Experiment", "", leglabels[0], "", leglabels[1], ""]
            for i, lab in enumerate(labels):
                ws.cell(ROW0, c0 + i, lab)
            xlab = "Pore Volumes" if kind == "BTEC" else "Distance (cm)"
            ylab = "Log10 C/C0" if kind == "BTEC" else "Log10 #"
            for i in range(0, 6, 2):
                ws.cell(ROW0 + 1, c0 + i, xlab)
                ws.cell(ROW0 + 1, c0 + i + 1, ylab)
            # data rows: col0/1 = exp, col2/3=model1, col4/5=model2
            r = DATA_ROW
            unit = 100.0 if kind == "RP" else 1.0  # m -> cm for RP
            for x, y, m1, m2 in arr:
                ws.cell(r, c0, x * unit)
                ws.cell(r, c0 + 1, y)
                ws.cell(r, c0 + 2, x * unit)
                ws.cell(r, c0 + 3, m1)
                ws.cell(r, c0 + 4, x * unit)
                ws.cell(r, c0 + 5, m2)
                r += 1
            col += BLOCK_W


def build_workbook(panels, out_name):
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    ws_data = wb.create_sheet("Data")
    write_data_sheet(ws_data, panels)
    ws_fig = wb.create_sheet("Figure")
    n = len(panels)
    row_anchor = 0
    for i, (key, corner_id, leglabels) in enumerate(panels):
        bottom = (i == n - 1)
        with_legend = (i == 1) if n > 1 else (i == 0)
        col0 = 1 + i * BLOCK_W * 2
        n_exp_btec = len(DATA[key]["btec"])
        n_exp_rp = len(DATA[key]["rp"])
        # model curves same length as measured here (they were computed at measured PV/depth points,
        # not on a dense grid, per Favorable_data_and_sim.xlsx's own layout)
        ch_btec = build_chart(ws_data, "BTEC", corner_id, *BTEC_X, *BTEC_Y,
                               col0=col0, data_row=DATA_ROW, n_exp=n_exp_btec, n_m1=n_exp_btec,
                               n_m2=n_exp_btec, bottom_row=bottom, with_legend=with_legend,
                               legend_labels=leglabels)
        ch_rp = build_chart(ws_data, "RP", corner_id, *RP_X, *RP_Y,
                             col0=col0 + BLOCK_W, data_row=DATA_ROW, n_exp=n_exp_rp, n_m1=n_exp_rp,
                             n_m2=n_exp_rp, bottom_row=bottom, with_legend=False,
                             legend_labels=leglabels)
        anchor_a = f"A{row_anchor+1}"
        anchor_b = f"J{row_anchor+1}"
        ws_fig.add_chart(ch_btec, anchor_a)
        ws_fig.add_chart(ch_rp, anchor_b)
        row_anchor += 13 if not bottom else 17
    ws_fig.sheet_view.zoomScale = 60
    ws_fig.sheet_view.zoomScaleNormal = 60
    ws_fig.page_setup.orientation = "landscape"
    ws_fig.page_setup.fitToPage = True
    ws_fig.page_setup.fitToWidth = 1
    ws_fig.page_setup.fitToHeight = 0
    ws_fig.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Figure", offset=-1 * wb.sheetnames.index("Figure"))
    wb.active = 0
    out_path = f"{OUTDIR}/{out_name}"
    wb.save(out_path)
    print("saved", out_path)


build_workbook(FIG4_PANELS, "Figure4_Favorable.xlsx")
build_workbook(SI5_PANELS, "FigureSI5_Favorable.xlsx")
