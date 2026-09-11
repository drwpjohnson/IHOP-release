"""
gl_is_free_test.py -- prior-fairness / corroboration test for alpha_s on the glass ionic-strength series.
alpha_s was originally PINNED to the Happel single-interception fraction, but that Happel value came from a
clean-bed simulation whose SCOV was tuned to the breakthrough plateau -- so the pin reuses the column's data.
This test frees alpha_s while holding k_f at the INDEPENDENT favorable clean-bed r_s (=22.9, from separate
alpha~1 experiments), so alpha_s is set by the retention-profile LEVEL (a different observable than the
plateau). If free alpha_s ~ Happel alpha_s, the pin was fair (internal consistency, not circularity).

RESULT: free alpha_s recovers Happel within 2% (6 mM), 2% (20 mM), 14% (50 mM) at equal-or-lower cost;
f_x essentially unchanged, alpha_m same branch. -> pin not load-bearing; glass IS now reported free-alpha_s
(gl_is_fx.py, canonical). Logged in identifiability_and_necessity_tests.md (iii) and fx_trend.xlsx
('alpha_s pin-vs-free'). Requires DataFromLi&Tong.xlsx under ../Data/.  Usage: python3 gl_is_free_test.py
"""
import numpy as np, openpyxl, re, warnings, functools, os
warnings.filterwarnings("ignore"); print = functools.partial(print, flush=True)
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
from openpyxl.utils import column_index_from_string as cix
_HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(_HERE, "..", "Data", "DataFromLi&Tong.xlsx")
DAY = 86400.; L = 0.2; theta = 0.375
REV, T0, Vref = 22.801836559387397, 3.58, 0.1667; INJPV = T0 / (L / Vref)
W_RP, W_SHAPE, W_PLAT, W_TAIL, W_TSLOPE = 6.0, 2.5, 10.0, 3.0, 16.0
isp = lambda v: any(abs(v - p) < 0.05 for p in {-6., -9.94, -15.33})
wb = openpyxl.load_workbook(DATA, data_only=True); ws = wb["Microspheres Glass Beads Li"]


def c0(dc):
    m = re.search(r'([\d.]+)\s*E\s*([+-]?\d+)', str(ws.cell(14, dc).value)); return float(m.group(1)) * 10 ** int(m.group(2)) if m else None
def load(dc):
    bt = [(ws.cell(r, dc).value, ws.cell(r, dc + 1).value) for r in range(16, 44)]
    bt = [(a, b) for a, b in bt if isinstance(a, (int, float)) and isinstance(b, (int, float)) and b < 0.5 and not isp(b)]
    rp = [(ws.cell(r, dc).value, ws.cell(r, dc + 1).value) for r in range(47, 57)]
    rp = [(a, b) for a, b in rp if isinstance(a, (int, float)) and isinstance(b, (int, float)) and b > 0.5 and not isp(b)]
    # ALL RP data has exactly 10 depth points, always (W.P.J., 2026-09-01) -- assert, don't silently truncate.
    assert len(rp) == 10, f"RP truncated: col {dc} got {len(rp)} points, expected 10"
    return np.array(bt), np.array(rp), c0(dc)


class Eng:
    def __init__(s, v): s.v = v / DAY; s.dx = L / 90; s.dt = s.dx / s.v; s.pv = L / s.v
    def run(s, k1, a_s, a_m, fx, kr, D, vns=0.05):
        kf = k1; k2 = fx * kf; kmw = a_m * kf; kmg = a_m * vns * kf; G = np.zeros((4, 4))
        G[0, 0] -= kf; G[3, 0] += a_s * kf; G[1, 0] += (1 - a_s) * kf; G[1, 1] -= (kmw + k2); G[3, 1] += kmw; G[2, 1] += k2; G[2, 2] -= kmg; G[3, 2] += kmg; G[1, 2] += kr; G[2, 2] -= kr
        s._k1 = k1; s._kmw = kmw; s._kmg = kmg; E = expm(G * s.dt); r = D * s.dt / s.dx**2; mn = (1 + 2 * r) * np.ones(90); of = -r * np.ones(89); mn[0] = 1 + r; mn[-1] = 1 + r
        lu = splu(csc_matrix(diags([of, mn, of], [-1, 0, 1], format="csc")))
        nt = 900; ti = INJPV * s.pv; c = vns; Y = np.zeros((4, 90)); C = np.zeros(nt); t = 0.; Yst = None
        for i in range(nt):
            Y = E @ Y; C[i] = (Y[0, -1] + Y[1, -1] + c * Y[2, -1]); Y[0, 1:] = Y[0, :-1]; Y[0, 0] = 0; Y[1, 1:] = Y[1, :-1]; Y[1, 0] = 0
            ym = Y[2].copy(); Y[2, 1:] = ym[1:] - c * (ym[1:] - ym[:-1]); Y[2, 0] = ym[0] * (1 - c); t += s.dt
            if t < ti: Y[0, 0] += 1.
            Y[0, :] = lu.solve(Y[0, :]); Y[1, :] = lu.solve(Y[1, :])
            if t < ti and t > 0.9 * ti: Yst = Y.copy()
        tp = (np.arange(nt) + 1) * s.dt / s.pv; x = (np.arange(90) + 0.5) * s.dx
        rp = (a_s * s._k1 * Yst[0, :] + s._kmw * Yst[1, :] + s._kmg * Yst[2, :]); m = (tp > 0.5 * INJPV) & (tp < INJPV)
        return dict(tp=tp, C=C, x=x, rp=rp, plat=float(C[m].mean()))


def sl(x, y): h = len(x) // 2; return float(np.polyfit(x[:h + 1], y[:h + 1], 1)[0]), float(np.polyfit(x[h:], y[h:], 1)[0])
def ts(pv, lc, lo=6, hi=10):
    m = (np.asarray(pv) >= lo) & (np.asarray(pv) <= hi); return float("nan") if np.sum(m) < 2 else float(np.polyfit(np.asarray(pv)[m], np.asarray(lc)[m], 1)[0])


def builder(dc, rs, vel):
    bt, rp, C0 = load(dc); col = Eng(vel); V_MS = vel / DAY; k1 = rs * V_MS; logK = np.log10(REV * T0 * theta * C0 * Vref); D = col.v * L / 150
    x = rp[:, 0]; lrp = rp[:, 1]; pv = bt[:, 0]; lc = bt[:, 1]; si, so = sl(x, lrp)
    pm = (pv > 1.3) & (pv < 3.5); pp = lc[pm]; plo, phi = float(pp.min()), float(pp.max()); tm = pv > 4.; sld = ts(pv, lc)
    def rr(a_s, a_m, fx, kr):
        r = col.run(k1, a_s, a_m, fx, kr, D)
        rl = np.interp(x, r['x'], np.maximum(r['rp'], 1e-300)) / V_MS; lS = logK + np.log10(np.maximum(rl, 1e-300))
        rpx = W_RP * (lS - lrp); sm, sn = sl(x, lS); shp = W_SHAPE * np.array([sm - si, sn - so]); p_ = np.log10(max(r['plat'], 1e-12)); ex = max(0., p_ - phi) + max(0., plo - p_); pl = W_PLAT * np.array([ex])
        tl = W_TAIL * (np.log10(np.maximum(np.interp(pv[tm], r['tp'], r['C']), 1e-12)) - lc[tm]); slm = ts(r['tp'], np.log10(np.maximum(r['C'], 1e-12))); tsl = W_TSLOPE * np.array([slm - sld])
        return np.concatenate([rpx, shp, pl, tl, tsl])
    return rr, col


def fit_free(dc, rs=22.9, vel=4.):
    rr, col = builder(dc, rs, vel)
    f = lambda p: rr(*[10 ** q for q in p])
    lo = [-3.3, -3.3, np.log10(1e-4), np.log10(1e-6)]; hi = [np.log10(.95), np.log10(.999), np.log10(1.0), np.log10(.1)]
    best = None
    for s0 in [[-1.5, -1., np.log10(.005), np.log10(1e-4)], [-.5, -.3, np.log10(.02), np.log10(3e-4)], [-2.3, -1.5, np.log10(.001), np.log10(1e-4)]]:
        r = least_squares(f, s0, bounds=(lo, hi), max_nfev=100)
        if best is None or r.cost < best.cost: best = r
    a_s, a_m, fx, kr = [10 ** q for q in best.x]; return dict(a_s=a_s, a_m=a_m, fx=fx, kr=kr * (L / col.v), cost=float(best.cost))


def fit_pinned(dc, a_s, rs=22.9, vel=4.):
    rr, col = builder(dc, rs, vel)
    f = lambda p: rr(a_s, *[10 ** q for q in p])
    lo = [-3.3, np.log10(1e-4), np.log10(1e-6)]; hi = [np.log10(.999), np.log10(1.0), np.log10(.1)]
    best = None
    for s0 in [[-1., np.log10(.005), np.log10(1e-4)], [-.3, np.log10(.02), np.log10(3e-4)]]:
        r = least_squares(f, s0, bounds=(lo, hi), max_nfev=100)
        if best is None or r.cost < best.cost: best = r
    a_m, fx, kr = [10 ** q for q in best.x]; return dict(a_s=a_s, a_m=a_m, fx=fx, kr=kr * (L / col.v), cost=float(best.cost))


if __name__ == "__main__":
    COLS = [("6 mM", "R", 0.0171), ("20 mM", "O", 0.1327), ("50 mM", "L", 0.4001)]  # a_s Happel single-fraction
    print(f"{'IS':>6} | {'a_s pin':>8}{'a_s FREE':>9}{'ratio':>7} | {'f_x pin':>8}{'f_x free':>9} | {'a_m pin':>8}{'a_m free':>9} | {'cost pin':>9}{'cost free':>10}")
    for IS, letter, asp in COLS:
        dc = cix(letter); fp = fit_pinned(dc, asp); ff = fit_free(dc)
        print(f"{IS:>6} | {asp:>8.4f}{ff['a_s']:>9.4f}{ff['a_s'] / asp:>7.2f} | {fp['fx']:>8.4f}{ff['fx']:>9.4f} | {fp['a_m']:>8.4f}{ff['a_m']:>9.4f} | {fp['cost']:>9.2f}{ff['cost']:>10.2f}")
