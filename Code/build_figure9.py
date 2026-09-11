#!/usr/bin/env python3
"""Build Figure 9 as a native openpyxl Excel chart: fitted k_r vs colloid size, at fixed IS=20 mM
(the 18 unfavorable conditions the caption cites), glass split by velocity (4 vs 8 m/day, circles,
two colors) and quartz (square), plus the DLVO/Kramers single-colloid escape prediction (dotted
grey curve, REFUTED -- the data are flat where the prediction sweeps ~584x) and per-medium
horizontal reference lines at the "clean" geometric-mean k_r (excluding each medium's own
fit-floor-hit column).

Data source: Manuscript/FigsExcelsUnfav/k_r_trend.xlsx -- 'k_r fitted (new)' sheet filtered to
IS_mM==20 (18 rows: 14 glass, 4 quartz -- matches the caption's "18 unfavorable conditions"
exactly), 'reference constants' sheet for the clean glass/quartz geomeans, and 'DLVO-Kramers
curve' sheet for the anchored prediction curve (40 points).

** FLAG, NOT SILENTLY FIXED: the embedded Figure 9 image currently in the manuscript
(word/media/image15.png, rId25) is STALE relative to this data, and stale in more than one place.
Its dashed reference lines read "glass k_r ~ 4.5e-05/s" and "quartz k_r ~ 2.1e-05/s" and its title
reads "necessity +4.3% (glass)" -- these are the PRE-2026-08-25 accumulated-solid-phase-convention
numbers (see kr_trend.py's own docstring audit trail). The CURRENT clean geomeans, from today's
k_r_trend.xlsx, are glass 3.75e-05/s and quartz 1.82e-05/s (ratio 2.06x), and current necessity is
+4.8% glass / +0.9% quartz. The 584x DLVO/Kramers size-swing number in the old title, by contrast,
IS still current (this script reproduces 584x from today's data too) -- so not everything in the
old image is wrong, just the two k_r reference constants and the necessity percentage. The biggest
visible consequence: under the OLD convention Li.M's k_r was railed at the 1e-6 floor and plotted
near the bottom of the old figure; under the CURRENT convention it is 5.63e-3 (the same degeneracy
already flagged for Figures 6 and 8) and appears near the TOP of this new chart instead --
a large, real, already-documented shift, not a build error.

Caption note, not acted on here: the caption text says the dashed lines are "per-medium median
kr values" -- the script and workbook both call them GEOMETRIC MEANS ("geomean"/"clean geomean"),
not medians. Flagging the wording difference; not changed since it's manuscript text, not data.

Run from Code/.
"""
import openpyxl
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text as ChartText
from openpyxl.chart.title import Title
from openpyxl.chart.legend import Legend
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (RichTextProperties, Paragraph, ParagraphProperties,
                                    CharacterProperties, Font as DrawFont, RegularTextRun)

SRC = "../Manuscript/FigsExcelsUnfav/k_r_trend.xlsx"   # run from Code/
OUTDIR = "../Manuscript/Figures"
APTOS = "Aptos"

GLASS4_COLOR = "1F77B4"   # matches kr_trend.py's own matplotlib colors exactly
GLASS8_COLOR = "2A9D8F"
QUARTZ_COLOR = "7B2D8B"
GREY = "888888"


def rich_run(text, sz=1400, bold=True, italic=False):
    cp = CharacterProperties(sz=sz, b=bold, i=italic)
    cp.latin = DrawFont(typeface=APTOS)
    return RegularTextRun(rPr=cp, t=text)


def axis_title(text, x, y):
    para = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True)),
                      r=[rich_run(text, sz=1400)])
    rt = RichText(bodyPr=RichTextProperties(), p=[para])
    t = Title(tx=ChartText(rich=rt))
    t.overlay = True
    t.layout = Layout(manualLayout=ManualLayout(xMode="edge", yMode="edge", x=x, y=y))
    return t


def read_source(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["k_r fitted (new)"]
    hdr = [ws.cell(4, c).value for c in range(1, 10)]
    rows = []
    r = 5
    while ws.cell(r, 1).value is not None:
        rows.append({h: ws.cell(r, c).value for c, h in enumerate(hdr, 1)})
        r += 1
    rows = [r for r in rows if r["IS_mM"] == 20]

    wref = wb["reference constants"]
    fg = wref.cell(6, 2).value   # glass geomean (clean)
    fq = wref.cell(8, 2).value   # quartz geomean (clean)

    wc = wb["DLVO-Kramers curve"]
    curve = []
    r = 5
    while wc.cell(r, 1).value is not None:
        curve.append((wc.cell(r, 1).value, wc.cell(r, 4).value))
        r += 1
    return rows, fg, fq, curve


def write_data_sheet(ws, rows, fg, fq, curve):
    g4 = [(r["size_um"], r["k_r_/s"], r["col"]) for r in rows if r["medium"] == "glass" and r["v_mday"] == 4]
    g8 = [(r["size_um"], r["k_r_/s"], r["col"]) for r in rows if r["medium"] == "glass" and r["v_mday"] == 8]
    qz = [(r["size_um"], r["k_r_/s"], r["col"]) for r in rows if r["medium"] == "quartz"]

    blocks = [(1, "glass 4 m/day", g4), (4, "glass 8 m/day", g8), (7, "quartz", qz)]
    for c0, label, pts in blocks:
        ws.cell(1, c0, f"{label}: size (um)"); ws.cell(1, c0 + 1, "k_r (/s)"); ws.cell(1, c0 + 2, "identifier (study.col)")
        for i, (x, y, ident) in enumerate(pts, start=2):
            ws.cell(i, c0, x); ws.cell(i, c0 + 1, y); ws.cell(i, c0 + 2, ident)

    ws.cell(1, 10, "DLVO/Kramers diam (um)"); ws.cell(1, 11, "DLVO/Kramers k_r anchored (/s)")
    for i, (x, y) in enumerate(curve, start=2):
        ws.cell(i, 10, x); ws.cell(i, 11, y)

    xlo, xhi = 0.08, 2.4
    ws.cell(1, 12, "glass ref x"); ws.cell(1, 13, "glass ref k_r (clean geomean)")
    ws.cell(2, 12, xlo); ws.cell(2, 13, fg); ws.cell(3, 12, xhi); ws.cell(3, 13, fg)
    ws.cell(1, 14, "quartz ref x"); ws.cell(1, 15, "quartz ref k_r (clean geomean)")
    ws.cell(2, 14, xlo); ws.cell(2, 15, fq); ws.cell(3, 14, xhi); ws.cell(3, 15, fq)

    ws.cell(1, 17, f"glass clean geomean = {fg:.3e} /s ; quartz clean geomean = {fq:.3e} /s ; ratio = {fg/fq:.2f}x")

    for col in ("C", "F", "I"):
        ws.column_dimensions[col].width = 24
    return dict(g4=len(g4), g8=len(g8), qz=len(qz), curve=len(curve))


def build_chart(ws, n):
    ch = ScatterChart()
    ch.style = 2
    ch.scatterStyle = "marker"
    ch.varyColors = False

    def colref(c, r1, r2):
        return Reference(ws, min_col=c, min_row=r1, max_col=c, max_row=r2)

    def add_line(xcol, ycol, count, title, color, dash="solid", w=19050):
        if count == 0: return
        xref = colref(xcol, 2, 1 + count); yref = colref(ycol, 2, 1 + count)
        s = Series(values=yref, xvalues=xref, title=title)
        s.marker = Marker(symbol="none")
        s.graphicalProperties = GraphicalProperties()
        s.graphicalProperties.line = LineProperties(solidFill=color, w=w, prstDash=dash)
        s.smooth = False
        ch.series.append(s)

    def add_points(c0, count, title, symbol, size, fill_color):
        if count == 0: return
        xref = colref(c0, 2, 1 + count); yref = colref(c0 + 1, 2, 1 + count)
        s = Series(values=yref, xvalues=xref, title=title)
        s.marker = Marker(symbol=symbol, size=size)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=fill_color)
        s.marker.graphicalProperties.line = LineProperties(solidFill="000000", w=9525)
        s.graphicalProperties = GraphicalProperties()
        s.graphicalProperties.line = LineProperties(noFill=True)
        s.smooth = False
        ch.series.append(s)

    # curve and reference lines first (drawn under the points)
    add_line(10, 11, n["curve"], "DLVO/Kramers (584x) -- REFUTED", GREY, dash="sysDot", w=12700)
    add_line(12, 13, 2, "glass k_r (clean geomean)", GLASS4_COLOR, dash="dash", w=15875)
    add_line(14, 15, 2, "quartz k_r (clean geomean)", QUARTZ_COLOR, dash="dash", w=15875)

    add_points(1, n["g4"], "glass, 4 m/day", "circle", 8, GLASS4_COLOR)
    add_points(4, n["g8"], "glass, 8 m/day", "circle", 8, GLASS8_COLOR)
    add_points(7, n["qz"], "quartz", "square", 8, QUARTZ_COLOR)

    ch.x_axis.scaling.min = 0.08
    ch.x_axis.scaling.max = 2.4
    ch.x_axis.scaling.logBase = 10
    ch.x_axis.delete = False
    ch.x_axis.axPos = "b"
    ch.x_axis.crossAx = 20
    ch.x_axis.majorTickMark = "out"
    ch.x_axis.minorTickMark = "none"
    ch.x_axis.tickLblPos = "nextTo"
    ch.x_axis.numFmt = "0.0"
    ch.x_axis.spPr = GraphicalProperties()
    ch.x_axis.spPr.line = LineProperties(solidFill="000000", w=12700)
    ch.x_axis.title = axis_title("colloid diameter (um)", 0.36, 0.90)

    ch.y_axis.scaling.min = 5e-7
    ch.y_axis.scaling.max = 1e-2
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.delete = False
    ch.y_axis.axPos = "l"
    ch.y_axis.crossAx = 10
    ch.y_axis.majorTickMark = "out"
    ch.y_axis.minorTickMark = "none"
    ch.y_axis.tickLblPos = "nextTo"
    ch.y_axis.numFmt = "0.0000000"
    ch.y_axis.spPr = GraphicalProperties()
    ch.y_axis.spPr.line = LineProperties(solidFill="000000", w=12700)
    ch.y_axis.title = axis_title("fitted k_r (/s)", 0.01, 0.35)

    ch.x_axis.majorGridlines = None
    ch.y_axis.majorGridlines = None

    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.14, y=0.04, w=0.66, h=0.90))

    ch.legend = Legend()
    ch.legend.position = "r"
    ch.legend.overlay = False
    ch.legend.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.81, y=0.10, w=0.19, h=0.75))
    legpara = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1000, b=False)))
    legpara.pPr.defRPr.latin = DrawFont(typeface=APTOS)
    ch.legend.txPr = RichText(bodyPr=RichTextProperties(), p=[legpara])

    ch.graphical_properties = GraphicalProperties()
    ch.graphical_properties.line = LineProperties(noFill=True)
    ch.graphical_properties.noFill = True

    top_para = Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True)))
    top_para.pPr.defRPr.latin = DrawFont(typeface=APTOS)
    ch.textProperties = RichText(bodyPr=RichTextProperties(), p=[top_para])

    ch.width = 19
    ch.height = 13
    return ch


def build_workbook(out_name):
    rows, fg, fq, curve = read_source(SRC)
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    ws_data = wb.create_sheet("Data")
    n = write_data_sheet(ws_data, rows, fg, fq, curve)
    ws_fig = wb.create_sheet("Figure")
    ch = build_chart(ws_data, n)
    ws_fig.add_chart(ch, "A1")
    ws_fig.sheet_view.zoomScale = 100
    ws_fig.sheet_view.zoomScaleNormal = 100
    ws_fig.page_setup.orientation = "landscape"
    ws_fig.page_setup.fitToPage = True
    ws_fig.page_setup.fitToWidth = 1
    ws_fig.page_setup.fitToHeight = 1
    ws_fig.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Figure", offset=-1 * wb.sheetnames.index("Figure"))
    wb.active = 0
    out_path = f"{OUTDIR}/{out_name}"
    wb.save(out_path)
    print("saved", out_path, n, f"glass_clean={fg:.3e} quartz_clean={fq:.3e}")


if __name__ == "__main__":
    build_workbook("Figure9_KrVsSize.xlsx")
