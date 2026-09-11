"""favorable_tail_clean.py -- clean test: k_r PINNED to the per-medium constant from the size series
(glass 5.58e-5/s, quartz 1.87e-5/s), v_ns=5%. k_r is the ONLY pinned kinetic parameter -- pinning it breaks
the f_x<->k_r near-surface degeneracy, so everything else (r_s, alpha_s, alpha_m, f_x) is fit. Does the FLAT
sustained shelf emerge under that single pin?

Figure: TWO panels per column -- LEFT = BTEC; RIGHT = RETENTION PROFILE (RP). BOTH are in the objective
(RP level W_RP=6 + two-segment slope W_SH=2.5; BTEC W_BT=3). FREE: delivery r_s (within +/-0.6 dex of the
plateau seed, so the RP slope can set delivery), single-interception alpha_s (via the MI fraction 1-alpha_s),
crawl attachment alpha_m, and recruitment f_x. PINNED: k_r (per-medium), v_ns=5%. Question: with the RP
constrained, does the sustained multi-PV tail still emerge?
RP amplitude S -> retained spheres via logK = log10(REV*T0*theta*Vref*C0); theta per medium (glass 0.375,
quartz 0.36). Writes favorable_tail_clean.png."""
import numpy as np, openpyxl, re
from openpyxl.utils import column_index_from_string as cix
from scipy.linalg import expm
from scipy.sparse import diags, csc_matrix
from scipy.sparse.linalg import splu
from scipy.optimize import least_squares
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
DATA="../Data/DataFromLi&Tong.xlsx"; DAY=86400.; L=0.2
REV,T0,Vref=22.801836559387397,3.58,0.1667; INJPV=T0/(L/Vref)
KR_MED={"glass":5.58e-5,"quartz":1.87e-5}   # per-second, from Serial-3 size series / quartz check; PINNED and superseded by 3.75e-5 / 1.82e-5 -- see favorable_both_models.py for why it is not updated
THETA_MED={"glass":0.375,"quartz":0.36}
PLACE={-6.,-9.94,-15.33}; isp=lambda v:any(abs(v-p)<0.05 for p in PLACE)
wb=openpyxl.load_workbook(DATA,data_only=True)
def c0cell(ws,c,r):
    m=re.search(r'([\d.]+)\s*E\s*([+-]?\d+)',str(ws.cell(r,c).value));return float(m.group(1))*10**int(m.group(2)) if m else None
def c0hdr(ws,c,br):
    H=" || ".join(str(ws.cell(r,c).value) for r in range(1,br) if ws.cell(r,c).value is not None)
    m=re.search(r'([\d.]+)\s*E\s*([+-]?\d+)',H);return float(m.group(1))*10**int(m.group(2)) if m else None
# RP_ROWS: per-sheet row range -- see favorable_both_models.py for the full explanation of why a single
# hardcoded range for both Li sheets is wrong (fixed 2026-09-01, W.P.J.). This file's COLS never actually
# points at Glass Beads Li today, so the bug was latent here, not live -- but the assert stays in place
# so it can't come back silently if that ever changes. ALL RP data has exactly 10 depth points, always.
RP_ROWS = {"Microspheres Glass Beads Li": (47, 57), "Microspheres Quartz Sand Li": (41, 51)}


def load_li(sh,dc):
    ws=wb[sh]
    bt=[(ws.cell(r,dc).value,ws.cell(r,dc+1).value) for r in range(10,45)];bt=[(a,b) for a,b in bt if isinstance(a,(int,float)) and isinstance(b,(int,float)) and b<0.5 and not isp(b)]
    r0,r1=RP_ROWS[sh]
    rp=[(ws.cell(r,dc).value,ws.cell(r,dc+1).value) for r in range(r0,r1)];rp=[(a,b) for a,b in rp if isinstance(a,(int,float)) and isinstance(b,(int,float)) and b>0.5 and not isp(b)]
    assert len(rp)==10, f"RP truncated: {sh} col {dc} got {len(rp)} points, expected 10 -- row range wrong for this sheet"
    return np.array(bt),np.array(rp),c0cell(ws,dc,8)
def load_tong(sh,c):
    ws=wb[sh];br=None
    for r in range(6,13):
        x=ws.cell(r,c).value
        if isinstance(x,str) and x.strip().upper().startswith("PV"):br=r;break
    pv=[];lc=[]
    for r in range(br+1,67):
        a=ws.cell(r,c).value;b=ws.cell(r,c+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isp(b) and b<0.5: pv.append(float(a));lc.append(float(b))
    xx=[];rp=[]
    for r in range(68,78):
        a=ws.cell(r,c).value;b=ws.cell(r,c+1).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and not isp(b): xx.append(float(a));rp.append(float(b))
    # ALL RP data has exactly 10 depth points, always (W.P.J., 2026-09-01) -- assert, don't silently truncate.
    if rp: assert len(rp)==10, f"RP truncated: {sh} col {c} got {len(rp)} points, expected 10"
    return np.array(list(zip(pv,lc))),np.array(list(zip(xx,rp))),c0hdr(ws,c,br)
def rs_plateau(bt):
    pv=bt[:,0];lc=bt[:,1];m=(pv>1.2)&(pv<3.4)
    if m.sum()<2:m=(pv>0.9)&(pv<3.6)
    return -np.log(10)*float(np.median(lc[m]))/L
def sl(x,y):h=len(x)//2;return float(np.polyfit(x[:h+1],y[:h+1],1)[0]),float(np.polyfit(x[h:],y[h:],1)[0])
class Eng:
    def __init__(s,vmday):s.v=vmday/DAY;s.dx=L/90;s.dt=s.dx/s.v;s.pv=L/s.v
    def run(s,k1,a_s,a_m,fx,kr,vns=0.05,tot=10):
        kf=k1;k2=fx*kf;kmw=a_m*kf;kmg=a_m*vns*kf;G=np.zeros((4,4))
        G[0,0]-=kf;G[3,0]+=a_s*kf;G[1,0]+=(1-a_s)*kf
        G[1,1]-=(kmw+k2);G[3,1]+=kmw;G[2,1]+=k2;G[2,2]-=kmg;G[3,2]+=kmg;G[1,2]+=kr;G[2,2]-=kr
        s._k1=k1;s._kmw=kmw;s._kmg=kmg;E=expm(G*s.dt);D=s.v*L/150;r=D*s.dt/s.dx**2;mn=(1+2*r)*np.ones(90);of=-r*np.ones(89);mn[0]=1+r;mn[-1]=1+r
        lu=splu(csc_matrix(diags([of,mn,of],[-1,0,1],format="csc")))
        nt=int(round(tot*90));ti=INJPV*s.pv;c=vns;Y=np.zeros((4,90));C=np.zeros(nt);t=0.;Yst=None
        for i in range(nt):
            Y=E@Y;C[i]=Y[0,-1]+Y[1,-1]+c*Y[2,-1];Y[0,1:]=Y[0,:-1];Y[0,0]=0;Y[1,1:]=Y[1,:-1];Y[1,0]=0
            ym=Y[2].copy();Y[2,1:]=ym[1:]-c*(ym[1:]-ym[:-1]);Y[2,0]=ym[0]*(1-c);t+=s.dt
            if t<ti:Y[0,0]+=1.
            Y[0,:]=lu.solve(Y[0,:]);Y[1,:]=lu.solve(Y[1,:])
            if t<ti and t>0.9*ti:Yst=Y.copy()
        tp=(np.arange(nt)+1)*s.dt/s.pv;x=(np.arange(90)+0.5)*s.dx
        rp=(a_s*s._k1*Yst[0,:]+s._kmw*Yst[1,:]+s._kmg*Yst[2,:])
        return dict(tp=tp,C=np.maximum(C,1e-300),x=x,rp=rp)
COLS=[("Glass Tong","L","tong","glass",0.5,4),("Glass Tong","AB","tong","glass",1.0,4),
      ("Quartz Li","B","li","quartz",0.98,2),("Quartz Li","E","li","quartz",0.98,4),("Quartz Li","I","li","quartz",0.98,8)]
SHEET={"Glass Tong":"Microspheres Glass Beads Tong","Quartz Li":"Microspheres Quartz Sand Li"}
fig,ax=plt.subplots(len(COLS),2,figsize=(11,2.7*len(COLS)))
print(f"{'cond':16}{'r_s':>6}{'k_r/PV':>9}{'MI%':>7}{'a_m':>7}{'f_x':>8}{'tailRMS':>9}{'rpRMS':>8}")
for i,(tab,letter,layout,med,size,vel) in enumerate(COLS):
    if layout=="li": bt,rp,C0=load_li(SHEET[tab],cix(letter))
    else:            bt,rp,C0=load_tong(SHEET[tab],cix(letter))
    rs=rs_plateau(bt); V_MS=vel/DAY; k1=rs*V_MS; eng=Eng(vel); kr=KR_MED[med]; theta=THETA_MED[med]
    logK=np.log10(REV*T0*theta*Vref*C0) if C0 else 0.0
    pv=bt[:,0]; lc=bt[:,1]; tailm=pv>4.2; rx=rp[:,0]; rlog=rp[:,1]
    si,so=sl(rx,rlog); W_RP,W_SH,W_BT=6.0,2.5,3.0; k1s=np.log10(k1)  # RP(level+slope)+BTEC in objective; r_s free +/-0.6 dex
    r0=eng.run(k1,1.0,0.0,0.0,0.0)
    def resid(p):
        kk=10**p[0]; a_s=1-10**p[1]; a_m=10**p[2]; fx=10**p[3]
        r=eng.run(kk,a_s,a_m,fx,kr)
        mC=np.log10(np.maximum(np.interp(pv,r['tp'],r['C']),1e-12))
        lS=logK+np.log10(np.maximum(np.interp(rx,r['x'],np.maximum(r['rp'],1e-300))/V_MS,1e-300))
        sm,sn=sl(rx,lS)
        return np.concatenate([W_RP*(lS-rlog), W_SH*np.array([sm-si,sn-so]), W_BT*(mC-lc)])
    best=None
    for s0 in [[k1s,-1.5,-1.5,np.log10(.005)],[k1s,-2.0,-2.0,np.log10(.02)]]:
        rr=least_squares(resid,s0,bounds=([k1s-0.6,-3,-3.3,np.log10(1e-4)],[k1s+0.6,np.log10(0.5),np.log10(0.9),np.log10(0.1)]),max_nfev=60)
        if best is None or rr.cost<best.cost: best=rr
    k1f=10**best.x[0]; a_s=1-10**best.x[1]; a_m=10**best.x[2]; fx=10**best.x[3]; rs=k1f/V_MS
    r1=eng.run(k1f,a_s,a_m,fx,kr); m1=np.log10(np.maximum(np.interp(pv,r1['tp'],r1['C']),1e-12))
    lS=logK+np.log10(np.maximum(np.interp(rx,r1['x'],np.maximum(r1['rp'],1e-300))/V_MS,1e-300))
    tr1=float(np.sqrt(np.mean((m1[tailm]-lc[tailm])**2))) if tailm.sum() else float('nan')
    rp_rms=float(np.sqrt(np.mean((lS-rlog)**2)))
    print(f"{tab+' '+letter:16}{rs:>6.1f}{kr*(L/eng.v):>9.3f}{100*(1-a_s):>7.1f}{a_m:>7.3f}{fx:>8.4f}{tr1:>9.3f}{rp_rms:>8.3f}")
    col=("#e07b00" if med=="glass" else "#7b2d8b"); aL,aR=ax[i]
    aL.plot(pv,lc,"ko",ms=5,zorder=5,label="data")
    aL.plot(r0['tp'],np.log10(r0['C']),"--",color="grey",lw=1.6,label="k_f only: no tail")
    aL.plot(r1['tp'],np.log10(r1['C']),"-",color=col,lw=2.2,label=f"+{100*(1-a_s):.1f}% MI (k_r {kr*(L/eng.v):.2f}/PV)")
    aL.set_xlim(0,12);aL.set_ylim(-6.2,0.4);aL.set_xlabel("pore volumes");aL.set_ylabel("log10 C/C0")
    aL.set_title(f"{tab} {letter}: {med} {size}um {vel} m/d - BTEC (r_s={rs:.0f}/m)",fontsize=9);aL.legend(fontsize=7,loc="upper right")
    aR.plot(rx*100,rlog,"ko",ms=5,zorder=5,label="data")
    aR.plot(r1['x']*100,logK+np.log10(np.maximum(r1['rp']/V_MS,1e-300)),"-",color=col,lw=2.2,label="Serial-3 (fit)")
    aR.set_xlim(0,20);aR.set_xlabel("distance (cm)");aR.set_ylabel("log10 spheres")
    aR.set_title(f"{tab} {letter}: RP (rms {rp_rms:.3f})",fontsize=9);aR.legend(fontsize=7,loc="upper right")
fig.suptitle("FAVORABLE - Serial-3, k_r PINNED (per-medium); r_s, MI fraction (1-a_s) & a_m fit to BTEC + RP (level+slope). Does the sustained tail survive?",fontsize=10)
fig.tight_layout(rect=[0,0,1,0.985]);fig.savefig("../Manuscript/FigsExcelsFavorable/favorable_tail_clean.png",dpi=125);print("wrote favorable_tail_clean.png")
